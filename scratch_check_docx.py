import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('MATATAG_Values_Education_Unit_3_Textbook (1).docx')

# Compare cover prompt in DOCX vs what markdown says
print("=== DOCX Table 1 (Cover Image Box) - Full AI Prompt ===")
cover_table = doc.tables[1]
for r in cover_table.rows:
    for c in r.cells:
        print(c.text)

print("\n\n=== DOCX Table 8 (L15 Figure 1) - Full AI Prompt ===")
t8 = doc.tables[8]
for r in t8.rows:
    for c in r.cells:
        print(c.text)
