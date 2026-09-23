# -*- coding: utf-8 -*-
"""Front matter builder: Pages 1, 2, and 3"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from generate_expanded_ebook import (
    format_paragraph, add_heading_1, add_heading_2, add_heading_3,
    add_body_p, add_prompt_box, add_callout_box, add_custom_table, set_cell_shading
)

def build_frontmatter(doc):
    # =========================================================================
    # PAHINA 1: PABALAT NG E-BOOK (FRONT COVER PAGE) - WALANG LEKTURA
    # =========================================================================
    p_deped = doc.add_paragraph()
    format_paragraph(p_deped, space_before=10, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_deped = p_deped.add_run("REPUBLIKA NG PILIPINAS • KAGAWARAN NG EDUKASYON\nDEPED MATATAG CURRICULUM • FILIPINO 7")
    r_deped.font.name = 'Arial'
    r_deped.font.size = Pt(11)
    r_deped.font.bold = True
    r_deped.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_title = doc.add_paragraph()
    format_paragraph(p_title, space_before=24, space_after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_title = p_title.add_run("GABAY PAMPAGKATUTO SA FILIPINO 7")
    r_title.font.name = 'Cambria'
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_sub = doc.add_paragraph()
    format_paragraph(p_sub, space_before=4, space_after=24, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_sub = p_sub.add_run("YUNIT II: MGA TULUYANG PANITIKAN AT TEKSTONG EKSPOSITORI SA PAGHUBOG NG MAKABAYANG KAMALAYAN AT MAPANURING PANANAW")
    r_sub.font.name = 'Cambria'
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    cover_prompt = (
        "PROMPT: A breathtaking, high-detail editorial book cover illustration capturing Philippine cultural heritage "
        "and 21st-century modern learning. In the foreground, an ancient giant Balete tree with majestic sprawling roots "
        "and warm golden sunlight filtering through the canopy. Filipino Grade 7 students (one boy and one girl in neat "
        "modern school uniforms) sitting together under the shade, engaged in thoughtful academic discussion while holding "
        "an illuminated digital tablet displaying indigenous Baybayin script and cultural folklore illustrations. "
        "In the background, a scenic Philippine rural landscape blending harmoniously with a modern, eco-friendly community "
        "library and distant misty mountains. Vibrant tropical color palette of emerald greens, warm ochre, and deep twilight "
        "navy blues. Style: Premium modern digital oil painting, highly detailed textures, warm cinematic lighting, artistic realism, "
        "8k resolution, suitable for a national educational textbook cover --ar 8.5:11 --v 6.0"
    )
    add_prompt_box(doc, "PROMPT SA PAGBUO NG LARAWAN SA PABALAT (FRONT COVER AI IMAGE PROMPT)", cover_prompt)

    p_meta = doc.add_paragraph()
    format_paragraph(p_meta, space_before=36, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_meta = p_meta.add_run("CERE Book Series • Edisyong Pampagkatuto at Pangguro\nNakaayon sa MATATAG 21st Century Literacy Framework")
    r_meta.font.name = 'Arial'
    r_meta.font.size = Pt(10)
    r_meta.font.bold = True
    r_meta.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    doc.add_page_break()

    # =========================================================================
    # PAHINA 2: PAUNANG SALITA, KARAPATANG-SIPI, AT PEDAGOHIKAL NA BALANGKAS
    # =========================================================================
    add_heading_1(doc, "PAUNANG SALITA AT PEDAGOHIKAL NA BALANGKAS")
    add_body_p(
        doc,
        "Malugod na pagtanggap sa mga mag-aaral at kapuwa mga gurong gagamit ng Gabay Pampagkatuto na ito para sa Ikapitong Baitang ng Bagong MATATAG Kurikulum! Ang aklat na ito ay maingat na binuo upang magsilbing tulay sa pagitan ng ating mayamang pamanang kultural at ng mga kasanayang pang-ika-21 siglo na kailangang-kailangan ng bawat kabataang Pilipino.",
        bold_prefix="Isang Pagpupugay sa Karunungan ng Bayan: "
    )
    add_body_p(
        doc,
        "Sa ilalim ng pinalakas na MATATAG Kurikulum, ang pagkatuto sa Filipino ay lumalagpas sa tradisyonal na pagsasaulo ng mga kahulugan at elemento ng panitikan. Sa halip, binibigyang-diin ang integratibong pagkatuto kung saan ang panitikan ay tinitingnan bilang buhay na salamin ng sosyo-kultural at pangkasaysayang konteksto ng ating pamayanan. Hindi lamang tayo nagbabasa ng kuwento; sinusuri natin ang mga kapangyarihang humubog dito, ang mga pananaw ng ating mga ninuno, at ang aral na maiaambag nito sa ating kasalukuyang lipunan."
    )
    add_body_p(
        doc,
        "Ang Yunit II ay nakatuon sa 'Mga Tuluyang Panitikan at Tekstong Ekspositori.' Sa pamamagitan ng masusing pagbasa sa maikling kuwento, alamat, at kuwentong-bayan, matutuklasan ng mag-aaral kung paano ginagamit ng mga manunulat ang wika upang maglahad ng katotohanan, magpatibay ng pananaw, at maglinaw ng mga komplikadong kaisipan. Kaalinsabay nito, sasanayin ang mga mag-aaral sa paggamit ng mga transisyong gramatikal, retorikal na pang-ugnay, at kohesiyong gramatikal (anapora at katapora) upang maging matatas, organisado, at lohikal sa pagsulat ng sariling mga teksto."
    )
    add_body_p(
        doc,
        "Espesyal na tampok ng e-book na ito ang pagiging interaktibo at multimedia-ready. Bawat aralin ay nagtataglay ng mga piling video documentary links mula sa mga mapagkakatiwalaang pampublikong tsanel tulad ng DepEd TV, Knowledge Channel, at PTV, na may kalakip na QR Code representations, timestamps, at mga gabay sa mapanuring panonood. Bukod dito, ang bawat visual na elemento ay may katapat na AI Image Generation Prompt upang mabigyang-daan ang visual literacy at malikhaing paglalapat ng teknolohiya sa silid-aralan."
    )
    add_body_p(
        doc,
        "Ang aklat na ito ay sumusunod sa gintong pamantayan ng '70% Malalim na Aralin at 30% Mapanuring Gawain.' Tinitiyak nito na bago sumabak ang mag-aaral sa mga pagsasanay at pagsusulit, taglay na nila ang malawak na teoretikal at konseptwal na pag-unawa. Mayroon ding kompletong Susi sa Pagwawasto sa dulo ng aklat na nagbibigay ng masusing paliwanag sa bawat sagot, upang maging katuwang ng guro sa epektibong pagtataya at magsilbing kagamitan sa sariling pagkatuto ng mag-aaral."
    )

    p_copy = doc.add_paragraph()
    format_paragraph(p_copy, space_before=12, space_after=4)
    r_c = p_copy.add_run("Pahina ng Karapatang-Sipi: © 2026 CERE Books Publishing. Reserbado ang lahat ng karapatan. Alinsunod sa Batas Republika Blg. 8293, walang bahagi ng aklat na ito ang maaaring kopyahin nang walang nakasulat na pahintulot mula sa may-akda at tagapaglathala, maliban sa mga sipi para sa layuning pampagkatuto sa silid-aralan.")
    r_c.font.name = 'Cambria'
    r_c.font.size = Pt(8.5)
    r_c.font.italic = True
    r_c.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_page_break()

    # =========================================================================
    # PAHINA 3: DETALYADONG TALAAN NG NILALAMAN AT MATATAG CURRICULUM MATRIX
    # =========================================================================
    add_heading_1(doc, "TALAAN NG NILALAMAN AT KURIKULUM MATRIX")
    add_body_p(doc, "Nakaayon sa MATATAG Filipino 7 Curriculum Guide (Kagawaran ng Edukasyon). Ang Yunit II ay binubuo ng tatlong pangunahing aralin na nagtataglay ng 70% lektura at 30% mapanuring gawain:")

    toc_headers = ["Aralin / Paksa", "Kasanayang MATATAG (Competency Codes)", "Pahina"]
    toc_data = [
        ["PAUNANG SALITA AT BALANGKAS NG PAGKATUTO", "Pilosopiya, Pedagohiya, at Gabay sa Paggamit", "Pahina 2"],
        ["TALAAN NG NILALAMAN AT KURIKULUM MATRIX", "Struktura ng Yunit II at Talaan ng mga Aralin", "Pahina 3"],
        ["ARALIN 7: Pagbasa sa Panitikan at Pananaw ng Pamayanan", "Pag-unawa sa Binasa, Tekstong Tuluyan, at Konteksto", "Pahina 4 - 15"],
        ["  • Paksa 7.1: Ang Panitikang Tuluyan at Kahalagahan ng Konteksto", "F7PN-IIa-1: Pagsusuri sa Kontekstong Kultural at Pangkasaysayan", "Pahina 5 - 8"],
        ["  • Paksa 7.2: Paglalahad, Katotohanan, Opinyon, at Pananaw", "F7PB-IIa-2: Pagkilatis sa Katotohanan vs. Opinyon sa Teksto", "Pahina 9 - 10"],
        ["  • Paksa 7.3: Transisyong Gramatikal at Pang-ugnay na Retorikal", "F7WG-IIa-3: Paggamit ng Retorikal na Pang-ugnay sa Lohika", "Pahina 11 - 12"],
        ["  • Mga Gawain, Pagsasanay, at Formative Assessment sa Aralin 7", "Mapanuring Paglalapat, Talakayan, at Formative Quiz (15 Aytem)", "Pahina 13 - 15"],
        ["ARALIN 8: Kuwento ng Pinagmulan, Linaw ng Pagpapaliwanag", "Alamat, Tekstong Ekspositori, at Pagsusuri ng Sanggunian", "Pahina 16 - 27"],
        ["  • Paksa 8.1: Ang Alamat Bilang Tuluyang Etnograpiko", "F7PN-IIb-4: Pagsusuri sa Katutubong Epistemolohiya at Alamat", "Pahina 17 - 19"],
        ["  • Paksa 8.2: Faktuwal na Paliwanag at Media & Info Literacy (MIL)", "F7PB-IIb-5: Pagkilatis sa Sanggunian at Pagsugpo sa Fake News", "Pahina 20 - 22"],
        ["  • Paksa 8.3: Kohesiyong Gramatikal: Anapora at Katapora", "F7WG-IIb-6: Paggamit ng Reperensiya sa Pag-iwas sa Redundancy", "Pahina 23 - 24"],
        ["  • Mga Gawain, Fact-Checking Matrix, at Formative Quiz sa Aralin 8", "Etnograpikong Pagsusuri, Drills, at Formative Quiz (15 Aytem)", "Pahina 25 - 27"],
        ["ARALIN 9: Kuwento ng Pamayanan, Mensahe para sa Bayan", "Kuwentong-Bayan, Organisadong Talata, at Multimodalidad", "Pahina 28 - 39"],
        ["  • Paksa 9.1: Ang Kuwentong-Bayan Bilang Salamin ng Pamayanan", "F7PB-IIc-7: Pagsusuri sa Pagpapahalaga at Kultura ng Bayan", "Pahina 29 - 31"],
        ["  • Paksa 9.2: Arkitektura ng Organisadong Talata at Transisyon", "F7PU-IIc-8: Paghabi ng Talatang may Kaisahan at Koherensiya", "Pahina 32 - 33"],
        ["  • Paksa 9.3: Multimodal na Komunikasyon at Makatarungang Representasyon", "F7PD-IIc-9: Semiotic Modes at Pagbuwag sa Cultural Stereotypes", "Pahina 34 - 36"],
        ["  • Mga Gawain, Advocacy Plan, at Formative Quiz sa Aralin 9", "Paglikha ng Multimodal Plan at Formative Quiz (15 Aytem)", "Pahina 37 - 39"],
        ["PANGWAKAS NA PAGTATAYA SA YUNIT II (Unit Assessment)", "Komprehensibong Pagsusuri, Pagsulat, at Rubrik sa Pagganap", "Pahina 40 - 41"],
        ["PANGKALAHATANG BUOD AT GLOSARYO NG YUNIT II", "Sintesis ng Konsepto, Mapa ng Kaisipan, at Talasalitaan", "Pahina 42"],
        ["KOMPLETONG SUSI SA PAGWAWASTO (Aralin 7, 8, 9, at Yunit)", "Detalyadong Paliwanag ng Tamang Sagot at Rubrik ng Guro", "Pahina 43 - 44"]
    ]
    add_custom_table(doc, toc_headers, toc_data, col_widths=[3.0, 2.5, 1.0])

    doc.add_page_break()
