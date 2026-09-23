# -*- coding: utf-8 -*-
"""Aralin 9 Builder: Pages 28 to 39 (12 Full Pages)"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from generate_expanded_ebook import (
    format_paragraph, add_heading_1, add_heading_2, add_heading_3,
    add_body_p, add_prompt_box, add_callout_box, add_custom_table,
    add_multimedia_box, set_cell_shading
)

def build_aralin9(doc):
    # =========================================================================
    # PAHINA 28: PANIMULA, LAYUNIN, AT MULTIMEDIA CORNER #5
    # =========================================================================
    add_heading_1(doc, "ARALIN 9: Kuwento ng Pamayanan, Mensahe para sa Bayan")
    add_heading_2(doc, "MGA LAYUNIN AT BALANGKAS NG PAGKATUTO")
    add_body_p(
        doc,
        "Ang Aralin 9 ang rurok ng Yunit II kung saan pinagsasama ang kultural na pamana ng Kuwentong-Bayan at ang makabagong kasanayan sa Multimodal na Komunikasyon at Makatarungang Representasyon. Sa araling ito, susuriin kung paano nagiging salamin ng kolektibong mithiin at sistemang panlipunan ang mga kuwentong-bayan. Matututuhan din ng mga mag-aaral ang arkitektura ng pagbuo ng organisadong talata na may mahigpit na kaisahan at koherensiya, at ang pagsusuri sa iba't ibang semiotic modes (teksto, visual, audio, espasyal, at kilos) upang tiyaking walang diskriminasyon o stereotype sa paglalarawan sa mga katutubo, kababaihan, at marhinalisadong sektor ng lipunan."
    )

    add_body_p(
        doc,
        "1. F7PB-IIc-7: Nasusuri ang kuwentong-bayan bilang salamin ng tradisyon, pagpapahalaga, at pamumuhay ng pamayanan sa kontekstong panlipunan at kultural.\n"
        "2. F7PU-IIc-8: Nakabubuo ng organisadong mga talata na nagtataglay ng paksang pangungusap, sumusuportang kaisipan, at epektibong transisyon tungkol sa isang napapanahong isyung panlipunan.\n"
        "3. F7PD-IIc-9: Nasusuri ang mga elementong multimodal (visual, tekstwal, audio) sa mga kagamitang pampahayag at natitiyak ang makatarungang representasyon nang walang stereotyping.\n"
        "4. F7EP-IIc-10: Nakabubuo ng isang plano para sa multimodal advocacy campaign na nagtataguyod ng pagkakaisa at paggalang sa pagkakaiba-iba ng kultura sa Pilipinas.",
        bold_prefix="Mga Kasanayang Pampagkatuto (MATATAG Competencies):\n"
    )

    add_body_p(
        doc,
        "• Paano sumasalamin sa kuwentong-bayan ang mga pangarap, hinanakit, at paraan ng pamamahala ng ating mga ninuno?\n"
        "• Ano ang semiotic modes sa multimodal na komunikasyon at paano ito nakaiimpluwensiya sa damdamin ng mambabasa?\n"
        "• Bakit mahalagang buwagin ang mga 'cultural stereotypes' sa media at paano makalilikha ng makatarungang representasyon?",
        bold_prefix="Mga Susing Katanungan para sa Aralin:\n"
    )

    add_multimedia_box(
        doc,
        mod_title="Kuwentong-Bayan at Multimodal Advocacy",
        vid_title="Ang Tradisyon ng Kuwentong-Bayan at Makatarungang Representasyon sa Bagong Media",
        channel="DepEd Filipino 7 Series / Cultural Center of the Philippines (CCP)",
        link="https://www.youtube.com/watch?v=MATATAG_Fil7_Aralin9_Bayan",
        qr_code_text="QR-FIL7-ARALIN9-MOD5",
        timestamps=[
            "00:40 - 04:10: Ang Kuwentong-Bayan bilang Katutubong Paraan ng Hustisya at Pamamahala.",
            "04:11 - 08:25: Pagsusuri sa mga Stereotype sa mga Lumang Patalastas at Palabas sa Telebisyon.",
            "08:26 - 12:30: Pagbuo ng Mapanuring Multimodal Campaign (Infographics at Video Advocacy)."
        ],
        questions=[
            "Ayon sa bidyo, paano nagiging 'buhay na aklat ng batas' ang kuwentong-bayan para sa isang pamayanang walang nasusulat na kodigo?",
            "Ano ang mga panganib na idinudulot ng stereotyping sa mga katutubong pangkat kapag ipinapakita sila sa media na laging atrasado o kakaiba?",
            "Paano mo magagamit ang iyong cellphone at social media upang magbahagi ng mga kuwentong nagtataguyod ng pantay na pagtingin sa lahat ng uri ng tao?"
        ]
    )

    doc.add_page_break()

    # =========================================================================
    # PAHINA 29: PAKSA 9.1 - ANG KUWENTONG-BAYAN BILANG SALAMIN NG NAYON
    # =========================================================================
    add_heading_2(doc, "Paksa 9.1: Ang Kuwentong-Bayan Bilang Salamin ng Pamayanan")
    add_body_p(
        doc,
        "Ang kuwentong-bayan (folklore) ay isa sa pinakamatatandang anyo ng panitikang tuluyan sa Pilipinas. Hindi tulad ng alamat na nakapokus sa pinagmulan ng mga bagay, ang kuwentong-bayan ay nakatuon sa mismong pang-araw-araw na buhay, pakikipagkapuwa, pagpapahalaga, at mga suliraning kinakaharap ng mga karaniwang mamamayan sa isang komunidad."
    )
    add_body_p(
        doc,
        "Sa pamamagitan ng kuwentong-bayan, naipapasa mula sa isang henerasyon patungo sa susunod ang mga unibersal na aral tungkol sa katapatan, kasipagan, paggalang sa mga pinuno, at ang kaparusahan sa mga mang-aapi at mapagsamantala. Ang bawat rehiyon sa Pilipinas—mula sa Batanes hanggang Tawi-Tawi—ay may sariling kaban ng mga kuwentong-bayan na nagpapakilala sa kanilang heograpiya, hanapbuhay, at panlipunang istruktura."
    )
    add_body_p(
        doc,
        "1. Kuwentong Katatawanan at Tusong Tauhan (Trickster Tales): Mga kuwento tungkol sa mga tauhang tulad ni Pilandok o Juan Tamad na gumagamit ng talino o katalasan ng isip upang madaig ang mga makapangyarihan o dambuhala.\n"
        "2. Pabula (Fables): Mga kuwentong gumagamit ng mga hayop bilang tauhan upang magturo ng moral na aral tungkol sa lipunan.\n"
        "3. Kuwento ng Karunungan at Pamamahala (Wisdom and Governance Tales): Mga naratibo tungkol sa tamang paggamit ng kapangyarihan, pagkakapantay-pantay, at pag-iingat sa kaban ng komunidad.",
        bold_prefix="Mga Uri ng Kuwentong-Bayan ayon sa Tema:\n"
    )

    prompt_p29 = (
        "PROMPT: A rich, culturally authentic scene of a pre-colonial Philippine village assembly (pulong ng nayon). "
        "Under the shade of a magnificent ancestral Banyan tree, villagers of diverse ages sit on handwoven mats. "
        "The respected village datu and female babaylan are listening intently to an honest young farmer presenting "
        "a communal harvest offering from an ornate hand-carved wooden chest (kaban). Warm dappled sunlight, detailed woven textiles "
        "(inabel/t'nalak-inspired patterns), bamboo architecture with nipa thatch roofs in the background. "
        "High artistic realism, narrative depth, vivid historical accuracy, 8k resolution --ar 16:9 --v 6.0"
    )
    add_prompt_box(doc, "PROMPT SA PAGBUO NG LARAWAN: ANG PAGPUPULONG NG NAYON SA SILONG NG PUNONG GABAY", prompt_p29)

    add_body_p(
        doc,
        "Sa pagsusuri sa kuwentong-bayan, tinitingnan natin kung paano ipinaglalaban ng mga marhinalisadong tauhan ang kanilang karapatan laban sa mga abusadong pinuno. Ipinakikita rito na bago pa man dumating ang mga Kanluraning ideya ng demokrasya, ang ating mga ninuno ay mayroon nang umiiral na sistema ng pampublikong konsultasyon at pananagutan sa kapuwa."
    )

    doc.add_page_break()

    # =========================================================================
    # PAHINA 30: KUWENTONG-BAYAN LUNSARAN - ANG KABAN NG BAYAN
    # =========================================================================
    add_heading_2(doc, "Akdang Pampanitikan para sa Aralin 9: Kuwentong-Bayan Lunsaran")
    add_heading_3(doc, "ANG PAGPUPULONG NG NAYON AT ANG KABAN NG BAYAN")
    add_body_p(
        doc,
        "Noong unang panahon sa Nayon ng Pinagsanghan, ang mga mamamayan ay nabubuhay sa pamamagitan ng pagtutulungan. Mayroon silang isang tradisyon na tinatawag na 'Bayanihan sa Palay.' Bawat magsasaka na umani ay nag-aambag ng tatlong takal ng butil sa isang malaking baul na yari sa matigas na kahoy ng molave—ang tinatawag na 'Kaban ng Bayan.' Ang kaban na ito ay inilalagay sa gitna ng Liwasan ng Nayon at binabantayan ng pinunong hinirang ng taumbayan.",
        italic_prefix="Panimula: "
    )
    add_body_p(
        doc,
        "Ang layunin ng Kaban ng Bayan ay napakadakila: kapag may dumating na bagyo, tagtuyot, o kapag may isang pamilyang nasunugan, binubuksan ang kaban sa harap ng lahat upang mabigyan sila ng pagkain at binhi nang walang anumang bayad. Sa ganitong paraan, walang sinumang taganayon ang nagugutom o naiiwan sa gitna ng sakuna."
    )
    add_body_p(
        doc,
        "Subalit nang pumanaw ang matandang pinuno, nahalal bilang bagong tagapangalaga ang isang tusong mangangalakal na nagngangalang Kabisang Mateo. Si Mateo ay magaling magsalita at nangakong dodoblehin ang laman ng kaban. Ngunit pagkalipas ng ilang buwan, napansin ng mga mamamayan na unti-unting lumiliit ang bahagi ng butil na ibinabalik sa mga nangangailangan. Tuwing may hihingi ng tulong, sinasabi ni Mateo: 'Ubos na ang kaban dahil sa mga gastusin sa pamamahala.' Samantala, ang sariling kamalig ni Mateo ay sumasabog sa dami ng sako ng palay.",
        italic_prefix="Tunggalian: "
    )
    add_body_p(
        doc,
        "Isang araw ng malakas na ulan, isang matandang balo na si Tandang Maria ang lumapit kay Mateo upang humingi ng kahit isang salop na bigas para sa kaniyang maysakit na apo. Ngunit tinaboy lamang siya ni Mateo at sinabing tamad ang kaniyang pamilya. Nasaksihan ito ng isang binatang magsasaka na nagngangalang Maypag-asa. Hindi siya pumayag sa ganitong kawalang-katarungan."
    )
    add_body_p(
        doc,
        "Tinugtog ni Maypag-asa ang tambuli ng nayon upang tawagin ang lahat sa isang pangkalahatang pagpupulong sa ilalim ng lumang puno. Sa harap ng buong nayon, buong-tapang na inihayag ni Maypag-asa ang kaniyang pagsusuri: 'Mga kababayan, ang Kaban ng Bayan ay hindi pag-aari ng iisang tao! Ito ay pinagpawisan ng bawat isa sa atin! Nararapat lamang na buksan ang baul ngayon din at bilangin ang bawat butil sa harap ng araw!'",
        italic_prefix="Kasukdulan: "
    )
    add_body_p(
        doc,
        "Nang buksan ang baul sa tulong ng mga nakatatanda, napatunayan na walang laman ang Kaban ng Bayan maliban sa mga buhangin na inilagay ni Mateo upang magmukhang mabigat. Hiyang-hiya at nahubaran ng maskara si Mateo. Sa bisa ng pagkakaisa ng nayon, binawi nila ang kanilang mga inambag mula sa kamalig ni Mateo, pinalayas siya sa nayon, at itinatag ang isang bagong patakaran: ang Kaban ng Bayan ay pamamahalaan na ng tatlong magkakaibang sektor (kabataan, kababaihan, at magsasaka) na may bukas na talaan para sa lahat.",
        italic_prefix="Kakalasan at Wakas: "
    )

    doc.add_page_break()

    # =========================================================================
    # PAHINA 31: MAPANURING PAGSUSURI AT SOCRATIC SEMINAR SA KUWENTONG-BAYAN
    # =========================================================================
    add_heading_2(doc, "Mapanuring Pagsusuri at Socratic Seminar sa Kuwentong-Bayan")
    add_body_p(
        doc,
        "Panoorin at basahin ang talakayan ng guro at mag-aaral na naglalatag ng malalim na sosyolohikal na interpretasyon sa kuwento:"
    )

    add_body_p(
        doc,
        "Guro: Magandang araw! Sa kuwentong 'Ang Kaban ng Bayan,' ano ang kinakatawan ng baul ng palay sa konteksto ng ating modernong gobyerno at lipunan?",
        bold_prefix="Guro: "
    )
    add_body_p(
        doc,
        "Daria (Mag-aaral): Ang Kaban ng Bayan po ay sumasagisag sa 'Taxes' o Buwis ng mga mamamayan at sa 'Pambansang Badyet.' Ang bawat butil ng palay ay kumakatawan sa pawis at pagod ng mga manggagawa na inilalaan para sa pampublikong serbisyo tulad ng ospital, kalsada, at edukasyon.",
        bold_prefix="Daria (Mag-aaral): "
    )
    add_body_p(
        doc,
        "Guro: Napakagaling na pagsusuri, Daria! Joshua, ano naman ang ipinakita ng karakter ni Kabisang Mateo tungkol sa panganib ng korapsiyon sa pamumuno?",
        bold_prefix="Guro: "
    )
    add_body_p(
        doc,
        "Joshua (Mag-aaral): Ipinakikita po nito na kapag ang isang lider ay walang 'transparency' at nag-iisa sa pamamahala ng pondo, madali siyang matukso ng kasakiman. Ang paglalagay niya ng buhangin sa baul ay sumasagisag sa 'ghost projects' o panlilinlang sa taumbayan.",
        bold_prefix="Joshua (Mag-aaral): "
    )
    add_body_p(
        doc,
        "Guro: Tumpak! At ano ang naging pinal na solusyon ng pamayanan na nagpapakita ng isang matatag na sistemang demokratiko?",
        bold_prefix="Guro: "
    )
    add_body_p(
        doc,
        "Daria: Hindi lamang po nila pinalitan ang lider; binago nila ang 'sistema.' Hinati nila ang pamamahala sa tatlong sektor (checks and balances) at ginawang 'open book' o bukas sa publiko ang talaan ng Kaban ng Bayan. Ito po ang diwa ng mabuting pamamahala (good governance).",
        bold_prefix="Daria: "
    )

    add_callout_box(
        doc,
        title="PAGLALAPAT SA PAGKAMAMAMAYAN: TRANSPARENCY AT ACCOUNTABILITY",
        body_lines=[
            "1. Transparency (Bukas na Pamamahala): Ang karapatan ng bawat mamamayan na malaman kung saan napupunta ang pampublikong pondo ng barangay at bayan.",
            "2. Accountability (Pananagutan): Ang pananagutan ng bawat pinuno na magpaliwanag at harapin ang kaparusahan sa batas kung siya ay nagtaksil sa tiwala ng bayan.",
            "3. People's Participation (Pakikilahok ng Mamamayan): Ang paggising sa kamalayan ng kabataan na hindi dapat manatiling tahimik kapag may nakikitang inhustisya."
        ],
        color_fill="F4F7FA",
        border_color="1B365D"
    )

    doc.add_page_break()

    # =========================================================================
    # PAHINA 32: PAKSA 9.2 (BAHAGI 1) - ARKITEKTURA NG ORGANISADONG TALATA
    # =========================================================================
    add_heading_2(doc, "Paksa 9.2: Arkitektura ng Organisadong Talata at Transisyon (Bahagi 1)")
    add_body_p(
        doc,
        "Ang isang maayos na sulatin ay maihahalintulad sa isang matibay na bahay-kubo o gusali: kailangan nito ng matibay na pundasyon, matitigas na haligi, at maayos na bubong. Sa pagsulat ng tekstong ekspositori, ang bawat talata ay may tiyak na arkitektura na dapat sundin upang hindi maligaw ang mambabasa."
    )

    add_body_p(
        doc,
        "1. Paksang Pangungusap (Topic Sentence): Ito ang utak o kapitan ng talata. Karaniwan itong matatagpuan sa simula ng talata at naglalaman ng pangunahing ideya o sentral na argumento na tatalakayin.\n"
        "2. Mga Pansuportang Pangungusap (Supporting Details): Ito ang mga sundalo na nagpapatibay sa kapitan. Naglalaman ito ng mga paliwanag, tiyak na halimbawa, datos, at lohikal na rason na nagpapatunay sa paksang pangungusap.\n"
        "3. Pangwakas o Transisyonal na Pangungusap (Concluding / Transition Sentence): Ito ang naglalagom sa natalakay at naghahanda sa mambabasa patungo sa susunod na talata.",
        bold_prefix="Tatlong Pangunahing Bahagi ng isang Huwarang Talata:\n"
    )

    para_headers = ["Bahagi ng Talata", "Tungkulin sa Diskurso", "Halimbawa mula sa Aralin"]
    para_data = [
        [
            "Paksang Pangungusap",
            "Nagsasaad ng sentral na punto ng talata.",
            "Ang Kaban ng Bayan ay hindi lamang imbakan ng palay kundi ang buhay na sagisag ng katapatan sa pamamahala."
        ],
        [
            "Pansuportang Detalye 1",
            "Nagbibigay ng rason at paliwanag.",
            "Kapag pinangalagaan ito nang may bukas na talaan, nagkakaroon ng tiwala ang mga mamamayan na mag-ambag ng kanilang butil."
        ],
        [
            "Pansuportang Detalye 2",
            "Nagbibigay ng kongkretong ebidensiya/halimbawa.",
            "Tulad ng ipinakita sa kasaysayan ng San Jose, naging ligtas ang mga pamilya sa taggutom dahil sa tapat na pamamahagi."
        ],
        [
            "Pangwakas na Pangungusap",
            "Naglalagom at nagpapatibay sa posisyon.",
            "Samakatuwid, ang tunay na sukatan ng isang maunlad na nayon ay ang kalinisan ng kaban at dangal ng kaniyang mga pinuno."
        ]
    ]
    add_custom_table(doc, para_headers, para_data, col_widths=[1.8, 2.2, 2.5])

    add_body_p(
        doc,
        "Tandaan: Ang isang epektibong talata ay dapat magtaglay ng 'Kaisahan' (Unity)—ibig sabihin, lahat ng pangungusap ay tumatalakay lamang sa IISANG paksang pangungusap. Kung may nais kang ipasok na bagong ideya, nararapat lamang na lumikha ng bagong talata."
    )

    doc.add_page_break()

    # =========================================================================
    # PAHINA 33: PAKSA 9.2 (BAHAGI 2) - KOHERENSIYA SA PAGITAN NG MGA TALATA
    # =========================================================================
    add_heading_2(doc, "Paksa 9.2: Koherensiya sa Pagitan ng mga Talata (Bahagi 2)")
    add_body_p(
        doc,
        "Hindi sapat na maayos ang bawat indibidwal na talata; kailangan ding maging madulas ang pagkakadugtong-dugtong ng mga talata sa loob ng buong sanaysay. Ito ang tinatawag nating 'Macro-level Coherence.'"
    )

    add_body_p(
        doc,
        "Upang magkaroon ng macro-level coherence, gumagamit tayo ng tatlong pangunahing teknik:\n"
        "1. Paggamit ng Transisyonal na Panimula sa Talata: Simulan ang bagong talata sa pamamagitan ng pagbanggit sa kaisipang tinalakay sa dulo ng naunang talata (hal. 'Kaugnay ng nabanggit na suliranin sa kaban, isa pang hamon ang kinakaharap ng nayon...').\n"
        "2. Paggamit ng Kronolohikal o Lohikal na Pagsusunod-sunod: Ayusin ang mga talata ayon sa kahalagahan (mula sa pinakamahalaga patungo sa detalye, o mula sa problema patungo sa solusyon).\n"
        "3. Pagpapanatili ng Matatag na Tono: Panatilihin ang pormal at mapanuring tono sa kabuuan ng sulatin.",
        bold_prefix="Mga Teknik sa Pagdurugtong ng mga Talata:\n"
    )

    add_callout_box(
        doc,
        title="HUWARANG BALANGKAS NG ISANG MAIKLING POSISYONG PAPEL",
        body_lines=[
            "• Talata 1 (Introduksiyon): Paglalahad ng Isyu + Konteksto + Thesis Statement (Pangunahing Pananaw).",
            "• Talata 2 (Unang Katwiran): Unang Haligi ng Pananaw na may Paliwanag at Ebidensiya.",
            "• Talata 3 (Ikalawang Katwiran): Pagsusuri sa Salungat na Pananaw (Counter-argument) at Pagpapabulaan Dito.",
            "• Talata 4 (Konklusyon at Pagkilos): Paglalagom sa mga Katwiran at Malinaw na Panawagan sa Pamayanan."
        ],
        color_fill="F4F7FA",
        border_color="1B365D"
    )

    add_body_p(
        doc,
        "Pansinin na sa balangkas na ito, bawat talata ay may malinaw na gampanin. Kapag nabasa ito ng kahit sinong mamamayan o opisyal ng pamahalaan, madali nilang mauunawaan ang iyong katwiran at hindi sila maliligaw sa iyong diskurso."
    )

    doc.add_page_break()

    # =========================================================================
    # PAHINA 34: PAKSA 9.3 (BAHAGI 1) - MULTIMODAL NA KOMUNIKASYON
    # =========================================================================
    add_heading_2(doc, "Paksa 9.3: Multimodal na Komunikasyon at Makatarungang Representasyon")
    add_body_p(
        doc,
        "Sa ika-21 siglo, hindi na lamang tayo nagbabasa ng purong nakalimbag na teksto. Araw-araw, nakikipag-ugnayan tayo sa mga 'Multimodal Texts'—mga komunikasyong pinagsasama ang dalawa o higit pang semiotic modes upang maghatid ng mas malalim at mabisang mensahe. Ang salitang 'semiotic' ay tumutukoy sa pag-aaral ng mga tanda, simbolo, at kahulugan."
    )

    modes_headers = ["Semiotic Mode", "Mga Elementong Kinabibilangan", "Halimbawa sa Araw-araw na Media"]
    modes_data = [
        [
            "1. Linguistic (Tekstwal)",
            "Bokabularyo, gramatika, estruktura ng pangungusap, metapora.",
            "Mga artikulo, caption, subtitle, pamagat ng balita."
        ],
        [
            "2. Visual (Biswal)",
            "Kulay, liwanag, anggulo ng kamera, linya, porma, layout, typography.",
            "Larawan sa poster, infographics, memes, thumbnail sa YouTube."
        ],
        [
            "3. Audio (Pandinig)",
            "Musika, sound effects, tono ng boses, bilis ng pagsasalita, katahimikan.",
            "BGM sa podcast, voiceover sa balita, tunog ng kulog sa bidyo."
        ],
        [
            "4. Gestural (Kilos at Galaw)",
            "Kumpas ng kamay, ekspresyon ng mukha, tindig, eye contact.",
            "Galaw ng host sa vlog, reaksiyon ng aktor sa maikling pelikula."
        ],
        [
            "5. Spatial (Espasyal)",
            "Layo ng mga bagay, pag-aayos ng elements sa pahina o screen.",
            "Disenyo ng website, espasyo ng mga text boxes sa brochure."
        ]
    ]
    add_custom_table(doc, modes_headers, modes_data, col_widths=[1.8, 2.3, 2.4])

    add_body_p(
        doc,
        "Bakit Napakalakas ng Multimodal Texts? Ayon sa mga pag-aaral sa sikolohiya ng pagkatuto, ang utak ng tao ay mas mabilis magproseso ng mga impormasyong pinagsama ang teksto at biswal. Subalit dahil napakalakas ng kapangyarihan nito, madali rin itong gamitin sa pagmamanipula ng kaisipan ng tao. Dito pumapasok ang napakahalagang konsepto ng 'Makatarungang Representasyon.'"
    )

    doc.add_page_break()

    # =========================================================================
    # PAHINA 35: PAKSA 9.3 (BAHAGI 2) - MAKATARUNGANG REPRESENTASYON
    # =========================================================================
    add_heading_2(doc, "Paksa 9.3: Makatarungang Representasyon at Pagbuwag sa Stereotype (Bahagi 2)")
    add_body_p(
        doc,
        "Ang 'Makatarungang Representasyon' (Fair Representation) ay ang patas, makatotohanan, at may paggalang na paglalarawan sa lahat ng tao sa media at panitikan—anuman ang kanilang etnisidad, relihiyon, kasarian, edad, o katayuan sa lipunan. Kaakibat nito ang pagbuwag sa 'Stereotyping.'"
    )

    add_body_p(
        doc,
        "Ang stereotype ay isang pinasimple, labis na pangkalahatan, at madalas ay maling pagtingin o paglalarawan sa isang partikular na grupo ng tao. Halimbawa:\n"
        "• Stereotype: Kapag ipinapakita ang mga katutubong Aeta o Badjao sa telebisyon, laging madudumi, namamalimos, o walang pinag-aralan ang kanilang papel.\n"
        "• Katotohanan: Ang ating mga katutubo ay may sariling mayamang kultura, propesyonal na mga doktor at guro, at mataas na karunungan sa pangangalaga ng gubat at dagat.\n"
        "• Stereotype: Ang paglalarawan sa mga matatanda bilang laging makukulit, mahihina, o sagabal sa kaunlaran.\n"
        "• Katotohanan: Ang mga nakatatanda ang bukal ng karunungan at nagbibigay ng moral na direksiyon sa pamayanan.",
        bold_prefix="Ano ang Stereotyping at Bakit ito Mapanganib?\n"
    )

    prompt_p35 = (
        "PROMPT: An empowering, dignified, and culturally authentic editorial illustration representing Filipino indigenous youth "
        "and modern empowerment. An indigenous Lumad / Cordilleran young woman wearing modern academic attire tastefully accented "
        "with authentic handwoven textile scarves (hablon/inabel). She is confidently presenting a modern renewable solar energy plan "
        "to a diverse community forum in an eco-friendly municipal hall. Beside her are diverse Filipino students taking notes on digital tablets. "
        "Bright natural daylight, dignified expressions, celebrating cultural diversity and academic excellence without any stereotypes. "
        "Master digital illustration, realistic proportions, high-end editorial graphic art, 8k resolution --ar 16:9 --v 6.0"
    )
    add_prompt_box(doc, "PROMPT SA PAGBUO NG LARAWAN: MAKATARUNGANG REPRESENTASYON NG KABATAANG KATUTUBO", prompt_p35)

    add_callout_box(
        doc,
        title="GABAY SA PAGTIYAK NG MAKATARUNGANG REPRESENTASYON",
        body_lines=[
            "1. Iwasan ang Pagsasawalang-bahala (No Erasure): Isama ang boses ng mga katutubo, kababaihan, at PWD sa iyong mga naratibo.",
            "2. Iwasan ang Tokenism: Huwag isama ang isang karakter para lamang masabing 'diverse' ang iyong kuwento; bigyan sila ng totoong papel at lalim.",
            "3. Irespeto ang Sagradong Kasuotan: Huwag gamitin ang mga tradisyunal na kasuotang katutubo bilang 'costume' o katatawanan sa mga palabas."
        ],
        color_fill="F4F7FA",
        border_color="1B365D"
    )

    doc.add_page_break()

    # =========================================================================
    # PAHINA 36: MULTIMEDIA CORNER #6 AT ADVOCACY CASE ANALYSIS
    # =========================================================================
    add_heading_2(doc, "Multimedia Corner #6: Multimodal Advocacy at Etikal na Media")
    
    add_multimedia_box(
        doc,
        mod_title="Makatarungang Representasyon sa Bagong Media",
        vid_title="Pagbuwag sa Stereotypes: Makatarungang Representasyon ng mga Katutubo at Marhinalisadong Pangkat sa Media",
        channel="DepEd Literacy Program / Commission on Human Rights (CHR) Education",
        link="https://www.youtube.com/watch?v=MATATAG_Fil7_Aralin9_Media",
        qr_code_text="QR-FIL7-ARALIN9-MOD6",
        timestamps=[
            "00:30 - 03:50: Pagsusuri sa mga Maling Representasyon sa mga Tradisyunal na Pelikula at Palabas.",
            "03:51 - 07:40: Paano Gumawa ng Advocacy Poster na may Paggalang at Dangal.",
            "07:41 - 11:15: Ang Papel ng Kabataan sa Paglikha ng Inklusibong Online Spaces."
        ],
        questions=[
            "Paano inilarawan sa bidyo ang tinatawag na 'cultural caricature' at bakit ito nakasasakit sa mga katutubong komunidad?",
            "Anong mga elemento ng kulay at layout ang dapat iwasan upang hindi magmukhang kaawa-awa ang representasyon ng isang sektor?",
            "Kung gagawa ka ng isang patalastas para sa iyong paaralan, paano mo titiyakin na ang lahat ng uri ng mag-aaral ay mabibigyan ng pantay na pagkilala?"
        ]
    )

    add_heading_3(doc, "Pagsusuri sa Sitwasyon (Case Analysis): Stereotype sa Poster ng Barangay")
    add_body_p(
        doc,
        "Isang barangay poster para sa Linggo ng Kalinisan ang nagpapakita ng isang karikatura ng isang katutubong Aeta na may maruming damit at nakaturo sa tambak ng basura na may caption na: 'Huwag Maging Parang Katutubo, Maglinis ng Kapaligiran!'"
    )
    add_body_p(
        doc,
        "Mapanuring Pagsusuri:\n"
        "1. Nilabag na Prinsipyo: Tahasang Stereotyping at Diskriminasyon. Iniuugnay ang katutubo sa dumi at kawalan ng disiplina, na lubhang taliwas sa katotohanan dahil ang mga katutubo ang pinakamahusay na tagapangalaga ng kalikasan sa kasaysayan.\n"
        "2. Tamang Aksiyon: Agarang ipabawi ang poster, maglabas ng pormal na paghingi ng paumanhin, at palitan ito ng isang multimodal poster na nagpapakita ng pagkakaisa ng lahat ng mamamayan sa paglilinis nang walang paninira sa anumang lahi o sektor."
    )

    doc.add_page_break()

    # =========================================================================
    # PAHINA 37: GAWAIN 9.1 AT GAWAIN 9.2 (30% ACTIVITY)
    # =========================================================================
    add_heading_2(doc, "Bahaging Gawain at Paglalapat para sa Aralin 9 (30% Activity)")
    
    add_callout_box(
        doc,
        title="Gawain 9.1: Sosyo-Kultural na Pagsusuri sa Kuwentong-Bayan",
        body_lines=[
            "Panuto: Balikan ang akdang 'Ang Pagpupulong ng Nayon at ang Kaban ng Bayan.' Sagutin ang mga sumusunod na tanong nang may lalim at gumamit ng mga transisyong gramatikal sa iyong mga pangungusap.",
            "1. Ano ang sinasagisag ng 'buhangin' na inilagay ni Kabisang Mateo sa loob ng Kaban ng Bayan? Paano ito maihahambing sa mga anomalya sa ating lipunan ngayon?",
            "2. Bakit mahalaga ang naging desisyon ng nayon na hatiin ang pamamahala ng kaban sa tatlong magkakaibang sektor (checks and balances)?",
            "3. Kung ikaw si Maypag-asa, anong mga hakbang ang iyong gagawin upang matiyak na hindi na muling mauulit ang panlilinlang sa inyong pamayanan?"
        ],
        color_fill="F4F7FA",
        border_color="1B365D"
    )

    add_callout_box(
        doc,
        title="Gawain 9.2: Pagsasaayos at Paghabi ng Organisadong Talata",
        body_lines=[
            "Panuto: Ang sumusunod na mga pangungusap ay nagtataglay ng isang paksang pangungusap, dalawang pansuportang detalye, at isang pangwakas na pangungusap, ngunit magulo ang pagkakaayos. Ayusin ang mga ito upang makabuo ng isang lohikal at organisadong talata. Lagyan ng bilang 1 hanggang 4.",
            "_____ A. Halimbawa, napatunayan sa maraming komunidad na ang mga proyektong may bukas na ulat sa badyet ay mas mabilis natatapos at mas pinakikinabangan ng tao.",
            "_____ B. Ang pagkakaroon ng bukas na pamamahala o transparency ay pundasyon ng matatag at maunlad na pamayanan.",
            "_____ C. Samakatuwid, karapatan at tungkulin ng bawat mamamayan na magtanong at makialam sa pamamalakad ng kanilang bayan.",
            "_____ D. Ito ay sapagkat nawawalan ng puwang ang korapsiyon kapag ang bawat sentimo ng pondo ay nalalaman ng taumbayan."
        ],
        color_fill="F4F7FA",
        border_color="1B365D"
    )

    doc.add_page_break()

    # =========================================================================
    # PAHINA 38: GAWAIN 9.3 - PLANO NG MULTIMODAL ADVOCACY AT ANTI-STEREOTYPE
    # =========================================================================
    add_heading_2(doc, "Bahaging Pagsasanay at Paglalapat para sa Aralin 9")
    
    add_callout_box(
        doc,
        title="Gawain 9.3: Pagbuo ng Plano para sa Multimodal Advocacy Campaign",
        body_lines=[
            "Panuto: Ikaw ay inatasang magdisenyo ng isang 'Advocacy Infographic o Social Media Campaign' para sa inyong paaralan na nagtataguyod ng Makatarungang Representasyon at Paggalang sa Pagkakaiba-iba. Punan ang sumusunod na blueprint ng inyong proyekto."
        ],
        color_fill="F4F7FA",
        border_color="1B365D"
    )

    act9_headers = ["Elemento ng Plano", "Deskripsiyon at Detalye ng Iyong Disenyo"]
    act9_data = [
        ["Pamagat ng Kampanya", "(Mag-isip ng nakapupukaw at makabayang pamagat...)"],
        ["Target na Audience", "Mga mag-aaral sa Baitang 7 at kabataan sa komunidad."],
        ["Linguistic Mode (Teksto)", "(Isulat dito ang pangunahing slogan at 2 paksang pangungusap...)"],
        ["Visual Mode (Biswal)", "(Ilarawan ang mga kulay, larawan ng tao, at layout upang maiwasan ang stereotype...)"],
        ["Audio Mode (Kung may video/podcast)", "(Ilarawan ang uri ng musika, sound effects, at tono ng boses...)"],
        ["Anti-Stereotype Checklist", "Tiyaking ang bawat sektor ay may dignidad, maayos na anyo, at may aktibong papel."]
    ]
    add_custom_table(doc, act9_headers, act9_data, col_widths=[2.5, 4.0])

    add_heading_3(doc, "Pamantayan sa Pagmamarka ng Multimodal Advocacy Plan:")
    rubric9_headers = ["Pamantayan", "Napakahusay (5)", "Mahusay (4)", "Kailangan ng Gabay (3)", "Puntos"]
    rubric9_data = [
        ["Makatarungang Representasyon", "Ganap na walang stereotype; may mataas na paggalang sa lahat", "Walang hayagang stereotype ngunit kulang sa lalim", "May bakas ng stereotype o maling paglalarawan", "/5"],
        ["Integrasyon ng Semiotic Modes", "Mahusay na pinagsama ang teksto, visual, at audio", "May teksto at visual ngunit hindi magkatugma", "Teksto lamang, walang multimodal na elemento", "/5"],
        ["Linaw ng Mensahe at Layon", "Napakalinaw ng adbokasiya at nakapupukaw ng pagkilos", "Malinaw ang mensahe ngunit karaniwan lamang", "Magulo at mahirap intindihin ang layon", "/5"],
        ["Gramatika at Balarila", "Walang kamalian sa bantas, transisyon, at baybay", "May 1-2 bahagyang kamalian", "Maraming maling baybay at bantas", "/5"]
    ]
    add_custom_table(doc, rubric9_headers, rubric9_data, col_widths=[2.0, 1.5, 1.5, 1.5, 0.7])

    doc.add_page_break()

    # =========================================================================
    # PAHINA 39: MABILISANG PAGTATAYA SA ARALIN 9 (FORMATIVE ASSESSMENT)
    # =========================================================================
    add_heading_2(doc, "Mabilisang Pagtataya sa Aralin 9 (Formative Assessment)")
    add_body_p(
        doc,
        "Panuto: Piliin ang titik ng tamang sagot at isulat ito sa patlang bago ang bawat bilang.",
        bold_prefix="Pangkalahatang Panuto: "
    )

    quiz9_items = [
        ("1. Ano ang pangunahing layunin ng kuwentong-bayan sa isang sinaunang pamayanan?\n"
         "A. Magbenta ng mga produkto sa palengke.\n"
         "B. Magsilbing salamin ng kultura, tradisyon, at mga kodigo ng pamumuhay at hustisya sa nayon.\n"
         "C. Mang-aliw lamang ng mga dayuhang bisita.\n"
         "D. Magpalaganap ng mga batas trapiko."),

        ("2. Sa kuwentong 'Ang Kaban ng Bayan,' ano ang naging kasalanan ni Kabisang Mateo?\n"
         "A. Hindi siya marunong magbasa at magsulat.\n"
         "B. Ninakaw niya ang butil sa kaban at pinalitan ito ng buhangin upang linlangin ang taumbayan.\n"
         "C. Nasunog niya ang liwasan ng bayan.\n"
         "D. Hindi siya nagbayad ng buwis sa kaniyang lupain."),

        ("3. Aling bahagi ng talata ang nagsisilbing sentral na kaisipan at gumagabay sa buong direksiyon ng sulatin?\n"
         "A. Pansuportang Pangungusap   B. Paksang Pangungusap   C. Bantas   D. Pamagat"),

        ("4. Ano ang tawag sa komunikasyong pinagsasama ang teksto, larawan, tunog, at layout upang maghatid ng kahulugan?\n"
         "A. Monomodal Text   B. Multimodal Text   C. Purong Diksiyonaryo   D. Tradisyonal na Talumpati"),

        ("5. Alin sa mga sumusunod ang halimbawa ng 'Visual Mode' sa isang poster?\n"
         "A. Ang laki at kulay ng font, pati ang liwanag at anggulo ng larawan.\n"
         "B. Ang background music na maririnig sa bidyo.\n"
         "C. Ang kumpas ng kamay ng tagapagsalita.\n"
         "D. Ang distansiya ng upuan sa entablado."),

        ("6. Ano ang 'Stereotyping'?\n"
         "A. Ang pagkuha ng magagandang larawan gamit ang DSLR camera.\n"
         "B. Ang paglalapat ng pangkalahatan at madalas ay maling katangian sa isang buong pangkat ng tao.\n"
         "C. Ang pagsasalin ng wika mula sa Ingles patungong Filipino.\n"
         "D. Ang pagbili ng mga orihinal na aklat sa tindahan."),

        ("7. Bakit mahalaga ang Makatarungang Representasyon sa pagbuo ng mga educational media?\n"
         "A. Upang maging mas mura ang gastos sa produksiyon.\n"
         "B. Upang mabigyan ng pantay na dignidad, boses, at paggalang ang lahat ng sektor nang walang diskriminasyon.\n"
         "C. Upang sumikat sa TikTok ang gumawa ng bidyo.\n"
         "D. Upang maging mas mahaba ang oras ng klase."),

        ("8. Alin sa mga sumusunod ang nagpapakita ng 'Macro-level Coherence' sa isang sanaysay?\n"
         "A. Ang paggamit ng makukulay na papel sa pag-print.\n"
         "B. Ang maayos, madulas, at lohikal na pagdurugtong ng mga talata gamit ang mga transisyonal na kaisipan.\n"
         "C. Ang pagsulat ng napakaliit na mga titik.\n"
         "D. Ang paglalagay ng drawing sa bawat sulok ng pahina."),

        ("9. Ano ang ipinahihiwatig ng 'Checks and Balances' na ipinatupad sa Kaban ng Bayan?\n"
         "A. Ang pamamahala ay hindi dapat ipinauubaya sa iisang tao lamang kundi dapat ibahagi sa iba't ibang sektor para sa pananagutan.\n"
         "B. Kailangang laging magdala ng tseke kapag bibili ng palay.\n"
         "C. Dapat ipagbawal ang pag-aambag ng palay sa nayon.\n"
         "D. Ang kabataan lamang ang may karapatang mamahala sa bayan."),

        ("10. Paano nakatutulong ang kasanayan sa multimodalidad sa mga mag-aaral ng Baitang 7?\n"
         "A. Nagiging eksperto sila sa paglalaro ng video games buong araw.\n"
         "B. Nagiging mapanuri silang konsyumer at responsableng tagalikha ng media sa modernong lipunan.\n"
         "C. Hindi na nila kailangang mag-aral ng balarila at pagsulat.\n"
         "D. Nagiging madali para sa kanila na maniwala sa anumang nakikita online.")
    ]

    for item in quiz9_items:
        add_body_p(doc, item)

    add_callout_box(
        doc,
        title="EXIT TICKET / PAGNINILAY SA ARALIN 9",
        body_lines=[
            "1. Isang bagong pananaw na natutuhan ko tungkol sa makatarungang representasyon ay: ____________________",
            "2. Bilang mag-aaral, paano ko gagamitin ang multimodal texts upang makatulong sa aking kapuwa? ____________"
        ],
        color_fill="FAFAFA",
        border_color="CCCCCC"
    )

    doc.add_page_break()
