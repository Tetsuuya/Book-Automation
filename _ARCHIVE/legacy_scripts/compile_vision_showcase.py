# -*- coding: utf-8 -*-
"""
COMPILATION SCRIPT: MATATAG INTERACTIVE VISION SHOWCASE
Generates the publication-grade textbook implementing full-spectrum interactivity:
- Natural continuous flow (zero orphan pages or single-line spillovers)
- Bulletproof YouTube search links (never 404)
- In-text 'Pause & Ponder' (Tigil at Mag-isip) checkpoints
- 3-2-1 Metacognitive Exit Tickets
- Concept matrices, Socratic inquiry dialogues, 4x4 analytic rubrics, 15-item quizzes with [ ] brackets
"""

import os
import sys
import shutil
import urllib.parse
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

import generate_expanded_ebook
from sections_yunit3_frontmatter import build_yunit3_frontmatter
from sections_yunit3_aralin13 import build_yunit3_aralin13
from sections_yunit3_aralin14 import build_yunit3_aralin14
from sections_yunit3_aralin15 import build_yunit3_aralin15
from sections_yunit3_assessment_keys import build_yunit3_assessment_and_keys

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_callout_border(cell, color_hex="1B365D", sz="36"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

def format_paragraph(p, space_before=0, space_after=4, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align

# Monkey-patch generate_expanded_ebook.add_multimedia_box with bulletproof zero-hallucination search link
def patched_add_multimedia_box(doc, mod_title, vid_title, channel, link, qr_code_text, timestamps, questions):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    set_cell_shading(c, "FFF5F5")
    set_callout_border(c, color_hex="CC0000", sz="36")
    set_cell_margins(c, top=140, bottom=140, left=180, right=180)
    
    # 1. Header
    p0 = c.paragraphs[0]
    format_paragraph(p0, space_after=2)
    r0 = p0.add_run(f"🎥 INTERAKTIBONG MULTIMEDIA: {mod_title.upper()}")
    r0.font.name = 'Cambria'
    r0.font.size = Pt(10.5)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
    
    # 2. Metadata
    p1 = c.add_paragraph()
    format_paragraph(p1, space_after=2)
    r1 = p1.add_run(f"Pamagat ng Panoorin: \"{vid_title}\"  |  Opisyal na Tsanel: {channel}")
    r1.font.name = 'Cambria'
    r1.font.size = Pt(9.5)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    # 3. Bulletproof Search Query URL (Guaranteed to NEVER 404!)
    search_keywords = f"{channel.split('/')[0].strip()} {vid_title}".replace('"', '')
    encoded_query = urllib.parse.quote_plus(search_keywords)
    bulletproof_url = f"https://www.youtube.com/results?search_query={encoded_query}"
    
    p2 = c.add_paragraph()
    format_paragraph(p2, space_after=1)
    r2a = p2.add_run("🔗 Direktang Search Link (Garantisadong Aktibo): ")
    r2a.font.name = 'Cambria'
    r2a.font.size = Pt(9.0)
    r2a.font.bold = True
    r2b = p2.add_run(bulletproof_url)
    r2b.font.name = 'Cambria'
    r2b.font.size = Pt(9.0)
    r2b.font.color.rgb = RGBColor(0x00, 0x44, 0xCC)
    r2b.font.underline = True
    
    # 4. Plain-text Search Keywords & QR Code Frame
    p_qr = c.add_paragraph()
    format_paragraph(p_qr, space_after=2)
    r_qr = p_qr.add_run(f"🔍 YouTube Search Keywords: \"{search_keywords}\"  |  [📱 QR Code: I-scan para sa YouTube]")
    r_qr.font.name = 'Cambria'
    r_qr.font.size = Pt(8.5)
    r_qr.font.italic = True
    r_qr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    
    # 5. Timestamps
    p3 = c.add_paragraph()
    format_paragraph(p3, space_before=2, space_after=1)
    r3 = p3.add_run("Mahahalagang Bahagi at Timestamps na Pag-aaralan:")
    r3.font.name = 'Cambria'
    r3.font.size = Pt(9.5)
    r3.font.bold = True
    r3.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    for ts in timestamps:
        pts = c.add_paragraph()
        format_paragraph(pts, space_after=1)
        r_ts = pts.add_run(f"⏱️ {ts}")
        r_ts.font.name = 'Cambria'
        r_ts.font.size = Pt(9.0)
        r_ts.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        
    # 6. Socratic Viewing Questions
    p4 = c.add_paragraph()
    format_paragraph(p4, space_before=3, space_after=1)
    r4 = p4.add_run("Panonood at Mapanuring Pagsusuri (Gabay sa Mapanuring Pag-iisip):")
    r4.font.name = 'Cambria'
    r4.font.size = Pt(9.5)
    r4.font.bold = True
    r4.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    for i, q in enumerate(questions, 1):
        pq = c.add_paragraph()
        format_paragraph(pq, space_after=1)
        rq = pq.add_run(f"{i}. {q}")
        rq.font.name = 'Cambria'
        rq.font.size = Pt(9.0)
        rq.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

generate_expanded_ebook.add_multimedia_box = patched_add_multimedia_box

def main():
    print("=== Bumubuo ng MATATAG Interactive Vision Showcase (.docx) ===")
    
    doc = generate_expanded_ebook.create_base_document()
    
    # Section margins and running headers
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.different_first_page_header_footer = True
    
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("DepEd MATATAG Curriculum • Filipino 7 | Yunit III")
    hr.font.name = 'Cambria'
    hr.font.size = Pt(8.5)
    hr.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fr = fp.add_run("DepEd MATATAG Curriculum • Public Edition / Edisyong Pampubliko")
    fr.font.name = 'Cambria'
    fr.font.size = Pt(8.5)
    fr.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    
    # 1. Front Matter
    print("1. Bumubuo ng Front Matter (Pahina 1-3)...")
    build_yunit3_frontmatter(doc)
    
    # 2. Aralin 13
    print("2. Bumubuo ng Aralin 13...")
    build_yunit3_aralin13(doc)
    
    # 3. Aralin 14
    print("3. Bumubuo ng Aralin 14...")
    build_yunit3_aralin14(doc)
    
    # 4. Aralin 15
    print("4. Bumubuo ng Aralin 15...")
    build_yunit3_aralin15(doc)
    
    # 5. Assessment and Keys
    print("5. Bumubuo ng Assessment, Buod, at Susi sa Pagwawasto...")
    build_yunit3_assessment_and_keys(doc)
    
    output_filename = "MATATAG_Interactive_Vision_Final.docx"
    doc.save(output_filename)
    print(f"Tagumpay na na-save ang: {output_filename}")
    
    # Copy to Knowledge Files benchmark
    benchmark_path = os.path.join(r"ALL_IN_ONE_GPT_SETUP", "KNOWLEDGE_FILES", "00_REFERENCE_TEXTBOOK_BENCHMARK.docx")
    try:
        shutil.copyfile(output_filename, benchmark_path)
        print(f"Kinopya sa Benchmark Knowledge: {benchmark_path}")
    except Exception as e:
        print(f"Hindi nakopya sa benchmark: {e}")

    # Copy to Downloads (handle file locks dynamically)
    downloads_dir = r"C:\Users\Rhenel Jhon Sajol\Downloads"
    saved_copy = None
    for attempt in ["MATATAG_Interactive_Vision_Final.docx", "MATATAG_Interactive_Vision_v2.docx", "MATATAG_Interactive_Vision_v3.docx", "MATATAG_Interactive_Vision_v4.docx"]:
        target_path = os.path.join(downloads_dir, attempt)
        try:
            shutil.copyfile(output_filename, target_path)
            saved_copy = target_path
            print(f"Kinopya sa Downloads: {saved_copy}")
            break
        except PermissionError:
            continue
    if not saved_copy:
        print("Paunawa: Naka-lock ang mga dokumento sa Word. Isara ang Word upang ma-update.")
    
    # Verification & Statistics
    saved_doc = Document(output_filename)
    total_paras = len(saved_doc.paragraphs)
    total_tables = len(saved_doc.tables)
    words_in_p = sum(len(p.text.split()) for p in saved_doc.paragraphs)
    words_in_t = sum(len(c.text.split()) for t in saved_doc.tables for row in t.rows for c in row.cells)
    total_words = words_in_p + words_in_t
    
    print("\n=== ISTATISTIKA NG INTERACTIVE VISION SHOWCASE ===")
    print(f"Kabuuang Salita (Total Words): {total_words}")
    print(f"Kabuuang Talata: {total_paras}")
    print(f"Kabuuang Talahanayan at Interactive Boxes: {total_tables}")
    print(f"Matagumpay na nailapat ang: Zero-Hallucination YouTube Links, Socratic Seminars, Concept Matrices, at 4x4 Rubrics!")

if __name__ == "__main__":
    main()
