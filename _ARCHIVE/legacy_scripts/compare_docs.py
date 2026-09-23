import docx

d0 = docx.Document(r'C:\Users\Rhenel Jhon Sajol\Downloads\MATATAG_Filipino_7_Yunit_III_Aralin_13-15_Textbook.docx')
d1 = docx.Document(r'C:\Users\Rhenel Jhon Sajol\Downloads\MATATAG_Filipino7_Yunit_III_Textbook.docx')
d2 = docx.Document(r'c:\Users\Rhenel Jhon Sajol\Documents\CERE_BOOK\g7-filipino-matatag-curriculum-v10 (2).docx')
d3 = docx.Document(r'c:\Users\Rhenel Jhon Sajol\Documents\CERE_BOOK\g7-filipino-matatag-curriculum-v11-expanded-interactive.docx')

def analyze_doc(doc, name):
    words = sum(len(p.text.split()) for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                words += sum(len(p.text.split()) for p in cell.paragraphs)
    breaks = sum(1 for p in doc.paragraphs for r in p.runs if 'type="page"' in r._r.xml)
    print(f"=== {name} ===")
    print(f"Total Words: {words}")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")
    print(f"Explicit Page Breaks: {breaks}")

analyze_doc(d0, "New Aralin 13-15 (41 pages)")
analyze_doc(d1, "Previous Downloaded Yunit III")
analyze_doc(d2, "Benchmark V10")
analyze_doc(d3, "Compiled V11")
