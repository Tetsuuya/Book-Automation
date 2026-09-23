# -*- coding: utf-8 -*-
"""Master Compilation Script for the 44-Page MATATAG Filipino 7 Yunit III E-Book"""

import sys
import os
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from generate_expanded_ebook import create_base_document
from sections_yunit3_frontmatter import build_yunit3_frontmatter
from sections_yunit3_aralin13 import build_yunit3_aralin13
from sections_yunit3_aralin14 import build_yunit3_aralin14
from sections_yunit3_aralin15 import build_yunit3_aralin15
from sections_yunit3_assessment_keys import build_yunit3_assessment_and_keys

def main():
    print("=== Sinisimulan ang Pagsasama-sama ng 44-Pahinang MATATAG Yunit III E-Book ===")
    
    doc = create_base_document()
    
    # Configure Section Margins and Headers
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.different_first_page_header_footer = True
    
    # Running Header on Pages 2+
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("DepEd MATATAG Curriculum • Filipino 7 | Yunit III")
    hr.font.name = 'Cambria'
    hr.font.size = Pt(8.5)
    hr.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    
    # Running Footer on Pages 2+
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fr = fp.add_run("DepEd MATATAG Curriculum • Edisyong Pampubliko (Baitang 7)")
    fr.font.name = 'Cambria'
    fr.font.size = Pt(8.5)
    fr.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    
    # 1. Front Matter (Pages 1-3)
    print("Bumubuo ng Front Matter (Pahina 1-3)...")
    build_yunit3_frontmatter(doc)
    
    # 2. Aralin 13 (Pages 4-15)
    print("Bumubuo ng Aralin 13 (Pahina 4-15)...")
    build_yunit3_aralin13(doc)
    
    # 3. Aralin 14 (Pages 16-27)
    print("Bumubuo ng Aralin 14 (Pahina 16-27)...")
    build_yunit3_aralin14(doc)
    
    # 4. Aralin 15 (Pages 28-39)
    print("Bumubuo ng Aralin 15 (Pahina 28-39)...")
    build_yunit3_aralin15(doc)
    
    # 5. Assessment and Keys (Pages 40-44)
    print("Bumubuo ng Assessment, Buod, Talatinigan, at Susi sa Pagwawasto (Pahina 40-44)...")
    build_yunit3_assessment_and_keys(doc)
    
    output_filename = "MATATAG_Filipino7_Yunit_III_NaturalFlow.docx"
    doc.save(output_filename)
    print(f"Tagumpay na na-save ang: {output_filename}")
    
    # Also copy to Downloads directory for the user's immediate access
    downloads_path = os.path.join(r"C:\Users\Rhenel Jhon Sajol\Downloads", output_filename)
    shutil.copyfile(output_filename, downloads_path)
    print(f"Kinopya rin sa Downloads: {downloads_path}")
    
    # Verification & Statistics
    saved_doc = Document(output_filename)
    total_paras = len(saved_doc.paragraphs)
    total_tables = len(saved_doc.tables)
    
    words_in_p = sum(len(p.text.split()) for p in saved_doc.paragraphs)
    words_in_t = sum(len(c.text.split()) for t in saved_doc.tables for row in t.rows for c in row.cells)
    total_words = words_in_p + words_in_t
    
    # Count page breaks
    page_breaks = 0
    for p in saved_doc.paragraphs:
        for r in p.runs:
            if '<w:br w:type="page"/>' in r._r.xml or 'w:type="page"' in r._r.xml:
                page_breaks += 1
                
    print("\n=== ISTATISTIKA NG DOKUMENTO ===")
    print(f"Kabuuang Bilang ng Salita (Total Word Count): {total_words} salita")
    print(f"Kabuuang Talata: {total_paras}")
    print(f"Kabuuang Talahanayan at Callout Boxes: {total_tables}")
    print(f"Kabuuang Nilatag na Pahina (Explicit Page Count): {page_breaks + 1} buong pahina")
    print(f"Average Words per Page: {total_words // (page_breaks + 1)} salita bawat pahina")

if __name__ == "__main__":
    main()
