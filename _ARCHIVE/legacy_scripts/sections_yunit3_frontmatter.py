# -*- coding: utf-8 -*-
"""Yunit III Front Matter Builder: Pages 1, 2, and 3"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from generate_expanded_ebook import (
    format_paragraph, add_heading_1, add_heading_2, add_heading_3,
    add_body_p, add_prompt_box, add_callout_box,
    set_cell_shading, set_callout_border, set_cell_margins
)

def build_yunit3_frontmatter(doc):
    # =========================================================================
    # PAHINA 1: PABALAT NG E-BOOK (FRONT COVER PAGE) - ZERO BRANDING
    # =========================================================================
    p_deped = doc.add_paragraph()
    format_paragraph(p_deped, space_before=10, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_deped = p_deped.add_run("REPUBLIKA NG PILIPINAS • KAGAWARAN NG EDUKASYON\nDEPED MATATAG CURRICULUM • FILIPINO 7")
    r_deped.font.name = 'Arial'
    r_deped.font.size = Pt(11)
    r_deped.font.bold = True
    r_deped.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_title = doc.add_paragraph()
    format_paragraph(p_title, space_before=24, space_after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_title = p_title.add_run("GABAY PAMPAGKATUTO SA FILIPINO 7")
    r_title.font.name = 'Cambria'
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_sub = doc.add_paragraph()
    format_paragraph(p_sub, space_before=4, space_after=24, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_sub = p_sub.add_run("YUNIT III: AKO AT ANG AKING PAGKATAO, TANGLAW NG KATATAGAN\nPanitikan sa Panahon ng Pananakop ng Espanya • Tekstong Pampahayagan • Wika at Multimodal")
    r_sub.font.name = 'Cambria'
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    cover_prompt = (
        "PAMAGAT NG AKLAT: Gabay Pampagkatuto sa Filipino 7\n"
        "BAITANG AT YUNIT: Filipino Baitang 7 • Yunit III: 'Ako at ang Aking Pagkatao, Tanglaw ng Katatagan'\n"
        "PAKSA: Panitikan sa Panahon ng Espanya • Tekstong Pampahayagan • Wika at Multimodal\n\n"
        "AI IMAGE PROMPT: A prestigious, high-end educational textbook cover layout. At the top of the cover, a sleek deep navy blue graphic banner featuring clean, crisp, perfectly rendered typography. "
        "The primary headline in bold white capital letters reads: 'GABAY PAMPAGKATUTO SA FILIPINO 7'. "
        "Below it, a secondary title in elegant warm gold letters reads: 'YUNIT III: AKO AT ANG AKING PAGKATAO, TANGLAW NG KATATAGAN'. "
        "Beneath in subtle light silver text: 'Panitikan sa Panahon ng Espanya • Tekstong Pampahayagan • Wika at Multimodal'. "
        "In the lower two-thirds illustration, two Filipino Grade 7 students (a 13-year-old boy and girl in neat modern school uniforms) "
        "studying together at a wooden table under warm golden morning light with an open book and study materials. "
        "In the background, a historic Spanish-colonial stone church belfry and majestic Philippine mountains under a luminous sunrise. "
        "Clean textbook graphic design, perfectly centered professional typography, ultra-high resolution, no gibberish text, no watermarks, no logos --ar 8.5:11 --v 6.0"
    )
    add_prompt_box(doc, "PROMPT SA PAGBUO NG LARAWAN SA PABALAT (FRONT COVER AI IMAGE PROMPT)", cover_prompt)

    p_meta = doc.add_paragraph()
    format_paragraph(p_meta, space_before=36, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_meta = p_meta.add_run("Edisyong Pampubliko para sa mga Paaralang MATATAG\nNakaayon sa MATATAG 21st Century Literacy Framework")
    r_meta.font.name = 'Arial'
    r_meta.font.size = Pt(10)
    r_meta.font.bold = True
    r_meta.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    doc.add_page_break()

    # =========================================================================
    # PAHINA 2: PAUNANG SALITA AT BALANGKAS PEDAGOHIKAL
    # =========================================================================
    add_heading_1(doc, "PAUNANG SALITA AT BALANGKAS PEDAGOHIKAL")
    
    add_body_p(doc, 
        "Malugod na inihahandog ang Yunit III ng Gabay Pampagkatuto sa Filipino 7 na may paksang \"Ako at ang Aking Pagkatao, "
        "Tanglaw ng Katatagan.\" Ang kagamitang pampagtuturong ito ay masusing idinisenyo upang magsilbing komprehensibo, siyentipiko, "
        "at makabuluhang patnubay para sa mga mag-aaral ng Ikapitong Baitang sa ilalim ng bagong Kurikulum na MATATAG ng Kagawaran "
        "ng Edukasyon. Sa pamamagitan ng masinsing pag-aaral sa mga tekstong umusbong at lumaganap sa Panahon ng Pananakop ng Espanya, "
        "layunin ng aklat na ito na linangin ang kritikal na pag-unawa, mapanuring pag-iisip, at pambansang kamalayan ng mga kabataang Pilipino."
    )

    add_body_p(doc, 
        "Ang Yunit III ay buong-buong nakaangkla sa mga Pamantayang Pangnilalaman at Pamantayan sa Pagganap para sa Ikatlong Markahan. "
        "Binibigyang-diin sa yunit na ito ang masusing pagsusuri sa panitikan hindi lamang bilang hiwalay na mga likhang-sining, kundi bilang "
        "mga buháy na bakas ng ating kasaysayan, kultura, wika, at pagkakakilanlan. Dito ay tutuklasin ng mga mag-aaral kung paano nakipagtagpo "
        "ang katutubong kamalayan sa mga banyagang impluwensiya, kung paano ginamit ang panitikan sa pananakop at pananampalataya, at kung "
        "paano patuloy na nagpamalas ng katatagan at sariling tinig ang mga Pilipino sa kabila ng kolonyalismong Espanyol."
    )

    add_body_p(doc, 
        "Mahigpit na ipinapatupad sa aklat na ito ang Prinsipyo ng 70/30 Balangkas Pedagohikal: 70% ng bawat aralin ay nakatuon sa malalim na "
        "masterclass lecture, kontekstuwalisadong pagsusuri ng lunsarang teksto, teoretikal na pagpapaliwanag, at Socratic na talakayan; "
        "samantalang ang 30% ay inilaan sa masinsing aplikasyon, pagsasanay, autentikong gawaing pagganap, at komprehensibong pagtataya. "
        "Sa bawat hakbang, nauuna ang matibay na ebidensiya bago ang interpretasyon, upang sanayin ang mga mag-aaral na maging responsable "
        "at mapanuring mambabasa sa digital at multimodal na panahon."
    )

    add_body_p(doc, 
        "Sa Aralin 13, magsisilbing saligan ang kaligirang pangkasaysayan ng panitikan, nailathalang balita, panghihiram at pagtutumbas ng salita, "
        "at representasyon ng etnisidad. Sa Aralin 14, hihimayin ang Pasyon bilang akdang panrelihiyon, mga pahayagan noong panahon ng Espanyol, "
        "wikang nasusulat, at antas ng pamumuhay. Sa Aralin 15, susuriin ang Urbana at Feliza tungkol sa Kalinisan, pagkuha ng impormasyon para "
        "sa balita, wikang gamit sa panayam, representasyon ng kababaihan, at pagwawasto ng gramatika sa comic book brochure."
    )

    add_callout_box(doc, 
        "PAALALA SA KARAPATANG-SIPI AT PAGGAMIT NG MATERYAL:\n"
        "Ang aklat na ito ay nilikha para sa layuning pang-edukasyon at pampubliko alinsunod sa mga pamantayan ng DepEd MATATAG Curriculum. "
        "Ang mga tekstong pampanitikan, sipi mula sa makasaysayang mga dokumento, at multimodal na mga sanggunian ay kinikilala bilang pangunahing "
        "batis at pinagkunan. Ipinagbabawal ang paggamit ng materyal na ito para sa komersiyal na kapakinabangan nang walang kaukulang pahintulot."
    )

    doc.add_page_break()

    # =========================================================================
    # PAHINA 3: TALAAN NG NILALAMAN AT BALANGKAS NG PAGKATUTO (FULL DETAILED PAGE)
    # =========================================================================
    tbl_banner = doc.add_table(rows=1, cols=1)
    tbl_banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_ban = tbl_banner.rows[0].cells[0]
    set_cell_shading(c_ban, "1B365D")
    set_cell_margins(c_ban, top=140, bottom=140, left=180, right=180)
    p_ban = c_ban.paragraphs[0]
    format_paragraph(p_ban, space_before=2, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_ban = p_ban.add_run("TALAAN NG NILALAMAN AT BALANGKAS NG PAGKATUTO")
    r_ban.font.name = 'Arial'
    r_ban.font.size = Pt(12)
    r_ban.font.bold = True
    r_ban.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    tbl_toc = doc.add_table(rows=1, cols=1)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_toc = tbl_toc.rows[0].cells[0]
    set_cell_shading(c_toc, "FAFAFA")
    set_callout_border(c_toc, color_hex="1B365D", sz="36")
    set_cell_margins(c_toc, top=100, bottom=100, left=160, right=160)

    p0 = c_toc.paragraphs[0]
    format_paragraph(p0, space_before=1, space_after=2)
    r0 = p0.add_run("YUNIT III: AKO AT ANG AKING PAGKATAO, TANGLAW NG KATATAGAN")
    r0.font.name = 'Cambria'
    r0.font.size = Pt(10.5)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    toc_entries = [
        ("ARALIN 13: Ang Panitikan sa Panahong Sakop ng Espanya ang Pilipinas", "4", True),
        ("  • Lunsaran: Kaligirang Pangkasaysayan at Tekstong Paglalahad", "", False),
        ("  • Paksa 13.1: Panitikan at Kaligirang Pangkasaysayan sa Panahon ng Espanyol", "5", False),
        ("  • Paksa 13.2: Nailathalang Balita at Panghihiram o Pagtutumbas ng Salita", "6", False),
        ("  • Paksa 13.3: Etnisidad, Tekstong Biswal, at Gawaing Multimodal", "9", False),
        ("  • Mga Gawaing Pampagkatuto 13.1–13.3 at Mabilisang Pagtataya (15 Puntos)", "10", False),
        ("  • Alokasyon ng Oras: 5 Sesyon (Sesyon 1 hanggang 5)", "15", False),
        ("", "", False),
        ("ARALIN 14: Pasyon at mga Tekstong Panrelihiyon sa Panahong Kolonyal", "16", True),
        ("  • Lunsaran: Ang Pasyon bilang Panitikan at Kulturang Pasalaysay", "", False),
        ("  • Paksa 14.1: Akdang Panrelihiyon—Anyo, Mensahe, Konteksto, at Tradisyon", "17", False),
        ("  • Paksa 14.2: Mga Pahayagan at Wikang Nasusulat noong Panahon ng Espanyol", "18", False),
        ("  • Paksa 14.3: Antas ng Pamumuhay at Elementong Biswal ng Comic Book Brochure", "21", False),
        ("  • Mga Gawaing Pampagkatuto 14.1–14.3 at Mabilisang Pagtataya (15 Puntos)", "22", False),
        ("  • Alokasyon ng Oras: 5 Sesyon (Sesyon 6 hanggang 10)", "27", False),
        ("", "", False),
        ("ARALIN 15: Ang Liham ni Urbana kay Feliza tungkol sa Kalinisan", "28", True),
        ("  • Lunsaran: Akdang Pangkagandahang-asal at Wikang Didaktiko", "", False),
        ("  • Paksa 15.1: Mensahe, Pagpapahalaga, at Konteksto ng Akdang Pangkagandahang-asal", "29", False),
        ("  • Paksa 15.2: Pagkuha ng Impormasyon para sa Balita at Wikang Gamit sa Panayam", "30", False),
        ("  • Paksa 15.3: Kababaihan sa Panahon ng Espanyol at Pagwawasto ng Diyalogo", "33", False),
        ("  • Mga Gawaing Pampagkatuto 15.1–15.3 at Mabilisang Pagtataya (15 Puntos)", "34", False),
        ("  • Alokasyon ng Oras: 5 Sesyon (Sesyon 11 hanggang 15)", "39", False),
        ("", "", False),
        ("PANGWAKAS NA PAGTATAYA SA YUNIT III (KABUUANG 40 PUNTOS)", "40", True),
        ("  • Bahagi I–IV: Pagsusuri ng Batis, Konsepto, Sanaysay, at GRASPS Comic Brochure", "40", False),
        ("BUOD NG YUNIT III AT KOMPREHENSIBONG TALATINIGAN (25 KONSEPTO)", "43", True),
        ("KOMPLETONG SUSI SA PAGWAWASTO, RASYONAL, AT GABAY SA GURO", "44", True)
    ]

    for title, page, is_bold in toc_entries:
        p = c_toc.add_paragraph()
        format_paragraph(p, space_before=0, space_after=1, line_spacing=1.05)
        r_t = p.add_run(title)
        r_t.font.name = 'Cambria'
        r_t.font.size = Pt(8.5 if not is_bold else 9)
        if is_bold:
            r_t.font.bold = True
            r_t.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        else:
            r_t.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            
        if page:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r_dots = p.add_run(f" {'.' * max(2, 60 - len(title))} ")
            r_dots.font.name = 'Cambria'
            r_dots.font.size = Pt(8)
            r_dots.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            r_p = p.add_run(page)
            r_p.font.name = 'Cambria'
            r_p.font.size = Pt(8.5 if not is_bold else 9)
            r_p.font.bold = True
            r_p.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    doc.add_page_break()
