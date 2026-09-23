# -*- coding: utf-8 -*-
"""Front matter builder: Pages 1, 2, and 3"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from generate_expanded_ebook import (
    format_paragraph, add_heading_1, add_heading_2, add_heading_3,
    add_body_p, add_prompt_box, add_callout_box, add_custom_table,
    set_cell_shading, set_callout_border, set_cell_margins
)

def build_frontmatter(doc):
    # =========================================================================
    # PAHINA 1: PABALAT NG E-BOOK (FRONT COVER PAGE) - WALANG LEKTURA
    # =========================================================================
    p_deped = doc.add_paragraph()
    format_paragraph(p_deped, space_before=10, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_deped = p_deped.add_run("REPUBLIKA NG PILIPINAS • KAGAWARAN NG EDUKASYON\nDEPED MATATAG CURRICULUM • FILIPINO 7")
    r_deped.font.name = 'Arial'
    r_deped.font.size = Pt(11)
    r_deped.font.bold = True
    r_deped.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_title = doc.add_paragraph()
    format_paragraph(p_title, space_before=24, space_after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_title = p_title.add_run("GABAY PAMPAGKATUTO SA FILIPINO 7")
    r_title.font.name = 'Cambria'
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_sub = doc.add_paragraph()
    format_paragraph(p_sub, space_before=4, space_after=24, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_sub = p_sub.add_run("YUNIT II: MGA TULUYANG PANITIKAN AT TEKSTONG EKSPOSITORI SA PAGHUBOG NG MAKABAYANG KAMALAYAN AT MAPANURING PANANAW")
    r_sub.font.name = 'Cambria'
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    cover_prompt = (
        "PROMPT: A breathtaking, high-detail editorial book cover illustration capturing Philippine cultural heritage "
        "and 21st-century modern learning. In the foreground, an ancient giant Balete tree with majestic sprawling roots "
        "and warm golden sunlight filtering through the canopy. Filipino Grade 7 students (one boy and one girl in neat "
        "modern school uniforms) sitting together under the shade, engaged in thoughtful academic discussion while holding "
        "an illuminated digital tablet displaying indigenous Baybayin script and cultural folklore illustrations. "
        "In the background, a scenic Philippine rural landscape blending harmoniously with a modern, eco-friendly community "
        "library and distant misty mountains. Vibrant tropical color palette of emerald greens, warm ochre, and deep twilight "
        "navy blues. Style: Premium modern digital oil painting, highly detailed textures, warm cinematic lighting, artistic realism, "
        "8k resolution, suitable for a national educational textbook cover --ar 8.5:11 --v 6.0"
    )
    add_prompt_box(doc, "PROMPT SA PAGBUO NG LARAWAN SA PABALAT (FRONT COVER AI IMAGE PROMPT)", cover_prompt)

    p_meta = doc.add_paragraph()
    format_paragraph(p_meta, space_before=36, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_meta = p_meta.add_run("CERE Book Series • Edisyong Pampagkatuto at Pangguro\nNakaayon sa MATATAG 21st Century Literacy Framework")
    r_meta.font.name = 'Arial'
    r_meta.font.size = Pt(10)
    r_meta.font.bold = True
    r_meta.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    doc.add_page_break()

    # =========================================================================
    # PAHINA 2: PAUNANG SALITA, KARAPATANG-SIPI, AT PEDAGOHIKAL NA BALANGKAS
    # =========================================================================
    add_heading_1(doc, "PAUNANG SALITA AT PEDAGOHIKAL NA BALANGKAS")
    add_body_p(
        doc,
        "Malugod na pagtanggap sa mga mag-aaral at kapuwa mga gurong gagamit ng Gabay Pampagkatuto na ito para sa Ikapitong Baitang ng Bagong MATATAG Kurikulum! Ang aklat na ito ay maingat na binuo upang magsilbing tulay sa pagitan ng ating mayamang pamanang kultural at ng mga kasanayang pang-ika-21 siglo na kailangang-kailangan ng bawat kabataang Pilipino.",
        bold_prefix="Isang Pagpupugay sa Karunungan ng Bayan: "
    )
    add_body_p(
        doc,
        "Sa ilalim ng pinalakas na MATATAG Kurikulum, ang pagkatuto sa Filipino ay lumalagpas sa tradisyonal na pagsasaulo ng mga kahulugan at elemento ng panitikan. Sa halip, binibigyang-diin ang integratibong pagkatuto kung saan ang panitikan ay tinitingnan bilang buhay na salamin ng sosyo-kultural at pangkasaysayang konteksto ng ating pamayanan. Hindi lamang tayo nagbabasa ng kuwento; sinusuri natin ang mga kapangyarihang humubog dito, ang mga pananaw ng ating mga ninuno, at ang aral na maiaambag nito sa ating kasalukuyang lipunan."
    )
    add_body_p(
        doc,
        "Ang Yunit II ay nakatuon sa 'Mga Tuluyang Panitikan at Tekstong Ekspositori.' Sa pamamagitan ng masusing pagbasa sa maikling kuwento, alamat, at kuwentong-bayan, matutuklasan ng mag-aaral kung paano ginagamit ng mga manunulat ang wika upang maglahad ng katotohanan, magpatibay ng pananaw, at maglinaw ng mga komplikadong kaisipan. Kaalinsabay nito, sasanayin ang mga mag-aaral sa paggamit ng mga transisyong gramatikal, retorikal na pang-ugnay, at kohesiyong gramatikal (anapora at katapora) upang maging matatas, organisado, at lohikal sa pagsulat ng sariling mga teksto."
    )
    add_body_p(
        doc,
        "Espesyal na tampok ng e-book na ito ang pagiging interaktibo at multimedia-ready. Bawat aralin ay nagtataglay ng mga piling video documentary links mula sa mga mapagkakatiwalaang pampublikong tsanel tulad ng DepEd TV, Knowledge Channel, at PTV, na may kalakip na QR Code representations, timestamps, at mga gabay sa mapanuring panonood. Bukod dito, ang bawat visual na elemento ay may katapat na AI Image Generation Prompt upang mabigyang-daan ang visual literacy at malikhaing paglalapat ng teknolohiya sa silid-aralan."
    )
    add_body_p(
        doc,
        "Ang aklat na ito ay sumusunod sa gintong pamantayan ng '70% Malalim na Aralin at 30% Mapanuring Gawain.' Tinitiyak nito na bago sumabak ang mag-aaral sa mga pagsasanay at pagsusulit, taglay na nila ang malawak na teoretikal at konseptwal na pag-unawa. Mayroon ding kompletong Susi sa Pagwawasto sa dulo ng aklat na nagbibigay ng masusing paliwanag sa bawat sagot, upang maging katuwang ng guro sa epektibong pagtataya at magsilbing kagamitan sa sariling pagkatuto ng mag-aaral."
    )

    p_copy = doc.add_paragraph()
    format_paragraph(p_copy, space_before=12, space_after=4)
    r_c = p_copy.add_run("Pahina ng Karapatang-Sipi: © 2026 CERE Books Publishing. Reserbado ang lahat ng karapatan. Alinsunod sa Batas Republika Blg. 8293, walang bahagi ng aklat na ito ang maaaring kopyahin nang walang nakasulat na pahintulot mula sa may-akda at tagapaglathala, maliban sa mga sipi para sa layuning pampagkatuto sa silid-aralan.")
    r_c.font.name = 'Cambria'
    r_c.font.size = Pt(8.5)
    r_c.font.italic = True
    r_c.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_page_break()

    # =========================================================================
    # PAHINA 3: DETALYADONG TALAAN NG NILALAMAN AT MATATAG CURRICULUM MATRIX
    # =========================================================================
    # 1. Solid Navy Blue Banner
    p_banner = doc.add_paragraph()
    p_banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_banner, space_before=14, space_after=8)
    shd_xml = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1B365D"/>')
    p_banner._p.get_or_add_pPr().append(shd_xml)
    r_b = p_banner.add_run("TALAAN NG NILALAMAN AT BALANGKAS NG PAGKATUTO")
    r_b.font.name = 'Cambria'
    r_b.font.size = Pt(16)
    r_b.font.bold = True
    r_b.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # 2. Bordered Callout Box with Thick Navy Left Stripe (#1B365D, 4.5pt) and #FAFAFA Shading
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    cell = tbl.rows[0].cells[0]
    set_cell_shading(cell, "FAFAFA")
    set_callout_border(cell, color_hex="1B365D", sz="36")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)

    toc_lines = [
        ("YUNIT II: MGA TULUYANG PANITIKAN AT TEKSTONG EKSPOSITORI", True, "1B365D", 11.5, 0),
        ("ARALIN 7: Pagbasa sa Panitikan at Pananaw ng Pamayanan .......................................... 4", True, "1B365D", 10.5, 0),
        ("   • Panitikan: Sanaysay/Tuluyan (Lunsaran: \"Ang Puno sa Liwasan\")", False, "333333", 9.5, 0),
        ("   • Paksa 7.1: Ang Panitikang Tuluyan at Kahalagahan ng Konteksto ............................... 5", False, "333333", 9.5, 0),
        ("   • Paksa 7.2: Paglalahad, Opinyon, at Pananaw ................................................................ 7", False, "333333", 9.5, 0),
        ("   • Paksa 7.3: Transisyong Gramatikal at Pang-ugnay na Retorikal ............................... 9", False, "333333", 9.5, 0),
        ("   • Gawaing Pampagkatuto (Mga Gawain 7.1, 7.2, at 7.3) ............................................... 11", False, "333333", 9.5, 0),
        ("   • Mabilisang Pagtataya sa Aralin 7 (Maikling Pagsusulit - 15 Puntos) ........................ 13", False, "333333", 9.5, 0),
        ("   • Alokasyon ng Oras: 5 Sesyon (Sesyon 1 hanggang 5)", False, "555555", 9.5, 6),
        ("ARALIN 8: Kuwento ng Pinagmulan, Linaw ng Pagpapaliwanag .................................... 14", True, "1B365D", 10.5, 0),
        ("   • Panitikan: Alamat (Lunsaran: \"Bakit May Liwanag sa Batis?\")", False, "333333", 9.5, 0),
        ("   • Paksa 8.1: Ang Alamat Bilang Panitikang Tuluyan at Malikhaing Paliwanag .................. 15", False, "333333", 9.5, 0),
        ("   • Paksa 8.2: Faktuwal na Paliwanag at Pagsusuri ng Sanggunian ................................ 17", False, "333333", 9.5, 0),
        ("   • Paksa 8.3: Kohesiyong Gramatikal: Anapora at Katapora ............................................. 19", False, "333333", 9.5, 0),
        ("   • Gawaing Pampagkatuto (Mga Gawain 8.1, 8.2, at 8.3) ............................................... 21", False, "333333", 9.5, 0),
        ("   • Mabilisang Pagtataya sa Aralin 8 (Maikling Pagsusulit - 15 Puntos) ........................ 23", False, "333333", 9.5, 0),
        ("   • Alokasyon ng Oras: 5 Sesyon (Sesyon 6 hanggang 10)", False, "555555", 9.5, 6),
        ("ARALIN 9: Kuwento ng Pamayanan, Mensahe para sa Bayan ....................................... 24", True, "1B365D", 10.5, 0),
        ("   • Panitikan: Kuwentong-Bayan (Lunsaran: \"Ang Kampana sa Gitna ng Nayon\") ......... 25", False, "333333", 9.5, 0),
        ("   • Paksa 9.1: Ang Kuwentong-Bayan Bilang Salamin ng Pamayanan ............................... 26", False, "333333", 9.5, 0),
        ("   • Paksa 9.2: Pagbuo ng Organisadong Talata at Transisyon ........................................... 28", False, "333333", 9.5, 0),
        ("   • Paksa 9.3: Multimodal na Komunikasyon at Makatarungang Representasyon ............ 30", False, "333333", 9.5, 0),
        ("   • Gawaing Pampagkatuto (Mga Gawain 9.1, 9.2, at 9.3) ............................................... 32", False, "333333", 9.5, 0),
        ("   • Mabilisang Pagtataya sa Aralin 9 (Maikling Pagsusulit - 15 Puntos) ........................ 34", False, "333333", 9.5, 0),
        ("   • Alokasyon ng Oras: 5 Sesyon (Sesyon 11 hanggang 15)", False, "555555", 9.5, 6),
        ("PANGWAKAS NA PAGTATAYA SA YUNIT II (COMPREHENSIVE EXAM - 40 Puntos) ...... 35", True, "1B365D", 10.5, 0),
        ("   • Bahagi I: Pagsusuring Pampanitikan at Konteksto (15 Puntos)", False, "333333", 9.5, 0),
        ("   • Bahagi II: Balarila, Transisyon, at Kohesiyon (15 Puntos)", False, "333333", 9.5, 0),
        ("   • Bahagi III: Biswal at Multimodal na Komunikasyon (5 Puntos)", False, "333333", 9.5, 0),
        ("   • Bahagi IV: Malikhaing Pagsulat na may Grapikong Rubrik (5 Puntos)", False, "333333", 9.5, 6),
        ("KOMPLETONG SUSI SA PAGWAWASTO AT GABAY SA PAGMAMARKA (PARA SA GURO) . 38", True, "1B365D", 10.5, 0),
        ("   • Susi sa Pagwawasto para sa Aralin 7 (Mga Gawain at Pagsusulit) ............................ 38", False, "333333", 9.5, 0),
        ("   • Susi sa Pagwawasto para sa Aralin 8 (Mga Gawain at Pagsusulit) ............................ 40", False, "333333", 9.5, 0),
        ("   • Susi sa Pagwawasto para sa Aralin 9 (Mga Gawain at Pagsusulit) ............................ 42", False, "333333", 9.5, 0),
        ("   • Susi sa Pagwawasto para sa Pangwakas na Pagtataya (Yunit II) .............................. 44", False, "333333", 9.5, 0)
    ]

    p_first = cell.paragraphs[0]
    first = True
    for text, bold, color, size, after_sp in toc_lines:
        p = p_first if first else cell.add_paragraph()
        first = False
        format_paragraph(p, space_before=1, space_after=after_sp, line_spacing=1.1)
        r = p.add_run(text)
        r.font.name = 'Cambria'
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = RGBColor(int(color[:2], 16), int(color[2:4], 16), int(color[4:], 16))

    doc.add_page_break()
