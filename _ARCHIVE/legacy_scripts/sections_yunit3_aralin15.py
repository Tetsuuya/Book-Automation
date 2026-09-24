# -*- coding: utf-8 -*-
"""Aralin 15 Builder: Pages 28 to 39 (12 Full Pages)"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from generate_expanded_ebook import (
    format_paragraph, add_heading_1, add_heading_2, add_heading_3,
    add_body_p, add_prompt_box, add_callout_box, add_custom_table,
    add_multimedia_box, set_cell_shading
)

def build_yunit3_aralin15(doc):
    # =========================================================================
    # PAHINA 28: PANIMULA, LAYUNIN, AT MULTIMEDIA CORNER #3
    # =========================================================================
    add_heading_1(doc, "ARALIN 15: Ang Liham ni Urbana kay Feliza tungkol sa Kalinisan")
    add_heading_2(doc, "MGA LAYUNIN AT BALANGKAS NG PAGKATUTO")
    add_body_p(
        doc,
        "Sa Aralin 15, hihimayin natin ang isa sa pinakamahahalagang akdang pangkagandahang-asal sa kasaysayan ng panitikang Pilipino: ang "
        "Pagsusulatan nina Urbana at Feliza (1864) na isinulat ni Padre Modesto de Castro. Tutuklasin natin kung paano ginamit ang anyong liham upang "
        "magtatag ng mga pamantayan ng wastong kilos, kalinisan ng katawan at tahanan, pakikitungo sa kapuwa, at gampanin ng kababaihan sa lipunang kolonyal. "
        "Kaagapay nito, pag-aaralan ang makabagong kasanayan sa pagkuha ng impormasyon para sa balita sa pamamagitan ng etikal na panayam, ang mapanuring "
        "pagsusuri sa representasyon ng kababaihan sa mga tekstong biswal, at ang masusing pagwawasto ng gramatika sa pagbuo ng mga diyalogo sa komiks."
    )

    add_body_p(
        doc,
        "1. F7PN-IIIc-1: Nasusuri ang mensahe, pagpapahalaga, at kontekstong historikal ng akdang pangkagandahang-asal.\n"
        "2. F7PB-IIIc-2: Nakabubuo ng mga neutral, bukas, at mabisang tanong sa panayam para sa pagsulat ng balita.\n"
        "3. F7PT-IIIc-3: Naipaliliwanag ang wastong gamit ng wika, antas ng pormalidad, at tono sa pagsasagawa ng panayam.\n"
        "4. F7PD-IIIc-4: Kritikal na nasusuri ang representasyon ng kababaihan sa mga historikal na tekstong biswal nang walang stereotype.\n"
        "5. F7PU-IIIc-5: Nagwawasto ng gramatika, bantas, at natural na daloy ng mga diyalogo sa pagbuo ng pinal na comic book brochure.",
        bold_prefix="Mga Kasanayang Pampagkatuto (MATATAG Competencies):\n"
    )

    add_body_p(
        doc,
        "• Paano naiiba ang simpleng pagsunod sa mga tuntunin ng kagandahang-asal sa malalim na pagkaunawa sa pagpapahalagang nasa likod nito?\n"
        "• Bakit itinuturing na pundasyon ng responsableng pamamahayag ang pagbuo ng mga tanong na neutral at hindi nangunguna sa panayam?\n"
        "• Paano nakatutulong ang wastong gramatika at maingat na boses ng tauhan sa pagtatanghal ng dignidad ng kababaihan sa komiks?",
        bold_prefix="Mga Susing Katanungan para sa Aralin:\n"
    )

    add_multimedia_box(
        doc,
        mod_title="Pagsusuri sa Urbana at Feliza at Asal-Kolonyal",
        vid_title="Urbana at Feliza ni Padre Modesto de Castro: Pagsusuri sa Kagandahang-Asal ng ika-19 na Siglo",
        channel="Knowledge Channel / PTV Educational Series",
        link="https://www.youtube.com/results?search_query=Knowledge+Channel+Urbana+at+Feliza+Modesto+de+Castro",
        qr_code_text="",
        timestamps=[
            "01:10 - 04:50: Ang Konteksto ng Pagsusulatan ng Dalawang Magkapatid (Maynila at Paombong).",
            "04:51 - 08:30: Pagsusuri sa Konsepto ng 'Kalinisan'—Mula sa Katawan Patungo sa Moralidad.",
            "08:31 - 13:15: Ang Papel ng Kababaihan sa Lipunang Kolonyal at Pagtaliwas sa Stereotype."
        ],
        questions=[
            "Ayon sa video lesson, paano naging patunay ang nobelang epistolaryo sa mataas na antas ng kasanayan sa wikang Tagalog noong ika-19 na siglo?",
            "Bakit binigyang-diin sa aralin na ang payo ni Urbana tungkol sa kalinisan ay may direktang kaugnayan sa kalusugang pampubliko ng pamayanan?",
            "Paano mo ihahambing ang mga tuntunin sa kagandahang-asal noon sa 'netiquette' o tamang pag-uugali sa social media ngayon?"
        ]
    )

    # doc.add_page_break() [REMOVED TO PREVENT ORPHAN BLANK PAGES]

    # =========================================================================
    # PAHINA 29: PAKSA 15.1 - AKDANG PANGKAGANDAHANG-ASAL
    # =========================================================================
    add_heading_2(doc, "PAKSA 15.1: Akdang Pangkagandahang-asal: Mensahe, Pagpapahalaga, at Konteksto")
    
    prompt_15_sala = (
        "PROMPT: A breathtaking, historically authentic interior view of the grand sala of a 19th-century Philippine Bahay na Bato in Manila, 1864. "
        "High timber ceilings, polished wide narra hardwood floors with warm reflections, sliding translucent capiz shell windows and open ventanillas "
        "revealing tropical garden greenery outside. Antique Viennese bentwood rocking chairs, a round marble-topped table with brass oil lamps, and potted palms. "
        "Pristine historical interior architecture, editorial photography aesthetic, warm natural afternoon light --ar 16:9 --v 6.0"
    )
    add_prompt_box(doc, "ARKITEKTURA AT LIPUNAN: ANG SALA NG BAHAY NA BATO SA MAYNILA (1864)", prompt_15_sala)
    
    add_body_p(
        doc,
        "Ang akdang pangkagandahang-asal ay isang uri ng panitikang didaktiko o pangaral na naglalayong magturo ng mga inaasahang gawi, wastong pagkilos, "
        "moralidad, at pakikipagkapuwa sa loob ng isang lipunan. Noong ika-19 na siglo sa Pilipinas, ang pinakatanyag na halimbawa nito ay ang "
        "Pagsusulatan ng Dalawang Binibini na si Urbana at si Feliza (1864) ni Padre Modesto de Castro. Ang aklat na ito ay naging 'bibliya ng kagandahang-asal' "
        "na binabasa sa mga paaralan at tahanan upang gabayan ang mga kabataan sa kanilang paglaki sa ilalim ng tradisyong Katoliko at kolonyal."
    )

    add_body_p(
        doc,
        "Ang anyong liham (epistolaryo) ng akda ay napakahalagang sangkap sa pagsusuri. Hindi ito isang malamig na listahan ng mga batas; bagkus, ito ay "
        "palitan ng sulat sa pagitan ng magkapatid: si Urbana na nag-aaral sa Maynila sa kolehiyo ng mga madre, at si Feliza na naiwan sa lalawigan sa Paombong, "
        "Bulacan. Ang boses ni Urbana ay boses ng nakatatandang kapatid na may malasakit, pagmamahal, at pagnanais na ihanda ang kaniyang nakababatang kapatid "
        "at bunsong kapatid na si Honesto sa maayos na pakikipamuhay sa kapuwa at sa bayan."
    )

    add_body_p(
        doc,
        "Sa pagsusuri ng liham tungkol sa kalinisan, makikita na ang 'kalinisan' ay hindi lamang tumutukoy sa paghuhugas ng kamay o pagsusuot ng malinis na "
        "damit. Mayroon itong tatlong antas: (1) Pisikal na Kalinisan—ang pangangalaga sa katawan upang maiwasan ang sakit; (2) Panlipunang Kalinisan—ang "
        "paggalang sa espasyo ng iba at kaayusan sa hapag-kainan upang hindi makapagbigay ng kahihiyan o pandidiri sa kapuwa; at (3) Moral na Kalinisan—ang "
        "kalinisan ng budhi, katapatan, at kadalisayan ng kalooban na siyang pinagmumulan ng tunay na dangal ng isang tao."
    )

    headers_15_1 = ["Lente sa Pagsusuri", "Kahulugan sa Akda", "Kontekstong Historikal (1864)", "Mapanuring Pagsusuri sa Kasalukuyan"]
    data_15_1 = [
        ["Tagapagsalita at Tono", "Si Urbana; mapagpaalala, mapagmahal, didaktiko", "Uring principalia; edukasyon sa Maynila sa ilalim ng mga madre", "Ipinakikita ang kapangyarihan ng kapatid bilang tagapagturo sa pamilya"],
        ["Pinatutungkulan", "Si Feliza at kapatid na si Honesto", "Kabataang nasa lalawigan na hinahanda sa pagtuntong sa lungsod", "Pundasyon ng edukasyong pangkapayapaan sa loob ng tahanan"],
        ["Payo sa Kalinisan", "Pag-iingat sa katawan, kasuotan, at pagkain", "Kakulangan sa modernong ospital; pag-iwas sa mga epidemya tulad ng kolera", "Nananatiling mahalagang aral sa kalusugang pampubliko at kalinisan"],
        ["Pagpapahalaga (Values)", "Paggalang sa kapuwa, hiya, disiplina, at dangal", "Pagpapahalagang kolonyal na naglalayong magpanatili ng kaayusan", "Dapat suriin kung aling asal ang makatao at alin ang labis na mapagpasailalim"]
    ]
    add_custom_table(doc, headers_15_1, data_15_1)

    add_body_p(
        doc,
        "Mahalagang paghiwalayin ang pag-unawa sa isang historikal na teksto laban sa awtomatikong pagsang-ayon sa lahat ng nilalaman nito. Nauunawaan natin "
        "kung bakit isinulat ni P. Modesto de Castro ang Urbana at Feliza sa konteksto ng ika-19 na siglo, ngunit may laya rin tayong suriin kung paano ang "
        "ilang payo ay nagtakda ng labis na limitasyon sa kalayaan ng kababaihan. Ang ganitong dalawahang pagbasa ang nagpapatunay ng pagiging mapanuri."
    )

    add_callout_box(
        doc,
        "MUNTING PAGSASANAY 15.1 (Pagsusuri sa Urbana at Feliza):\n"
        "1. Bakit naging napakabisa ng anyong liham sa pagtuturo ng kagandahang-asal kumpara sa isang pormal na aklat ng batas?\n"
        "2. Ipaliwanag ang tatlong antas ng kalinisan na matatagpuan sa akda: pisikal, panlipunan, at moral.\n"
        "3. Paano mo ipaliliwanag ang pagkakaiba ng pag-unawa sa konteksto ng 1864 laban sa pagpapasya kung aling aral ang ilalapat sa iyong buhay ngayon?"
    )

    # doc.add_page_break() [REMOVED TO PREVENT ORPHAN BLANK PAGES]

    # =========================================================================
    # PAHINA 30: PAKSA 15.2 - PANAYAM AT PAGKUHA NG IMPORMASYON
    # =========================================================================
    add_heading_2(doc, "PAKSA 15.2: Pagkuha ng Impormasyon para sa Balita at Wikang Gamit sa Panayam")
    
    add_body_p(
        doc,
        "Ang panayam ay isang sistematiko, layunin, at organisadong pakikipagtalastasan sa pagitan ng tagapanayam (interviewer) at kinakapanayam (interviewee) "
        "upang makakuha ng mapagkakatiwalaan, bago, at malalim na impormasyon tungkol sa isang tiyak na paksa o pangyayari. Sa pagsulat ng balita, "
        "ang panayam ang nagbibigay ng buhay at kredibilidad sa ulat. Kung walang panayam, ang balita ay magiging tuyo at maaaring mauwi sa sariling "
        "opinyon o haka-haka ng sumulat."
    )

    add_body_p(
        doc,
        "Upang maging matagumpay ang panayam, napakahalaga ng pagbuo ng mga tanong. May tatlong pangunahing uri ng tanong: (1) Saradong Tanong—nasasagot "
        "ng oo, hindi, petsa, o tiyak na datos upang kumpirmahin ang katotohanan; (2) Bukas na Tanong—nagbibigay-daan sa malawak na pagpapaliwanag, paglalahad "
        "ng karanasan, at pananaw ng kinakapanayam; at (3) Follow-up Question—tanong na sumusunod sa naging sagot upang humingi ng karagdagang halimbawa, "
        "linawin ang malabong pahayag, o beripikahin ang isang mahalagang detalye."
    )

    add_body_p(
        doc,
        "Ang pinakamalaking pagkakamali sa panayam ay ang paggamit ng nangungunang tanong (leading question). Ang nangungunang tanong ay naglalaman na ng "
        "inaasahang sagot o paghusga ng nagtatanong. Halimbawa: 'Hindi ba't napakahusay ng inyong proyekto sa kalinisan?' Ang tanong na ito ay hindi neutral "
        "dahil idinidiktang sumang-ayon ang kausap. Ang tamang neutral na tanong ay: 'Paano ninyo sinusuri ang naging resulta ng inyong proyektong pangkalinisan?' "
        "Sa neutral na tanong, may kalayaan ang kinakapanayam na maglahad ng tagumpay at ng mga hamon na nananatili."
    )

    headers_15_2 = ["Uri ng Tanong", "Katangian at Gamit", "Maling Halimbawa (Nangunguna)", "Tamang Halimbawa (Neutral at Mapanuri)"]
    data_15_2 = [
        ["Saradong Tanong", "Para sa tiyak na faktuwal na datos at beripikasyon", "\"Nagsimula ba ito noong Lunes?\"", "\"Kailan po opisyal na sinimulan ang kampanya sa paaralan?\""],
        ["Bukas na Tanong", "Para sa masusing pagpapaliwanag at malayang pananaw", "\"Hindi ba't masaya kayo sa programa?\"", "\"Ano po ang inyong naging karanasan sa pagpapatupad ng programa?\""],
        ["Follow-up Question", "Para sa pagpapalalim, paglilinaw, at paghahanap ng datos", "\"Talaga pong mahirap?\"", "\"Maaari po ba kayong magbahagi ng tiyak na halimbawa ng hamong naranasan?\""],
        ["Tanong sa Beripikasyon", "Para sa pagkumpirma ng mga datos at bilang", "\"Kayo ang pinakamagaling, tama?\"", "\"Tama po ba ang tala na 80% ng mga silid ang napanatiling malinis?\""]
    ]
    add_custom_table(doc, headers_15_2, data_15_2)

    add_body_p(
        doc,
        "Sa etika ng pamamahayag, may pananagutan ang tagapanayam na pakinggan ang buong pahayag nang may paggalang at huwag baluktutin ang sinabi ng kausap. "
        "Kung magpa-paraphrase (magpaliwanag sa sariling salita), dapat manatili ang orihinal na diwa. Ang mahusay na mamamahayag ay naghahanap ng "
        "katotohanan at hindi naghahanap lamang ng kumpirmasyon sa kaniyang sariling pananaw."
    )

    add_callout_box(
        doc,
        "MUNTING PAGSASANAY 15.2 (Pagsusuri sa Tanong ng Panayam):\n"
        "1. Bakit itinuturing na hindi etikal ang paggamit ng mga nangungunang tanong sa isang pormal na panayam para sa balita?\n"
        "2. Ayusin ang sumusunod na nangungunang tanong upang maging bukas at neutral: \"Hindi ba't tamad ang mga mag-aaral na hindi naglilinis?\"\n"
        "3. Paano nakatutulong ang follow-up question sa pagtuklas ng mga impormasyong hindi agad lumalabas sa unang tanong?"
    )

    # doc.add_page_break() [REMOVED TO PREVENT ORPHAN BLANK PAGES]

    # =========================================================================
    # PAHINA 31: LUNSARANG TEKSTO - ANG LIHAM NI URBANA KAY FELIZA
    # =========================================================================
    add_heading_2(doc, "LUNSARANG TEKSTO: Sipi mula sa Urbana at Feliza tungkol sa Kalinisan")
    
    prompt_15_dining = (
        "PROMPT: A staged educational historical illustration depicting proper 19th-century Filipino dining etiquette based on Urbana at Feliza. "
        "An ilustrado family seated politely at a long narra dining table set with white linen, fine blue-and-white porcelain plates, and silver flatware. "
        "Upright posture, courteous hand positioning, and respectful social interaction without elbows on the table. Warm candlelight and chandelier illumination, classical instructional art style --ar 16:9 --v 6.0"
    )
    add_prompt_box(doc, "KAGANDAHANG-ASAL SA HAPAG: PORMAL NA HAPAG-KAINAN NG ILUSTRADO", prompt_15_dining)
    add_body_p(doc, "Klasikong Akdang Didaktiko mula sa Ika-19 na Siglo (Pagsusulatan nina Urbana at Feliza, 1864)", italic_prefix="")

    add_body_p(
        doc,
        "\"KAPATID KONG FELIZA:\n\n"
        "Sa liham mong huli ay sinabi mo sa akin na si Honesto, ang bunso nating kapatid, ay nagsisimula nang pumasok sa paaralan at makihalubilo sa mga "
        "bata sa bayan. Malaki ang aking kagalakan, ngunit kalakip nito ang matinding pag-aalala para sa kaniyang kalagayan. Kaya ipinagbibilin ko sa iyo na "
        "ituro mo sa kaniya ang wastong pag-iingat sa katawan at sa mga gawi sa pakikiharap sa kapuwa tao, sapagkat ang kalinisan ng katawan ay salamin ng "
        "kalinisan ng kaluluwa.\n\n"
        "Ipagbilin mo sa kaniya na pagkagising sa umaga ay maghugas agad ng mukha, maglinis ng bibig at ngipin, at magsuklay ng buhok bago humarap sa altar "
        "upang magdasal. Huwag siyang haharap sa kaniyang mga magulang o sa guro na marumi ang mga kamay at kuko, o gusot ang kaniyang kasuotan. Ang "
        "karumihan ay nagbubunga ng pandidiri sa kapuwa at nagpapakita ng kawalan ng paggalang sa mga taong kaniyang kaharap.\n\n"
        "Sa paglalakad sa lansangan, huwag siyang magpapakita ng kadunguan o magtatakbo na parang walang pinag-aralan. Kung makikipag-usap sa mga nakatatanda, "
        "ayusin ang tindig, huwag magkakamot ng ulo o mangungulangot, at huwag tititig nang bastos sa mukha ng kausap. Sa hapag-kainan naman, huwag siyang "
        "mag-uunahan sa pagsandok, huwag pupunuin ang bibig ng pagkain na parang nagmamadali, at huwag magsasalita habang may laman ang bibig. Ang taong "
        "malinis kumilos ay kinalulugdan ng Diyos at iginagalang ng kaniyang kapuwa mamamayan.\n\n"
        "Ingatan mo, Feliza, na ituro sa kaniya na ang kalinisan ay hindi pagmamayabang o pagnanais na magmukhang mayaman, kundi isang banal na tungkulin "
        "upang maiwasan ang kapahamakan ng katawan at mapanatili ang kapayapaan ng loob. Hanggang dito na lamang, at nawa'y ingatan kayo ng Maykapal.\n\n"
        "ANG IYONG NAGMAMAHAL NA KAPATID,\nURBANA\""
    )

    add_body_p(
        doc,
        "Ang siping ito ay nagpapakita ng masinsing ugnayan ng personal na kalinisan at ng panlipunang kaayusan noong panahon ng kolonyalismo. "
        "Mapapansin na ang bawat payo ni Urbana ay may kalakip na dahilan: ang paglilinis ng katawan ay hindi lamang para sa sarili kundi tanda ng paggalang "
        "sa kapuwa at sa Diyos. Ipinapakita rin nito ang papel ng tahanan at ng kababaihan bilang unang guro ng disiplina at moralidad sa lipunan."
    )

    add_callout_box(
        doc,
        "MGA GABAY NA TANONG SA PAG-UNAWA SA LUNSARAN:\n"
        "1. Sino ang sumulat ng liham, kanino ito ipinadala, at sino ang pangunahing pinatutungkulan ng mga tagubilin sa kalinisan?\n"
        "2. Ano-anong tiyak na gawi sa umaga, sa lansangan, at sa hapag-kainan ang ibinilin ni Urbana para sa kapatid na si Honesto?\n"
        "3. Bakit sinabi ni Urbana na ang kalinisan ay hindi pagmamayabang kundi paggalang sa kapuwa at sa Diyos?\n"
        "4. Alin sa mga payong ito ang sa tingin mo ay makabuluhan pa rin sa kasalukuyang panahon, at alin ang nangangailangan ng panibagong pagsusuri?"
    )

    # doc.add_page_break() [REMOVED TO PREVENT ORPHAN BLANK PAGES]

    # =========================================================================
    # PAHINA 32: SOCRATIC NA TALAKAYAN - KAGANDAHANG-ASAL AT ETIKA
    # =========================================================================
    add_heading_2(doc, "SOCRATIC NA TALAKAYAN: Kagandahang-Asal, Wika, at Etika ng Komunikasyon")
    add_body_p(doc, "Mapanuring Talakayan sa Pagitan ng Guro at Mag-aaral tungkol sa Didaktikong Panitikan", italic_prefix="")

    socratic_dialogue_15 = [
        ("Guro", "Kung babasahin natin ang liham ni Urbana, masasabi ba nating ang kaniyang mga payo ay para lamang sa mga mayayaman o para sa lahat ng tao?"),
        ("Mag-aaral 1 (Paolo)", "Sa unang tingin po, parang para sa lahat dahil mahalaga ang kalinisan. Pero kung iisipin ang kasaysayan noong 1864, ang mga pamilyang nakapupunta sa kolehiyo sa Maynila tulad ni Urbana ay kabilang sa maykayang principalia."),
        ("Guro", "Napakatabil ng iyong obserbasyon, Paolo! Ibig sabihin, ang pamantayan ng 'kagandahang-asal' ay madalas ding hinuhubog ng uring panlipunan. Ngunit masama ba ang maging malinis at magalang kahit mahirap ang isang tao?"),
        ("Mag-aaral 2 (Lina)", "Hindi po masama, Guro. Sa katunayan, sinabi ni Urbana na ang kalinisan ay hindi pagmamayabang o pagnanais na magmukhang mayaman. Ang layunin ay paggalang sa sarili at sa kapuwa."),
        ("Guro", "Mahusay, Lina! Ngayon, iuugnay natin ito sa ating aralin sa panayam. Kung ikaw ay magsasagawa ng panayam tungkol sa kalinisan sa inyong paaralan, paano mo gagamitin ang aral ni Urbana sa paraan ng iyong pakikipag-usap?"),
        ("Mag-aaral 3 (Joshua)", "Kailangan pong maging malinis din ang ating pananalita at intensiyon. Ang tagapanayam ay dapat magalang, nakikinig, hindi nagmamataas, at hindi gumagamit ng mga tanong na nagpapahiya sa kapanayam."),
        ("Guro", "Eksakto, Joshua! Ang etika ng panayam ay maihahalintulad sa kalinisan ng budhi: hindi ka naghahanap ng paninira kundi naghahanap ka ng katotohanan. At paano naman sa pagwawasto ng diyalogo sa komiks?"),
        ("Mag-aaral 4 (Mika)", "Sa komiks po, ang wastong gramatika at bantas ay tanda rin ng kalinisan sa pagsulat. Kapag magulo ang diyalogo at mali ang bantas, parang marumi ang mensahe at nahihirapang magbasa ang kapuwa mag-aaral."),
        ("Guro", "Napakagandang paglalagom, Mika! Ang kalinisan ay hindi lamang sabon at tubig; ito ay kaayusan sa kilos, katapatan sa panayam, linaw sa gramatika, at dignidad sa pakikitungo sa kapuwa tao.")
    ]

    for speaker, text in socratic_dialogue_15:
        p = doc.add_paragraph()
        format_paragraph(p, space_before=2, space_after=3)
        r_spk = p.add_run(f"{speaker}: ")
        r_spk.font.name = 'Cambria'
        r_spk.font.size = Pt(10)
        r_spk.font.bold = True
        r_spk.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D) if "Guro" in speaker else RGBColor(0x8B, 0x1A, 0x1A)
        r_txt = p.add_run(f"\"{text}\"")
        r_txt.font.name = 'Cambria'
        r_txt.font.size = Pt(10)

    add_callout_box(
        doc,
        "HAMONG SOCRATIC (Etika at Wika):\n"
        "Sumulat ng isang masinsing talata (5–7 pangungusap) bilang tugon sa hamon:\n"
        "\"Bakit ang kasanayan sa pakikinig at pagbuo ng neutral na tanong sa panayam ay higit na makapangyarihan kaysa kakayahang magsalita nang mabulaklak?\"\n"
        "Gamitin ang mga konseptong: etikal na pamamahayag, katapatan sa datos, paggalang sa kapanayam, at pag-iwas sa pagkiling."
    )

    # doc.add_page_break() [REMOVED TO PREVENT ORPHAN BLANK PAGES]

    # =========================================================================
    # PAHINA 33: PAKSA 15.3 - KABABAIHAN AT PAGWAWASTO NG DIYALOGO
    # =========================================================================
    add_heading_2(doc, "PAKSA 15.3: Kababaihan sa Tekstong Biswal at Pagwawasto ng Diyalogo")
    
    add_body_p(
        doc,
        "Ang representasyon ng kababaihan sa Panahon ng Pananakop ng Espanya ay isa sa mga paksang nangangailangan ng masusing pagwawasto sa mga "
        "karaniwang stereotype. Madalas nating marinig ang pariralang 'parang Maria Clara' na ginagamit upang ilarawan ang babaeng mahina, tahimik, "
        "laging umiiyak, at nakakulong lamang sa tahanan. Subalit kung susuriin ang kasaysayan at iba't ibang batis, makikita na ang kababaihang Pilipino "
        "ay nagpamalas ng pambihirang tatag, talino sa pagnenegosyo, pamumuno sa pamayanan, at maging pakikibaka sa rebolusyon (tulad nina Gabriela Silang "
        "at Tandang Sora)."
    )

    add_body_p(
        doc,
        "Sa pagsusuri ng mga tekstong biswal para sa comic book brochure, kailangang maging maingat sa paglalarawan sa mga babaeng tauhan. Hindi sila dapat "
        "iguhit bilang pasibong dekorasyon o mga tauhang walang sariling desisyon. Ang mga babae sa komiks ay dapat may aktibong papel sa tagpo—nagsasalita, "
        "nagpapasya, nagtuturo, at nakikipagtulungan sa paglutas ng mga suliranin ng komunidad. Ang kanilang kasuotan (baro't saya, alampay) ay dapat "
        "itanghal nang may dignidad at historikal na katumpakan."
    )

    visual_prompt_15 = (
        "PROMPT: An empowering historical comic book panel illustration depicting two intelligent Filipino young women in late 19th-century Paombong, Bulacan. "
        "In the center foreground, Feliza, dressed in an elegant crisp pina baro and checkered saya, sitting at a wooden writing desk next to a capiz-shell window, "
        "holding an ink pen while thoughtfully discussing an open letter with her younger brother Honesto, who is looking up with genuine respect. "
        "Through the open window in the background, a sunny town plaza with native women actively managing market trade stalls and school children reading. "
        "Warm golden hour illumination, realistic textures of woven textiles and polished hardwood, dignified and historically authentic composition --ar 16:9 --v 6.0"
    )
    add_prompt_box(doc, "KONSEPTONG BISWAL AT PROMPT SA COMIC PANEL (DIGNIDAD AT KABABAIHAN)", visual_prompt_15)

    add_body_p(
        doc,
        "Kasabay ng visual framing ang pagwawasto ng gramatika sa speech balloons. Dahil limitado ang espasyo sa komiks, ang bawat salita ay dapat "
        "may bigat at tungkulin. Ang pagwawasto ay sumasaklaw sa: (1) Tamang baybay at bantas (wastong gamit ng kuwit, tandang pananong, at tuldok); "
        "(2) Angkop na gamit ng panghalip upang hindi maging malabo ang tinutukoy; (3) Konsistensi ng tono ayon sa ugnayan ng mga tauhan; at "
        "(4) Pag-iwas sa labis na haba upang maging madaling basahin."
    )

    headers_15_3 = ["Aspekto ng Diyalogo", "Karaniwang Problema sa Borador", "Wastong Pagwawasto (Rebisyon)", "Paliwanag sa Pagbabago"]
    data_15_3 = [
        ["Gamit ng Bantas", "\"Ate bakit ko kailangang maglinis araw araw\"", "\"Ate, bakit ko kailangang maglinis araw-araw?\"", "Nilagyan ng kuwit pagkatapos ng pantawag (Ate) at tandang pananong sa dulo."],
        ["Malabong Panghalip", "\"Sinabi niya sa kaniya na dapat sundin niya iyon.\"", "\"Sinabi ni Urbana kay Feliza na dapat sundin ang tagubilin.\"", "Pinalitan ng tiyak na pangngalan ang magkakasunod na panghalip upang luminaw."],
        ["Labis na Haba", "\"Gusto ko sanang sabihin sa iyo na napakahalaga talaga na huwag mong kalilimutan ang aking bilin...\"", "\"Feliza, ingatan mong huwag malimot ang aking bilin sa kalinisan.\"", "Pinaikli at ginawang matipid ang pananalita nang hindi nawawala ang diwa."],
        ["Kaangkupan ng Boses", "\"Edi wow ate ikaw na ang magaling!\"", "\"Opo, Ate Urbana, susundin ko po ang iyong payo.\"", "Inalis ang modernong balbal at ginamit ang magalang at natural na tono ng panahon."]
    ]
    add_custom_table(doc, headers_15_3, data_15_3)

    add_callout_box(
        doc,
        "MUNTING PAGSASANAY 15.3 (Pag-edit ng Diyalogo sa Komiks):\n"
        "1. Bakit kailangang maging maingat sa paggamit ng mga modernong balbal kapag sumusulat ng historikal na diyalogo sa komiks?\n"
        "2. Rebisahin ang sumusunod na linya upang maging matipid at wasto ang bantas: \"Honesto dapat ka maghugas ng kamay mo bago kumain kasi baka magkasakit ka eh!\"\n"
        "3. Paano nakatutulong ang speech balloon sa pagpapakita ng ahensiya at katalinuhan ng mga babaeng tauhan sa kuwento?"
    )

    # doc.add_page_break() [REMOVED TO PREVENT ORPHAN BLANK PAGES]

    # =========================================================================
    # PAHINA 34: MGA GAWAING PAMPAGKATUTO 15.1 AT 15.2
    # =========================================================================
    add_heading_2(doc, "MGA GAWAING PAMPAGKATUTO SA ARALIN 15")
    add_heading_3(doc, "Gawain 15.1: Payo, Pagpapahalaga, at Kontekstong Pangkasaysayan")
    add_body_p(
        doc,
        "Panuto: Pumili ng dalawang tiyak na tagubilin mula sa liham ni Urbana sa pahina 31 (halimbawa: gawi sa hapag-kainan o gawi sa lansangan). "
        "Suriin ang mga ito gamit ang sumusunod na analytical matrix upang maipakita ang lalim ng pagpapahalaga at konteksto."
    )

    headers_g15_1 = ["Tiyak na Tagubilin ni Urbana", "Dahilang Ibinigay sa Liham", "Pagpapahalagang Nasa Likod (Values)", "Pagsusuri sa Konteksto at Kasalukuyan"]
    data_g15_1 = [
        ["Huwag magsasalita habang may laman ang bibig sa hapag-kainan", "Upang hindi makapagdulot ng pandidiri sa kapuwa kumakain", "Paggalang sa kapuwa; disiplina sa sarili; kaayusan sa pagkain", "Nananatiling mahalagang pamantayan sa etika ng pagkain at kalinisan ngayon."],
        ["Maghugas ng kamay at magsuklay bago humarap sa magulang at guro", "Upang hindi magmukhang marumi at walang pinag-aralan", "Pagpapahalaga sa dignidad ng pamilya at paggalang sa awtoridad", "Pundasyon ng personal hygiene at pagpapahalaga sa sariling kaanyuan."],
        ["[Pumili ng sariling tagubilin mula sa teksto]", "[Isulat ang dahilan mula sa liham]", "[Tukuyin ang pagpapahalaga]", "[Ipaliwanag ang kabuluhan ngayon]"]
    ]
    add_custom_table(doc, headers_g15_1, data_g15_1)

    add_heading_3(doc, "Gawain 15.2: Mula Panayam tungo sa Pagsulat ng Maikling Balita")
    add_body_p(
        doc,
        "Sitwasyon: Isipin na ikaw ay tagapagbalita ng inyong pahayagang pangkampus. Nagsagawa ka ng panayam sa tagapangasiwa ng programang "
        "'Oplan Linis-Paaralan' at sa isang kinatawan ng mga mag-aaral. Narito ang iyong nakalap na mga tala:\n"
        "• Tagapangasiwa: 'Nagsimula ang programa noong Setyembre 1. Layunin naming turuan ang mga bata ng paghihiwalay ng basura (segregation). Sa unang linggo, 70% ng klase ang sumunod.'\n"
        "• Mag-aaral: 'Maganda po ang layunin, pero minsan po kulang ang mga basurahan sa bawat palapag kaya nahihirapan ang mga kaklase ko.'\n\n"
        "Gawain: Sumulat ng isang maikling balita (100–120 salita) batay LAMANG sa mga nakalap na impormasyon sa panayam. "
        "Gumamit ng angkop na ulo ng balita (headline), sumunod sa baligtad na piramide (lead paragraph), at tiyaking neutral at tapat ang ulat."
    )

    add_callout_box(
        doc,
        "TUNTUNIN SA PAGSULAT NG BALITA:\n"
        "Huwag magdagdag ng sariling opinyon o akusasyon na hindi sinabi ng mga kinapanayam. Ang integridad ng balita ay nakasalalay sa katapatan sa datos."
    )

    # doc.add_page_break() [REMOVED TO PREVENT ORPHAN BLANK PAGES]

    # =========================================================================
    # PAHINA 35: GAWAIN 15.3 - EDITOR NG COMIC BROCHURE AT RUBRIK
    # =========================================================================
    add_heading_2(doc, "Gawain 15.3: Editor ng Comic Book Brochure — Rebisyon ng Diyalogo")
    add_body_p(
        doc,
        "Gawain: Ikaw ang punong editor ng comic book brochure. Natanggap mo ang sumusunod na burador ng diyalogo para sa tatlong panel mula sa isang "
        "baguhang manunulat. Masyadong magulo ang gramatika, gumagamit ng mga modernong balbal, at hindi angkop ang tono sa historikal na tagpo. "
        "Isagawa ang propesyonal na rebisyon sa pamamagitan ng pagwawasto sa gramatika, bantas, at pagbabalik ng tamang boses ng mga tauhan."
    )

    add_body_p(
        doc,
        "Orihinal na Borador ng Manunulat:\n"
        "• Panel 1 (Honesto): \"Ate Feliza bakit ba kailangan ko pa maghugas ng kamay eh gutom na gutom na kaya ako grabe naman!\"\n"
        "• Panel 2 (Feliza): \"Honesto sinabi kasi ni Ate Urbana na dapat daw malinis ka palagi para hindi ka maging kadiri tignan ng mga tao sa labas.\"\n"
        "• Panel 3 (Honesto): \"Ah ganun ba sige na nga maghuhugas na ako para bida ako sa eskwela bukas!\"\n\n"
        "Iyong Pinahusay na Rebisyon (Wasto, Matipid, at may Dignidad):\n"
        "• Panel 1 (Honesto): _____________________________________________________________________________________________________\n"
        "• Panel 2 (Feliza): ______________________________________________________________________________________________________\n"
        "• Panel 3 (Honesto): _____________________________________________________________________________________________________"
    )

    headers_rubrik_15 = ["Pamantayan", "Napakahusay (4)", "Mahusay (3)", "Nalilinang (2)", "Nangangailangan ng Gabay (1)"]
    data_rubrik_15 = [
        ["Katumpakan ng Gramatika at Bantas (25%)", "Ganap na wasto ang baybay, bantas (kuwit, tuldok, tandang pananong), at kayarian ng pangungusap.", "Wasto ang halos lahat ng pangungusap; may isa o dalawang maliliit na pagkukulang sa bantas.", "May ilang kapansin-pansing kamaliang panggramatika na nakaaapekto sa daloy.", "Maraming mali sa gramatika at bantas na nakahahadlang sa pag-unawa."],
        ["Kaangkupan ng Boses at Tono (25%)", "Napakaganda at natural ng tono; angkop na angkop sa ugnayan ng magkapatid sa kontekstong historikal.", "Maayos at natural ang pananalita ng mga tauhan.", "Medyo pormal o hindi natural pakinggan para sa isang bata o nakatatandang kapatid.", "Hindi angkop ang tono; gumagamit ng hindi nararapat na modernong salita."],
        ["Kalinawan at Kaiklian ng Mensahe (25%)", "Napakalinaw at matipid sa salita; sakto ang haba para sa speech balloon nang hindi sumisikip.", "Malinaw ang mensahe at kasya sa karaniwang speech balloon.", "Medyo mahaba ang ilang linya kaya maaaring matakpan ang drowing.", "Masyadong mahaba at magulo ang pahayag para sa komiks."],
        ["Representasyon at Dignidad (25%)", "Mapanuri at may mataas na paggalang sa pagkatao at relasyon ng pamilya nang walang stereotype.", "Magalang at responsable ang pangkalahatang mensahe.", "May kaunting kababawan sa pagpapaliwanag ng aral.", "Nagpapakita ng kawalang-galang o mapanghusgang kaisipan sa tauhan."]
    ]
    add_custom_table(doc, headers_rubrik_15, data_rubrik_15)

    # doc.add_page_break() [REMOVED TO PREVENT ORPHAN BLANK PAGES]

    # =========================================================================
    # PAHINA 36: PAGNINILAY, EXIT TICKET, AT BUOD NG ARALIN 15
    # =========================================================================
    add_heading_2(doc, "PAGNINILAY, EXIT TICKET, AT BUOD NG ARALIN 15")
    add_heading_3(doc, "Metakognitibong Pagninilay sa Aralin 15")
    add_body_p(
        doc,
        "Panuto: Matapos ang komprehensibong pagtalakay sa Urbana at Feliza, etika ng panayam, at rebisyon ng komiks, pagnilayan ang iyong mga natutuhan "
        "sa pamamagitan ng pagkumpleto sa mga sumusunod na pangungusap."
    )

    reflection_prompts_15 = [
        "1. Ang pinakamahalagang aral na aking natutuhan tungkol sa tunay na kahulugan ng 'kalinisan' ay __________________________________________________________________________________________________.",
        "2. Sa pagsasagawa ng panayam para sa balita, titiyakin kong neutral ang aking mga tanong dahil __________________________________________________________________________________________________.",
        "3. Natutuhan ko na sa pag-edit ng diyalogo sa komiks, hindi sapat na tama ang gramatika; kailangan ding __________________________________________________________________________________________________.",
        "4. Binago ng araling ito ang aking pananaw sa kababaihan sa kasaysayan sa pamamagitan ng __________________________________________________________________________________________________."
    ]
    for prm in reflection_prompts_15:
        add_body_p(doc, prm)

    add_heading_3(doc, "Exit Ticket: 3 Bagay na Dadalhin Ko sa Pinal na Pagtataya")
    add_body_p(
        doc,
        "• 1 Prinsipyo sa Mapanuring Panayam: _____________________________________________________________________________________\n"
        "• 1 Tuntunin sa Pagwawasto ng Diyalogo: ___________________________________________________________________________________\n"
        "• 1 Pagpapahalaga mula kay Urbana na Aking Isasabuhay: ____________________________________________________________________"
    )

    add_callout_box(
        doc,
        "SINTESIS AT BUOD NG ARALIN 15:\n"
        "1. Ang Urbana at Feliza ay nagtatanghal ng kahalagahan ng kagandahang-asal bilang gabay sa pakikipagkapuwa at moral na kaayusan. Ang kalinisan ay "
        "hindi panlabas lamang kundi repleksiyon ng kalinisan ng budhi at paggalang sa dignidad ng kapuwa tao.\n"
        "2. Ang panayam ay pundasyon ng tumpak na balita. Ang responsableng tagapanayam ay gumagamit ng mga bukas at neutral na tanong, umiiwas sa pangunguna, "
        "at nagbeberipika ng datos upang maiwasan ang maling impormasyon.\n"
        "3. Ang representasyon ng kababaihan sa mga historikal na akda at biswal ay dapat magtampok sa kanilang ahensiya, talino, at katatagan sa halip "
        "na itali sila sa mga pasibong stereotype ng kahinaan.\n"
        "4. Ang pagwawasto ng gramatika at bantas sa comic book brochure ay mahalagang kasanayan upang matiyak na malinaw, matipid, natural, at mabisa "
        "ang ugnayan ng teksto at biswal sa paghahatid ng makabuluhang mensahe."
    )

    # doc.add_page_break() [REMOVED TO PREVENT ORPHAN BLANK PAGES]

    # =========================================================================
    # PAHINA 37: MABILISANG PAGTATAYA SA ARALIN 15 (BAHAGI I: AYTEM 1-8)
    # =========================================================================
    add_heading_2(doc, "MABILISANG PAGTATAYA SA ARALIN 15 (15 PUNTOS)")
    add_heading_3(doc, "Bahagi I: Mapanuring Pagpili (Aytem 1 hanggang 8)")
    add_body_p(doc, "Panuto: Piliin ang titik ng pinakatumpak na sagot. Isulat ang malaking titik sa patlang bago ang bawat bilang.")

    mcq_items_15_p1 = [
        ("1. Ano ang tawag sa anyong pampanitikan ng 'Urbana at Feliza' na gumagamit ng palitan ng mga liham sa pagitan ng mga tauhan?",
         ["A. Epistolaryo", "B. Alegoriko", "C. Pikaresko", "D. Tulang Pasalaysay"]),
        ("2. Sino ang paring may-akda ng klasikong aklat na 'Pagsusulatan ng Dalawang Binibini na si Urbana at si Feliza' (1864)?",
         ["A. P. Gaspar Aquino de Belen", "B. P. Mariano Pilapil", "C. P. Modesto de Castro", "D. P. Pedro Pelaez"]),
        ("3. Ayon sa liham ni Urbana kay Feliza, bakit dapat maglinis ng katawan at magsuklay ng buhok bago humarap sa magulang o guro?",
         ["A. Upang magmukhang mayaman at makahingi ng pabor.",
          "B. Upang gayahin ang mga banyagang nakikita sa mga larawan.",
          "C. Dahil may parusang kulong ang sinumang marumi ang damit.",
          "D. Sapagkat ang karumihan ay nagpapakita ng kawalan ng paggalang sa kapuwa at nagdudulot ng pandidiri."]),
        ("4. Alin sa mga sumusunod ang PINAKAMAHUSAY na halimbawa ng bukas at neutral na tanong sa isang panayam para sa balita?",
         ["A. \"Hindi ba't napakahusay ng inyong pamamalakad sa paaralan?\"",
          "B. \"Paano ninyo ilalarawan ang mga naging tagumpay at hamon sa pagpapatupad ng bagong patakaran?\"",
          "C. \"Kayo po ang may kasalanan kung bakit marumi ang liwasan, tama?\"",
          "D. \"Masaya ang lahat sa programa, hindi ba?\""]),
        ("5. Ano ang pangunahing layunin ng isang 'follow-up question' sa panayam?",
         ["A. Baguhin agad ang paksa kapag hindi nagustuhan ang sagot.",
          "B. Pilitin ang kinakapanayam na sumang-ayon sa sariling pananaw ng tagapagbalita.",
          "C. Palalimin, linawin, o kumuha ng kongkretong ebidensiya at datos batay sa naunang sagot ng kinakapanayam.",
          "D. Tapusin nang mabilis ang panayam upang makauwi na."]),
        ("6. Bakit mapanganib ang paggamit ng 'leading question' (nangungunang tanong) sa pamamahayag?",
         ["A. Dahil idinidikte nito ang sagot at nawawala ang pagiging patas, obhetibo, at mapagkakatiwalaan ng balita.",
          "B. Dahil nagiging masyadong mahaba ang artikulo sa pahayagan.",
          "C. Dahil baka hindi maintindihan ng mambabasa ang mga salita.",
          "D. Dahil mas mahal ang bayad sa tagapanayam kapag ganoon ang tanong."]),
        ("7. Paano wawasakin ng isang responsableng manunulat ng komiks ang stereotype tungkol sa kababaihan noong panahon ng Espanyol?",
         ["A. Huwag nang maglagay ng babaeng tauhan sa kuwento.",
          "B. Ipakita silang laging umiiyak at naghihintay ng tulong sa bawat panel.",
          "C. Baguhin ang kanilang kasuotan patungo sa makabagong kasuotan sa kasalukuyan.",
          "D. Ilarawan ang mga babae na may sariling ahensiya, talino, aktibong pagpapasya, at ambag sa pamayanan."]),
        ("8. Alin sa mga sumusunod ang tamang rebisyon ng linyang: \"Feliza maghugas ka ng kamay mo ngayon na para malinis ka!\"?",
         ["A. \"Feliza maghugas ka kamay ngayon na.\"",
          "B. \"Feliza, maghugas ka ng iyong mga kamay ngayon upang mapanatili ang kalinisan.\"",
          "C. \"Edi maghugas ka na lang Feliza para tapos na!\"",
          "D. \"Hugas kamay ka na Feliza dali!\""])
    ]

    for q_text, choices in mcq_items_15_p1:
        add_body_p(doc, q_text, bold_prefix="")
        for ch in choices:
            p_c = doc.add_paragraph()
            format_paragraph(p_c, space_before=1, space_after=1)
            p_c.paragraph_format.left_indent = Inches(0.2)
            r_c = p_c.add_run(ch)
            r_c.font.name = 'Cambria'
            r_c.font.size = Pt(9.5)

    # doc.add_page_break() [REMOVED TO PREVENT ORPHAN BLANK PAGES]

    # =========================================================================
    # PAHINA 38: MABILISANG PAGTATAYA (BAHAGI I KONT. AT BAHAGI II: PAGSUSURI)
    # =========================================================================
    add_heading_2(doc, "MABILISANG PAGTATAYA SA ARALIN 15 (PAGPAPATULOY)")
    add_heading_3(doc, "Bahagi I (Pagpapatuloy: Aytem 9 at 10)")

    mcq_items_15_p2 = [
        ("9. Ano ang ibig sabihin ng prinsipyo na ang 'kalinisan ay salamin ng kaluluwa' ayon sa kaisipan ni Urbana?",
         ["A. Na ang panlabas na kaayusan at disiplina sa katawan ay nagmumula sa dalisay na budhi at paggalang sa Diyos at kapuwa.",
          "B. Na ang taong may mamahaling sabon ay tiyak na maliligtas sa kabilang buhay.",
          "C. Na kailangang magsuot ng puting damit araw-araw.",
          "D. Na hindi na kailangang magdasal kung naligo na sa umaga."]),
        ("10. Sa pag-edit ng diyalogo sa isang speech balloon, bakit mahalagang maging matipid sa salita?",
         ["A. Upang makatipid sa tinta ng bolpen o printer.",
          "B. Dahil bawal ang mahahabang pangungusap sa wikang Filipino.",
          "C. Upang hindi matakpan ang mahalagang visual elements ng larawan at maging madaling basahin ang mensahe.",
          "D. Upang magmukhang misteryoso ang tauhan sa komiks."])
    ]

    for q_text, choices in mcq_items_15_p2:
        add_body_p(doc, q_text, bold_prefix="")
        for ch in choices:
            p_c = doc.add_paragraph()
            format_paragraph(p_c, space_before=1, space_after=1)
            p_c.paragraph_format.left_indent = Inches(0.2)
            r_c = p_c.add_run(ch)
            r_c.font.name = 'Cambria'
            r_c.font.size = Pt(9.5)

    add_heading_3(doc, "Bahagi II: Masinsing Pagsusuri at Pagbibigay-Katwiran (Aytem 11 hanggang 15)")
    add_body_p(doc, "Panuto: Sagutin ang bawat katanungan sa 3 hanggang 5 masinsing pangungusap. (1 puntos bawat aytem)")

    open_items_15 = [
        "11. Ipaliwanag kung bakit ang pag-unawa sa konteksto ng ika-19 na siglo ay hindi nangangahulugang dapat nating tanggapin ang lahat ng lumang pamantayan sa kasalukuyan.",
        "12. Sumulat ng isang bukas at neutral na tanong para sa panayam tungkol sa kalinisan sa inyong barangay, at ipaliwanag kung bakit neutral ito.",
        "13. Bakit itinuturing na may pananagutang moral ang isang tagapagbalita sa paraan ng kaniyang pag-quote o pag-paraphrase sa sinabi ng kinapanayam?",
        "14. Magbigay ng dalawang patunay mula sa kasaysayan na sumasalungat sa stereotype na ang mga kababaihan noong panahon ng Espanyol ay walang sariling lakas at talino.",
        "15. Sa pagwawasto ng diyalogo sa komiks, bakit sinasabing ang tamang bantas (tulad ng kuwit pagkatapos ng pantawag) ay nakatutulong sa pagiging natural ng boses ng tauhan?"
    ]

    for o_q in open_items_15:
        p_oq = doc.add_paragraph()
        format_paragraph(p_oq, space_before=3, space_after=2)
        r_oq = p_oq.add_run(o_q)
        r_oq.font.name = 'Cambria'
        r_oq.font.size = Pt(10)
        r_oq.font.bold = True

    add_callout_box(
        doc,
        "PAMANTAYAN SA PAGMAMARKA NG BAHAGI II:\n"
        "• 1.0 Puntos: Tumpak at malalim ang paliwanag, may maayos na lohika, at nagtataglay ng tiyak na ebidensiya mula sa aralin.\n"
        "• 0.5 Puntos: Bahagyang wasto ngunit kulang sa pagpapalalim o walang halimbawang sumusuporta.\n"
        "• 0.0 Puntos: Hindi tumutugon sa katanungan o mali ang ibinigay na katuwiran."
    )

    # doc.add_page_break() [REMOVED TO PREVENT ORPHAN BLANK PAGES]

    # =========================================================================
    # PAHINA 39: ARALIN 15 TRANSITION & PAGHAHANDA SA PANGWAKAS NA PAGTATAYA
    # =========================================================================
    add_heading_2(doc, "PAGLALAGOM NG ARALIN 15 AT PAGHAHANDA SA PANGWAKAS NA PAGTATAYA")
    
    add_body_p(
        doc,
        "Sa pagtatapos ng Aralin 15, nakumpleto natin ang tatlong pangunahing aralin ng Yunit III. Mula sa kaligirang pangkasaysayan at pamamahayag sa "
        "Aralin 13, patungo sa kapangyarihan ng Pasyon at tradisyong pasalita sa Aralin 14, hanggang sa kagandahang-asal, etika ng panayam, at rebisyon "
        "ng komiks sa Aralin 15—napatunayan natin na ang panitikan ay isang buháy, dinamiko, at makabuluhang salamin ng pagkataong Pilipino."
    )

    add_body_p(
        doc,
        "Nakahanda na ngayon ang mga mag-aaral na harapin ang Pangwakas na Pagtataya sa Yunit III (Pahina 40 hanggang 44). Ang pagtatayang ito ay "
        "susubok sa inyong kakayahang magbasa nang may konteksto, magsuri ng bagong datos, magpahayag ng sariling katuwiran sa isang organisadong "
        "sanaysay, at bumuo ng pinal na multimodal na comic book brochure alinsunod sa pamantayan ng DepEd MATATAG K-10 Kurikulum."
    )

    add_callout_box(
        doc,
        "MGA PAALALA BAGO HARAPIN ANG PANGWAKAS NA PAGTATAYA (40 PUNTOS):\n"
        "1. Bahagi I (10 Puntos): Pagsusuri sa isang bagong tekstong pangkasaysayan gamit ang paghihiwalay ng datos laban sa interpretasyon.\n"
        "2. Bahagi II (10 Puntos): Paglalapat ng mga konsepto sa wika, balita, at multimodal na komunikasyon.\n"
        "3. Bahagi III (10 Puntos): Pagsulat ng isang komprehensibo at estrukturadong sanaysay na nag-uugnay sa teksto, konteksto, at pagkatao.\n"
        "4. Bahagi IV (10 Puntos): Autentikong Gawaing Pagganap—Pinal na Comic Book Brochure gamit ang GRASPS Model at Analitikong Rubrik."
    )

    doc.add_page_break()
