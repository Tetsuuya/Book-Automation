# -*- coding: utf-8 -*-
"""
DEPED MATATAG PUBLICATION-GRADE DOCX BUILDER & COMPILER ENGINE
Use this Python script in OpenAI Code Interpreter or local Python environment
to compile complete 40+ page interactive multimedia textbooks (.docx) with
verbatim content preservation, clean mathematical typography, and benchmark styling.

Supports both:
1. Direct CLI Compilation:
   python 03_DOCX_BUILDER_SCRIPT.py <input.md> [output.docx]
2. Programmatic API:
   from 03_DOCX_BUILDER_SCRIPT import compile_markdown_to_docx, initialize_document, add_formula_box
"""

import os
import sys
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

# ==========================================
# 1. XML STYLING UTILITIES
# ==========================================

def set_cell_shading(cell, color_hex):
    """Set background color fill for a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    """Set internal cell padding (in dxa: 20 dxa = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_callout_border(cell, color_hex="1B365D", sz="36"):
    """Set a bold solid left border stripe for callout boxes (sz 36 = 4.5pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

def set_dashed_border(cell, color_hex="CCCCCC", sz="12"):
    """Set dashed border around figure specification containers."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="dashed" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
            <w:left w:val="dashed" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
            <w:bottom w:val="dashed" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
            <w:right w:val="dashed" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

def set_table_grid_borders(table, color_hex="D0D7DE", sz="4"):
    """Set crisp light-gray borders for tables."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
            <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
            <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
            <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
            <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def set_table_no_borders(table):
    """Remove all borders from a container table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="none"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

# ==========================================
# 2. MATHEMATICAL TYPOGRAPHY SANITIZER
# ==========================================

def clean_math_text(text):
    r"""
    Transforms LaTeX symbols and artifacts into clean, human-readable textbook typography.
    Never outputs raw LaTeX commands like \qquad or \frac in the Word document.
    """
    if not text:
        return ""
    
    # Horizontal spacing
    text = text.replace(r"\qquad", "    —    ")
    text = text.replace(r"\quad", "  ")
    text = text.replace(r"\hspace{1cm}", "    ")
    text = text.replace(r"\hspace{0.5cm}", "  ")
    
    # Operators and symbols
    text = text.replace(r"\times", "×")
    text = text.replace(r"\cdot", "·")
    text = text.replace(r"\div", "÷")
    text = text.replace(r"\pm", "±")
    text = text.replace(r"\pi", "π")
    text = text.replace(r"\theta", "θ")
    text = text.replace(r"\degree", "°")
    text = text.replace(r"^\circ", "°")
    text = text.replace(r"\le", "≤").replace(r"\leq", "≤")
    text = text.replace(r"\ge", "≥").replace(r"\geq", "≥")
    text = text.replace(r"\neq", "≠")
    text = text.replace(r"\approx", "≈")
    text = text.replace(r"\infty", "∞")
    
    # Fractions: \frac{a}{b} -> (a / b)
    text = re.sub(r'\\frac\{([^}]+)\}\{([^}]+)\}', r'(\1 / \2)', text)
    # Square roots: \sqrt{x} -> √(x)
    text = re.sub(r'\\sqrt\{([^}]+)\}', r'√(\1)', text)
    # Text wrappers
    text = re.sub(r'\\text\{([^}]+)\}', r'\1', text)
    # Exponents
    text = re.sub(r'\^2\b', '²', text)
    text = re.sub(r'\^3\b', '³', text)
    
    # Delimiters
    text = text.replace(r"\[", "").replace(r"\]", "")
    text = re.sub(r'\$(?!\d)', '', text.replace("$$", ""))   # keep currency like $5
    return text

# ==========================================
# 3. MODULAR API HELPERS
# ==========================================

def format_paragraph(p, space_before=0, space_after=4, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align

def set_outline_level(p, level):
    """Marks a paragraph as a heading level (0 = H1) without changing its banner styling."""
    p._p.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="{level}"/>'))

def add_page_number_field(paragraph, size=8.5):
    """Appends a live PAGE field so printed page numbers match the TOC."""
    run = paragraph.add_run()
    run.font.name = 'Cambria'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), kind)
        else:
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = text
        run._r.append(el)

def add_hyperlink(paragraph, text, url, size=10.5, bold=False, italic=False):
    """Adds a clickable external hyperlink run (used for YouTube Direct Links and sources)."""
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), r_id)
    r = OxmlElement('w:r')
    rPr = parse_xml(
        f'<w:rPr {nsdecls("w")}><w:rFonts w:ascii="Cambria" w:hAnsi="Cambria"/>'
        f'{"<w:b/>" if bold else ""}{"<w:i/>" if italic else ""}'
        f'<w:color w:val="0563C1"/><w:sz w:val="{int(size * 2)}"/><w:u w:val="single"/></w:rPr>'
    )
    r.append(rPr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    link.append(r)
    paragraph._p.append(link)

def add_heading_1(doc, text):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    set_cell_shading(c, "1B365D")
    set_cell_margins(c, top=120, bottom=120, left=180, right=180)
    set_table_no_borders(t)
    p = c.paragraphs[0]
    format_paragraph(p, space_before=2, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_outline_level(p, 0)
    run = p.add_run(clean_math_text(text).upper())
    run.font.name = 'Cambria'
    run.font.size = Pt(13.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p_sp = doc.add_paragraph()
    format_paragraph(p_sp, space_before=0, space_after=4)
    return p

def add_heading_2(doc, text):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    set_cell_shading(c, "EBF3FA")
    set_callout_border(c, color_hex="1B365D", sz="36")
    set_cell_margins(c, top=80, bottom=80, left=150, right=150)
    set_table_no_borders(t)
    p = c.paragraphs[0]
    format_paragraph(p, space_before=2, space_after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_outline_level(p, 1)
    run = p.add_run(clean_math_text(text))
    run.font.name = 'Cambria'
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    p_sp = doc.add_paragraph()
    format_paragraph(p_sp, space_before=0, space_after=3)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=6, space_after=2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(clean_math_text(text))
    run.font.name = 'Cambria'
    run.font.size = Pt(11.0)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x2A, 0x4B, 0x7C)
    return p

def add_body_p(doc, text, bold_prefix="", italic_prefix=""):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=0, space_after=3.5, line_spacing=1.15)
    if bold_prefix:
        r_b = p.add_run(clean_math_text(bold_prefix))
        r_b.font.name = 'Cambria'
        r_b.font.size = Pt(10.5)
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    if italic_prefix:
        r_i = p.add_run(clean_math_text(italic_prefix))
        r_i.font.name = 'Cambria'
        r_i.font.size = Pt(10.5)
        r_i.font.italic = True
    run = p.add_run(clean_math_text(text))
    run.font.name = 'Cambria'
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p

def add_callout_box(doc, title, items, bg_color="FAFAFA", border_color="1B365D"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, bg_color)
    set_callout_border(cell, border_color, "36")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    p = cell.paragraphs[0]
    format_paragraph(p, space_before=2, space_after=4)
    run = p.add_run(clean_math_text(title))
    run.font.name = 'Cambria'
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    for item in items:
        pi = cell.add_paragraph()
        format_paragraph(pi, space_before=1, space_after=2)
        ri = pi.add_run(clean_math_text(item))
        ri.font.name = 'Cambria'
        ri.font.size = Pt(9.5)
        ri.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    p_sp = doc.add_paragraph()
    format_paragraph(p_sp, space_before=0, space_after=3)

def add_formula_box(doc, title, formula, where_items=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, "F4F7FA")
    set_callout_border(cell, "1B365D", "36")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    p_title = cell.paragraphs[0]
    format_paragraph(p_title, space_before=2, space_after=3)
    r_t = p_title.add_run(clean_math_text(title))
    r_t.font.name = 'Cambria'
    r_t.font.size = Pt(10.5)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    p_eq = cell.add_paragraph()
    format_paragraph(p_eq, space_before=2, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_eq = p_eq.add_run(clean_math_text(formula))
    r_eq.font.name = 'Cambria'
    r_eq.font.size = Pt(12.0)
    r_eq.font.bold = True
    r_eq.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    if where_items:
        p_w = cell.add_paragraph()
        format_paragraph(p_w, space_before=2, space_after=1)
        rw = p_w.add_run("Where:")
        rw.font.name = 'Cambria'
        rw.font.size = Pt(9.5)
        rw.font.bold = True
        rw.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        for w in where_items:
            pw = cell.add_paragraph()
            format_paragraph(pw, space_before=1, space_after=1)
            rw2 = pw.add_run(f"• {clean_math_text(w)}")
            rw2.font.name = 'Cambria'
            rw2.font.size = Pt(9.0)
            rw2.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    p_sp = doc.add_paragraph()
    format_paragraph(p_sp, space_before=0, space_after=3)

def add_worked_example_box(doc, example_title, problem_text, given_items, formula_used, solution_steps, final_answer, verification=""):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, "F8FAFC")
    set_callout_border(cell, "2A4B7C", "36")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    p_title = cell.paragraphs[0]
    format_paragraph(p_title, space_before=2, space_after=3)
    r_t = p_title.add_run(f"💡 {clean_math_text(example_title)}")
    r_t.font.name = 'Cambria'
    r_t.font.size = Pt(11.0)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    p_prob = cell.add_paragraph()
    format_paragraph(p_prob, space_before=2, space_after=3)
    rp_lbl = p_prob.add_run("Problem: ")
    rp_lbl.font.bold = True
    rp_lbl.font.color.rgb = RGBColor(0x2A, 0x4B, 0x7C)
    rp_txt = p_prob.add_run(clean_math_text(problem_text))
    rp_txt.font.name = 'Cambria'
    rp_txt.font.size = Pt(10.0)
    
    if given_items:
        p_g = cell.add_paragraph()
        format_paragraph(p_g, space_before=1, space_after=2)
        r_g = p_g.add_run(f"Step 1 (Given): {', '.join(clean_math_text(g) for g in given_items)}")
        r_g.font.name = 'Cambria'
        r_g.font.size = Pt(9.5)
        r_g.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        
    if formula_used:
        p_f = cell.add_paragraph()
        format_paragraph(p_f, space_before=1, space_after=2)
        rf_lbl = p_f.add_run("Step 2 (Formula / Law): ")
        rf_lbl.font.bold = True
        rf_lbl.font.color.rgb = RGBColor(0x2A, 0x4B, 0x7C)
        rf_txt = p_f.add_run(clean_math_text(formula_used))
        rf_txt.font.name = 'Cambria'
        rf_txt.font.size = Pt(9.5)
        
    for s_idx, step in enumerate(solution_steps):
        ps = cell.add_paragraph()
        format_paragraph(ps, space_before=1, space_after=2)
        rs = ps.add_run(f"Step {s_idx + 3}: {clean_math_text(step)}")
        rs.font.name = 'Cambria'
        rs.font.size = Pt(9.5)
        rs.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        
    p_ans = cell.add_paragraph()
    format_paragraph(p_ans, space_before=3, space_after=2)
    ra_lbl = p_ans.add_run("✅ Final Answer: ")
    ra_lbl.font.bold = True
    ra_lbl.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    ra_txt = p_ans.add_run(clean_math_text(final_answer))
    ra_txt.font.name = 'Cambria'
    ra_txt.font.size = Pt(10.0)
    ra_txt.font.bold = True
    
    if verification:
        p_v = cell.add_paragraph()
        format_paragraph(p_v, space_before=1, space_after=2)
        rv_lbl = p_v.add_run("Verification: ")
        rv_lbl.font.bold = True
        rv_lbl.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        rv_txt = p_v.add_run(clean_math_text(verification))
        rv_txt.font.name = 'Cambria'
        rv_txt.font.size = Pt(9.0)
        rv_txt.font.italic = True
    p_sp = doc.add_paragraph()
    format_paragraph(p_sp, space_before=0, space_after=3)

# ==========================================
# 4. MARKDOWN TO DOCX PARSER ENGINE
# ==========================================

INLINE_MATH = re.compile(r'(\$\$.+?\$\$|\\\(.+?\\\)|\\\[.+?\\\]|\$(?=\S)[^$\n]+?(?<=\S)\$)')

def _is_inline_math(seg):
    if seg.startswith(('$$', '\\(', '\\[')):
        return True
    inner = seg[1:-1]
    return bool(re.search(r'[\\^_{}=+\-−×÷·<>≤≥()|]', inner) or re.fullmatch(r'[A-Za-z]', inner))  # plain "$5 and $10" money stays text

def format_inline_runs(paragraph, text, default_size=10.5, default_color=None, default_bold=False, default_italic=False):
    # Inline TeX ($...$, \(...\), \[...\]) becomes a real Word equation; other text keeps the normal run formatting.
    if '$' in text or '\\(' in text or '\\[' in text:
        for seg in INLINE_MATH.split(text):
            if not seg:
                continue
            if INLINE_MATH.fullmatch(seg) and _is_inline_math(seg):
                inner = seg[2:-2] if seg.startswith(('$$', '\\(', '\\[')) else seg[1:-1]
                add_inline_math(paragraph, inner, default_size + 0.5)
            else:
                _format_text_runs(paragraph, seg, default_size, default_color, default_bold, default_italic)
        return
    _format_text_runs(paragraph, text, default_size, default_color, default_bold, default_italic)

def _format_text_runs(paragraph, text, default_size=10.5, default_color=None, default_bold=False, default_italic=False):
    """Parses inline bold, italic, code, and symbols within a paragraph run."""
    text = clean_math_text(text)
    pattern = r'(\[[^\]]+\]\(https?://[^)\s]+\)|https?://[^\s)|]+|\*\*[^*]+\*\*|\*[^*]+\*|_[^_]+_|`[^`]+`)'
    parts = re.split(pattern, text)
    
    for part in parts:
        if not part:
            continue
        link = re.fullmatch(r'\[([^\]]+)\]\((https?://[^)\s]+)\)', part)
        if link or re.fullmatch(r'https?://[^\s)|]+', part):
            label, url = (link.group(1), link.group(2)) if link else (part, part)
            add_hyperlink(paragraph, label, url, size=default_size, bold=default_bold, italic=default_italic)
            continue
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            r = paragraph.add_run(part[2:-2])
            r.bold = True
            r.italic = default_italic
        elif ((part.startswith('*') and part.endswith('*')) or (part.startswith('_') and part.endswith('_'))) and len(part) >= 2:
            r = paragraph.add_run(part[1:-1])
            r.bold = default_bold
            r.italic = True
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            r = paragraph.add_run(part[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(default_size - 1)
            continue
        else:
            r = paragraph.add_run(part)
            r.bold = default_bold
            r.italic = default_italic
            
        r.font.name = 'Cambria'
        r.font.size = Pt(default_size)
        if default_color:
            r.font.color.rgb = default_color
        else:
            r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

def render_callout_box(doc, lines):
    """Renders blockquote callouts: YouTube modules, Figure specs, Formula cards, Worked examples."""
    all_text = " ".join(lines)
    is_youtube = any(k in all_text for k in ["INTERACTIVE MULTIMEDIA", "INTERAKTIBONG MULTIMEDIA", "🎥"])
    is_figure = any(k in all_text for k in ["Figure:", "Larawan:", "AI Generation Prompt", "Scene Specification", "🖼️", "📷"])
    is_formula = any(k in all_text for k in ["Formula Card", "FORMULA CARD", "📐"])
    is_example = any(k in all_text for k in ["Worked Example", "WORKED EXAMPLE", "💡"])
    
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    set_cell_margins(c, top=120, bottom=120, left=180, right=180)
    
    if is_youtube:
        set_cell_shading(c, "FFF5F5")
        set_callout_border(c, color_hex="CC0000", sz="36")
    elif is_figure:
        set_cell_shading(c, "FAFAFA")
        set_dashed_border(c, color_hex="CCCCCC", sz="12")
    elif is_formula:
        set_cell_shading(c, "F4F7FA")
        set_callout_border(c, color_hex="1B365D", sz="36")
    elif is_example:
        set_cell_shading(c, "F8FAFC")
        set_callout_border(c, color_hex="2A4B7C", sz="36")
    else:
        set_cell_shading(c, "F4F7FA")
        set_callout_border(c, color_hex="1B365D", sz="36")
        
    skip_to = 0
    for idx, line in enumerate(lines):
        clean_l = line.lstrip('> ').strip()
        if not clean_l or idx < skip_to:
            continue
        disp = _display_math_block([l.lstrip('> ') for l in lines], idx)
        if disp:
            add_display_math(c, disp[0])
            skip_to = disp[1]
            continue
        p = c.paragraphs[0] if idx == 0 and len(c.paragraphs) == 1 and not c.paragraphs[0].text else c.add_paragraph()
        format_paragraph(p, space_before=1, space_after=2, line_spacing=1.15)
        
        if clean_l.startswith('### '):
            clean_l = clean_l[4:]
            col = RGBColor(0xCC, 0x00, 0x00) if is_youtube else RGBColor(0x1B, 0x36, 0x5D)
            format_inline_runs(p, clean_l, default_size=10.5, default_color=col, default_bold=True)
        elif clean_l.startswith('## '):
            clean_l = clean_l[3:]
            format_inline_runs(p, clean_l, default_size=11.0, default_color=RGBColor(0x1B, 0x36, 0x5D), default_bold=True)
        elif clean_l.startswith('# '):
            clean_l = clean_l[2:]
            format_inline_runs(p, clean_l, default_size=11.5, default_color=RGBColor(0x1B, 0x36, 0x5D), default_bold=True)
        else:
            format_inline_runs(p, clean_l, default_size=9.5)
            
    p_sp = doc.add_paragraph()
    format_paragraph(p_sp, space_before=0, space_after=3)

def render_markdown_table(doc, table_lines):
    """Renders markdown tables with Navy header row, alternating zebra stripes, and clean cell padding."""
    rows_data = []
    for line in table_lines:
        line_s = line.strip()
        if not line_s.startswith('|'):
            continue
        if re.match(r'^\|[\s\-:|]+\|$', line_s):
            continue
        cells = [c.strip() for c in line_s.split('|')[1:-1]]
        if cells:
            rows_data.append(cells)
            
    if not rows_data:
        return
        
    num_cols = max(len(r) for r in rows_data)
    num_rows = len(rows_data)
    
    t = doc.add_table(rows=num_rows, cols=num_cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_grid_borders(t, color_hex="D0D7DE", sz="4")
    
    for r_idx, row in enumerate(rows_data):
        is_header = (r_idx == 0)
        for c_idx in range(num_cols):
            cell = t.cell(r_idx, c_idx)
            cell_text = row[c_idx] if c_idx < len(row) else ""
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            
            p = cell.paragraphs[0]
            format_paragraph(p, space_before=1, space_after=1, line_spacing=1.15)
            
            if is_header:
                set_cell_shading(cell, "1B365D")
                format_inline_runs(p, cell_text, default_size=9.5, default_color=RGBColor(0xFF, 0xFF, 0xFF), default_bold=True)
            else:
                bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                set_cell_shading(cell, bg)
                format_inline_runs(p, cell_text, default_size=9.0)
                
    p_sp = doc.add_paragraph()
    format_paragraph(p_sp, space_before=0, space_after=3)

def extract_clean_markdown(text):
    """
    Cleans raw markdown or ChatGPT export files.
    Extracts response bodies, strips chat prompts, timestamps, and model thought logs.
    """
    if '## Prompt:' in text and '## Response:' in text:
        turns = text.split('## Prompt:')
        content_parts = []
        for t in turns[1:]:
            if '## Response:' in t:
                resp = t.split('## Response:')[1]
                lines = resp.splitlines()
                start_idx = 0
                while start_idx < len(lines) and (
                    not lines[start_idx].strip() or 
                    'PM' in lines[start_idx] or 
                    'AM' in lines[start_idx] or 
                    'gpt-' in lines[start_idx] or 
                    'Worked for' in lines[start_idx]
                ):
                    start_idx += 1
                
                end_idx = len(lines)
                while end_idx > start_idx and (
                    not lines[end_idx-1].strip() or 
                    lines[end_idx-1].strip().startswith('[Proceed') or 
                    lines[end_idx-1].strip().startswith('[Compile') or 
                    'Type `Proceed`' in lines[end_idx-1] or 
                    'Type Compile' in lines[end_idx-1]
                ):
                    end_idx -= 1
                
                cleaned_part = '\n'.join(lines[start_idx:end_idx]).strip()
                if cleaned_part:
                    # Ignore final stub confirmation
                    if 'Compilation and QA results' in cleaned_part and 'Download the complete textbook' in cleaned_part:
                        continue
                    content_parts.append(cleaned_part)
        return '\n\n[PAGE BREAK]\n\n'.join(content_parts)
    return text

# ==========================================
# 3B. MATHEMATICS EXTENSIONS (see 07_MATHEMATICS_TEXTBOOK_STANDARDS.md)
#     $$ display formulas $$, ```steps, ```exercises, ```figure, ::: boxes
# ==========================================

_MATH_STATE = {'lesson': None, 'fig': 0}

NAVY_HEX, BLUE_HEX, RED_HEX, TEAL_HEX, GRAY_HEX = '#1B365D', '#2F6FDB', '#E04848', '#0F8C7E', '#5A6270'

def _xml_escape(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

def _true_minus(s):
    """Hyphen-minus used as a sign or operator becomes a real minus sign (−)."""
    s = re.sub(r'(?<![\w)\]])-(?=[\d(a-zA-Z₱])', '−', s)
    return s.replace(' - ', ' − ')

def _match_paren(s, i):
    depth = 0
    for j in range(i, len(s)):
        if s[j] == '(':
            depth += 1
        elif s[j] == ')':
            depth -= 1
            if depth == 0:
                return j
    return len(s) - 1

def _split_top_level(s, sep=','):
    depth = 0
    for i, ch in enumerate(s):
        if ch == '(':
            depth += 1
        elif ch == ')':
            depth -= 1
        elif ch == sep and depth == 0:
            return s[:i].strip(), s[i + 1:].strip()
    return s.strip(), ''

def _omml_runs(text, size):
    """Plain text -> OMML runs. Single-letter variables are italic; numbers, words and symbols upright."""
    out = []
    for tok in re.findall(r'(?<![A-Za-z])[A-Za-z](?![A-Za-z])|[A-Za-z]{2,}|[^A-Za-z]+', text):
        italic = len(tok) == 1 and tok.isalpha()
        sty = '' if italic else '<m:rPr><m:sty m:val="p"/></m:rPr>'
        out.append(
            f'<m:r>{sty}<w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/>'
            f'<w:color w:val="1A1A1A"/><w:sz w:val="{int(size * 2)}"/></w:rPr>'
            f'<m:t xml:space="preserve">{_xml_escape(tok)}</m:t></m:r>'
        )
    return ''.join(out)

def _omml(expr, size=12.0):
    """Converts plain Unicode math with frac(a, b), abs(x), sqrt(x) and x^(n) into Word OMML."""
    out, buf, i = [], '', 0
    expr = expr.replace('∣', '|')
    while i < len(expr):
        m = re.match(r'(frac|abs|sqrt)\(', expr[i:])
        if m:
            if buf:
                out.append(_omml_runs(buf, size)); buf = ''
            start = i + len(m.group(1))
            end = _match_paren(expr, start)
            inner = expr[start + 1:end]
            if m.group(1) == 'frac':
                a, b = _split_top_level(inner)
                out.append(f'<m:f><m:num>{_omml(a, size)}</m:num><m:den>{_omml(b, size)}</m:den></m:f>')
            elif m.group(1) == 'abs':
                out.append(f'<m:d><m:dPr><m:begChr m:val="|"/><m:endChr m:val="|"/></m:dPr><m:e>{_omml(inner, size)}</m:e></m:d>')
            else:
                out.append(f'<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/><m:e>{_omml(inner, size)}</m:e></m:rad>')
            i = end + 1
            continue
        if expr[i] in '^_' and i + 1 < len(expr) and expr[i + 1] == '(':
            tag = 'm:sSup' if expr[i] == '^' else 'm:sSub'
            part = 'm:sup' if expr[i] == '^' else 'm:sub'
            bm = re.search(r'([A-Za-z0-9]+|\S)$', buf)
            base = bm.group(1) if bm else ''
            buf = buf[:len(buf) - len(base)]
            if buf:
                out.append(_omml_runs(buf, size)); buf = ''
            end = _match_paren(expr, i + 1)
            out.append(f'<{tag}><m:e>{_omml_runs(base, size)}</m:e><{part}>{_omml(expr[i + 2:end], size)}</{part}></{tag}>')
            i = end + 1
            continue
        buf += expr[i]
        i += 1
    if buf:
        out.append(_omml_runs(buf, size))
    return ''.join(out)

TEX_SYMBOLS = {
    'times': '×', 'div': '÷', 'cdot': '·', 'pm': '±', 'mp': '∓', 'le': '≤', 'leq': '≤', 'ge': '≥', 'geq': '≥',
    'ne': '≠', 'neq': '≠', 'approx': '≈', 'equiv': '≡', 'sim': '∼', 'cong': '≅', 'propto': '∝',
    'lt': '<', 'gt': '>', 'pi': 'π', 'theta': 'θ', 'alpha': 'α', 'beta': 'β', 'gamma': 'γ', 'delta': 'δ',
    'Delta': 'Δ', 'lambda': 'λ', 'mu': 'μ', 'sigma': 'σ', 'Sigma': 'Σ', 'omega': 'ω', 'phi': 'φ',
    'infty': '∞', 'circ': '°', 'degree': '°', 'cdots': '⋯', 'ldots': '…', 'dots': '…', 'to': '→',
    'rightarrow': '→', 'leftarrow': '←', 'Rightarrow': '⇒', 'implies': '⇒', 'iff': '⇔', 'angle': '∠',
    'triangle': '△', 'perp': '⊥', 'parallel': '∥', 'in': '∈', 'notin': '∉', 'cup': '∪', 'cap': '∩',
    'subset': '⊂', 'emptyset': '∅', 'therefore': '∴', 'because': '∵', 'lvert': '|', 'rvert': '|',
    'vert': '|', 'mid': '|', '%': '%', '{': '{', '}': '}', '$': '$', '#': '#', '&': '&', '_': '_',
    ',': ' ', ';': ' ', ':': ' ', '!': '', ' ': ' ', 'quad': '  ', 'qquad': '    ',
    'displaystyle': '', 'textstyle': '', 'limits': '', 'nolimits': '',
}
TEX_TEXT_CMDS = ('text', 'textrm', 'textbf', 'textit', 'mathrm', 'mathbf', 'mathit', 'operatorname', 'mbox', 'boxed',
                 'overline', 'underline', 'mathbb', 'mathcal')
TEX_SIZERS = ('left', 'right', 'big', 'Big', 'bigg', 'Bigg', 'bigl', 'bigr', 'Bigl', 'Bigr')

def _tex_group(s, i):
    # Returns (content, next_index) for a {group}, or a single token, starting at s[i].
    while i < len(s) and s[i] == ' ':
        i += 1
    if i >= len(s):
        return '', i
    if s[i] == '{':
        depth = 0
        for j in range(i, len(s)):
            if s[j] == '{':
                depth += 1
            elif s[j] == '}':
                depth -= 1
                if depth == 0:
                    return s[i + 1:j], j + 1
        return s[i + 1:], len(s)
    if s[i] == '\\':
        m = re.match(r'\\([A-Za-z]+|.)', s[i:])
        return m.group(0), i + len(m.group(0))
    return s[i], i + 1

def tex_to_math(s):
    # Converts LaTeX (as ChatGPT writes it) into the builder's math syntax: frac(), sqrt(), ^(), _().
    out, i = '', 0
    while i < len(s):
        ch = s[i]
        if ch == '\\':
            if s[i:i + 2] == '\\\\':
                out += '\n'; i += 2; continue
            m = re.match(r'\\([A-Za-z]+|.)', s[i:])
            name = m.group(1); i += len(m.group(0))
            if name in ('frac', 'dfrac', 'tfrac'):
                a, i = _tex_group(s, i); b, i = _tex_group(s, i)
                out += f'frac({tex_to_math(a)}, {tex_to_math(b)})'
            elif name == 'sqrt':
                if i < len(s) and s[i] == '[':
                    i = s.index(']', i) + 1
                a, i = _tex_group(s, i)
                out += f'sqrt({tex_to_math(a)})'
            elif name in TEX_TEXT_CMDS:
                a, i = _tex_group(s, i)
                out += tex_to_math(a) if name.startswith('math') or name in ('boxed', 'overline', 'underline') else a
            elif name in TEX_SIZERS:
                if i < len(s) and s[i] == '.':
                    i += 1
            elif name in ('begin', 'end'):
                _, i = _tex_group(s, i)
                if i < len(s) and s[i] == '{':
                    _, i = _tex_group(s, i)
            elif name in TEX_SYMBOLS:
                out += TEX_SYMBOLS[name]
            else:
                out += name
            continue
        if ch in '^_':
            g, i = _tex_group(s, i + 1)
            g = tex_to_math(g)
            out += f'{ch}({g})'
            continue
        if ch in '{}':
            i += 1
            continue
        out += ch
        i += 1
    return out

def _prep_math(expr):
    expr = tex_to_math(expr.strip()).replace('&', '')
    return _true_minus(clean_math_text(expr))

def add_inline_math(paragraph, expr, size=11.0):
    paragraph._p.append(parse_xml(f'<m:oMath {nsdecls("m", "w")}>{_omml(_prep_math(expr), size)}</m:oMath>'))

def _display_math_block(lines, i):
    # Returns (expression, next_line_index) when lines[i] opens a $$ or \[ display block, else None.
    s = lines[i].strip()
    for op, cl in (('$$', '$$'), ('\\[', '\\]')):
        if not s.startswith(op):
            continue
        if len(s) > len(op) + len(cl) - 1 and s.endswith(cl) and s != op:
            return s[len(op):-len(cl)], i + 1
        body, j = [s[len(op):]], i + 1
        while j < len(lines):
            t = lines[j].strip()
            if t.endswith(cl):
                body.append(t[:-len(cl)])
                return '\n'.join(b for b in body if b.strip()), j + 1
            body.append(t)
            j += 1
    return None

def add_display_math(container, expr, size=12.5):
    """Centered display formula on its own line, typeset as real Word math."""
    if '\\\\' in expr or '\n' in expr.strip():
        rows = []
        for line in re.split(r'\\\\|\n', expr):
            line = re.sub(r'\\(begin|end)\{[^}]*\}(\{[^}]*\})?', '', line).strip()
            if not line:
                continue
            body, _, why = line.partition('&&')
            body = tex_to_math(body).replace('&', '').strip()
            why = re.sub(r'^\(|\)$', '', tex_to_math(why).strip())
            rows.append(body + (f' | {why}' if why else ''))
        return add_steps(container, rows, size=size - 1)
    p = container.add_paragraph()
    format_paragraph(p, space_before=3, space_after=5, align=WD_ALIGN_PARAGRAPH.CENTER)
    p._p.append(parse_xml(
        f'<m:oMathPara {nsdecls("m", "w")}><m:oMathParaPr><m:jc m:val="center"/></m:oMathParaPr>'
        f'<m:oMath>{_omml(_prep_math(expr), size)}</m:oMath></m:oMathPara>'
    ))
    return p

def _top_level_eq(s):
    depth = 0
    for i, ch in enumerate(s):
        if ch == '(':
            depth += 1
        elif ch == ')':
            depth -= 1
        elif ch == '=' and depth == 0:
            return i
    return -1

def add_steps(container, lines, size=11.0):
    """Solution steps aligned at '=' with the reason in a third column."""
    rows = []
    for raw in lines:
        s = raw.strip().strip('$')
        if not s:
            continue
        expr, reason = s, ''
        if ' | ' in s:
            a, b = s.rsplit(' | ', 1)
            if re.search(r'[A-Za-z]', b):
                expr, reason = a.strip(), b.strip()
        if expr.startswith('='):
            lhs, rhs = '', expr[1:].strip()
        else:
            k = _top_level_eq(expr)
            lhs, rhs = (expr[:k].strip(), expr[k + 1:].strip()) if k >= 0 else (expr, None)
        rows.append((lhs, rhs, reason))
    if not rows:
        return
    t = container.add_table(rows=len(rows), cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_table_no_borders(t)
    for r, (lhs, rhs, reason) in zip(t.rows, rows):
        for cell, w in zip(r.cells, (2.1, 2.3, 2.1)):
            cell.width = Inches(w)
        p = r.cells[0].paragraphs[0]
        format_paragraph(p, space_before=1, space_after=1, align=WD_ALIGN_PARAGRAPH.RIGHT)
        if lhs:
            add_inline_math(p, lhs, size)
        p = r.cells[1].paragraphs[0]
        format_paragraph(p, space_before=1, space_after=1)
        if rhs is not None:
            add_inline_math(p, '= ' + rhs, size)
        p = r.cells[2].paragraphs[0]
        format_paragraph(p, space_before=1, space_after=1)
        if reason:
            format_inline_runs(p, reason, default_size=9.0, default_color=RGBColor(0x5A, 0x62, 0x70), default_italic=True)
    format_paragraph(container.add_paragraph(), space_before=0, space_after=2)

def add_exercises(container, lines, cols=2, start=1, size=10.5):
    """Numbered exercise grid (borderless table) like a math-book practice set."""
    items = [re.sub(r'^\s*(\d+[.)]|[-*•])\s+', '', l).strip() for l in lines if l.strip()]
    if not items:
        return
    cols = max(1, min(cols, 4))
    nrows = (len(items) + cols - 1) // cols
    t = container.add_table(rows=nrows, cols=cols)
    t.autofit = False
    set_table_no_borders(t)
    for idx, item in enumerate(items):
        cell = t.rows[idx // cols].cells[idx % cols]
        cell.width = Inches(6.6 / cols)
        p = cell.paragraphs[0]
        format_paragraph(p, space_before=2, space_after=3)
        r = p.add_run(f"{start + idx}.  ")
        r.font.name = 'Cambria'; r.font.size = Pt(size); r.font.bold = True
        r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        format_inline_runs(p, _true_minus(item), default_size=size)
    format_paragraph(container.add_paragraph(), space_before=0, space_after=3)

BOX_STYLES = {
    'concept':  ("EAF2FB", "1B365D", "📐 "),
    'example':  ("F7F9FC", "2A4B7C", "💡 "),
    'watchout': ("FFF7E6", "C77700", "⚠️ "),
    'note':     ("F4F7FA", "1B365D", ""),
}

def add_box(container, kind, title, inner_markdown):
    """Boxed Key Concept / Worked Example / Watch Out containing full markdown (figures, steps, math)."""
    bg, border, icon = BOX_STYLES.get(kind, BOX_STYLES['note'])
    t = container.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    set_cell_shading(c, bg)
    set_callout_border(c, color_hex=border, sz="36")
    set_cell_margins(c, top=120, bottom=120, left=200, right=200)
    p = c.paragraphs[0]
    format_paragraph(p, space_before=1, space_after=4)
    p.paragraph_format.keep_with_next = True
    if title:
        col = RGBColor(int(border[0:2], 16), int(border[2:4], 16), int(border[4:6], 16))
        format_inline_runs(p, icon + title, default_size=11.0, default_color=col, default_bold=True)
    append_markdown_to_doc(c, inner_markdown)
    format_paragraph(container.add_paragraph(), space_before=0, space_after=4)

def _parse_figure_spec(lines):
    spec = {}
    for l in lines:
        if ':' not in l:
            continue
        k, v = l.split(':', 1)
        spec.setdefault(k.strip().lower(), []).append(v.strip())
    return spec

def _nums(s):
    return [float(x) for x in re.findall(r'-?\d+(?:\.\d+)?', s.replace('−', '-'))]

def _fmt(v):
    v = float(v)
    s = str(int(v)) if v == int(v) else ('%g' % v)
    return s.replace('-', '−')

def _label(s):
    return _true_minus(s.replace('−', '-')) if s else s

def _fig_numberline(plt, spec, vertical=False):
    from matplotlib.patches import FancyArrowPatch
    from matplotlib.path import Path
    lo, hi, *rest = _nums(spec.get('range', ['-6, 6'])[0])
    step = rest[0] if rest else 1
    hops = [h.split(',', 2) for h in spec.get('hop', [])]
    n_h = len(hops)
    if vertical:
        fig, ax = plt.subplots(figsize=(2.6 + 0.5 * n_h, 0.42 * ((hi - lo) / step + 2)))
    else:
        fig, ax = plt.subplots(figsize=(6.6, 1.2 + 0.42 * max(n_h, 1)))
    P = (lambda a, o: (o, a)) if vertical else (lambda a, o: (a, o))
    ext = 0.6 * step
    ax.annotate('', xy=P(hi + ext, 0), xytext=P(lo - ext, 0), annotation_clip=False,
                arrowprops=dict(arrowstyle='<|-|>', lw=1.6, color=NAVY_HEX, mutation_scale=12))
    ax.plot(*zip(P(lo - ext, 0), P(hi + ext, 0)), alpha=0)
    v, unit = lo, step * 0.12
    while v <= hi + 1e-9:
        ax.plot(*zip(P(v, -unit), P(v, unit)), color=NAVY_HEX, lw=1.4)
        kw = dict(ha='right', va='center') if vertical else dict(ha='center', va='top')
        ax.text(*P(v, -2.6 * unit), _fmt(v), fontsize=10, color='#1A1A1A',
                fontweight='bold' if v == 0 else 'normal', **kw)
        v += step
    if spec.get('zero'):
        ax.text(*P(0, 2.2 * unit), spec['zero'][0], fontsize=9, color=GRAY_HEX, style='italic',
                **(dict(ha='left', va='bottom') if vertical else dict(ha='center', va='bottom')))
    for k, h in enumerate(hops):
        s, e = float(h[0].replace('−', '-')), float(h[1].replace('−', '-'))
        lab = _label(h[2].strip()) if len(h) > 2 else ''
        height = step * (0.9 + 0.75 * k)
        path = Path([P(s, unit), P((s + e) / 2, 2 * height), P(e, unit)], [Path.MOVETO, Path.CURVE3, Path.CURVE3])
        ax.add_patch(FancyArrowPatch(path=path, arrowstyle='-|>', mutation_scale=13, lw=1.6,
                                     color=RED_HEX if e < s else BLUE_HEX))
        if lab:
            kw = dict(ha='left', va='center') if vertical else dict(ha='center', va='bottom')
            ax.text(*P((s + e) / 2, height + 0.15 * step), lab, fontsize=10, fontweight='bold',
                    color=RED_HEX if e < s else BLUE_HEX, **kw)
    for ptxt in spec.get('point', []):
        parts = ptxt.split(',', 1)
        pv = float(parts[0].replace('−', '-'))
        ax.plot(*P(pv, 0), 'o', ms=8, color=TEAL_HEX, zorder=5, mec='white', mew=1)
        if len(parts) > 1 and parts[1].strip():
            off = -5.2 * unit if (n_h or vertical) else 2.4 * unit
            kw = dict(ha='right', va='center') if vertical else dict(ha='center', va='top' if n_h else 'bottom')
            if vertical:
                off = -14 * unit
            ax.text(*P(pv, off), _label(parts[1].strip()), fontsize=9, color=TEAL_HEX, fontweight='bold', **kw)
    for dtxt in spec.get('distance', []):
        parts = dtxt.split(',', 2)
        a, b = float(parts[0].replace('−', '-')), float(parts[1].replace('−', '-'))
        o = -7.5 * unit if not vertical else 3 * unit + (0.75 * step if n_h else 0) * 0
        if vertical:
            o = -9.5 * unit
        ax.annotate('', xy=P(b, o), xytext=P(a, o), annotation_clip=False,
                    arrowprops=dict(arrowstyle='<->', lw=1.2, color=GRAY_HEX))
        ax.plot(*zip(P(a, o), P(b, o)), alpha=0)
        if len(parts) > 2:
            kw = dict(ha='right', va='center') if vertical else dict(ha='center', va='top')
            ax.text(*P((a + b) / 2, o - 1.2 * unit), _label(parts[2].strip()), fontsize=9, color=GRAY_HEX, **kw)
    ax.set_aspect('equal', adjustable='datalim')
    ax.axis('off')
    return fig, (2.4 + 0.45 * n_h if vertical else 6.0)

def _counter(ax, x, y, positive, r=0.36):
    from matplotlib.patches import Circle
    ax.add_patch(Circle((x, y), r, fc=BLUE_HEX if positive else RED_HEX, ec='white', lw=1.5, zorder=3))
    ax.text(x, y - 0.02, '+' if positive else '−', ha='center', va='center', color='white',
            fontsize=15, fontweight='bold', zorder=4)

def _fig_counters(plt, spec):
    from matplotlib.patches import FancyBboxPatch
    pos = int(_nums(spec.get('positive', ['0'])[0])[0])
    neg = int(_nums(spec.get('negative', ['0'])[0])[0])
    pairs = min(pos, neg) if spec.get('pairs', ['no'])[0].lower().startswith('y') else 0
    cols = max(pos, neg, 1)
    fig, ax = plt.subplots(figsize=(0.62 * cols + 1.2, 2.4))
    for i in range(pos):
        _counter(ax, i, 1, True)
    for i in range(neg):
        _counter(ax, i, 0, False)
    for i in range(pairs):
        ax.add_patch(FancyBboxPatch((i - 0.44, -0.44), 0.88, 1.88, boxstyle='round,pad=0.02,rounding_size=0.3',
                                    fc='none', ec=GRAY_HEX, lw=1.2, ls='--', zorder=2))
    for rtxt in spec.get('remove', []):
        n = int(_nums(rtxt)[0]); row = 1 if 'pos' in rtxt.lower() else 0
        count = pos if row == 1 else neg
        for i in range(max(0, count - n), count):
            ax.plot([i - 0.35, i + 0.35], [row - 0.35, row + 0.35], color='#222', lw=2.2, zorder=5)
            ax.plot([i - 0.35, i + 0.35], [row + 0.35, row - 0.35], color='#222', lw=2.2, zorder=5)
    _counter(ax, -0.2, -1.05, True, r=0.2); ax.text(0.1, -1.05, '= +1', va='center', fontsize=9.5)
    _counter(ax, 1.4, -1.05, False, r=0.2); ax.text(1.7, -1.05, '= −1', va='center', fontsize=9.5)
    if pairs:
        ax.add_patch(FancyBboxPatch((3.0, -1.3), 0.5, 0.5, boxstyle='round,pad=0.02,rounding_size=0.15',
                                    fc='none', ec=GRAY_HEX, lw=1.1, ls='--'))
        ax.text(3.7, -1.05, '= zero pair', va='center', fontsize=9.5, color=GRAY_HEX)
    ax.set_xlim(-0.8, cols + 0.2); ax.set_ylim(-1.4, 1.6)
    ax.set_aspect('equal'); ax.axis('off')
    return fig, min(0.5 * cols + 1.6, 5.6)

def _fig_groups(plt, spec):
    from matplotlib.patches import FancyBboxPatch
    g = int(_nums(spec.get('groups', ['3'])[0])[0])
    each = int(_nums(spec.get('each', ['2'])[0])[0])
    k, per_row = abs(each), min(abs(each), 4) or 1
    rows = (k + per_row - 1) // per_row
    w = per_row + 0.6
    fig, ax = plt.subplots(figsize=(g * (w + 0.4) * 0.6 + 0.4, rows * 0.6 + 1.0))
    for gi in range(g):
        x0 = gi * (w + 0.4)
        ax.add_patch(FancyBboxPatch((x0, 0), w, rows + 0.2, boxstyle='round,pad=0.05,rounding_size=0.25',
                                    fc='#F4F7FA', ec=NAVY_HEX, lw=1.2))
        for j in range(k):
            _counter(ax, x0 + 0.8 + (j % per_row), rows - 0.4 - (j // per_row), each > 0)
        ax.text(x0 + w / 2, -0.35, f'group {gi + 1}', ha='center', va='top', fontsize=8.5, color=GRAY_HEX)
    ax.set_xlim(-0.2, g * (w + 0.4)); ax.set_ylim(-0.9, rows + 0.5)
    ax.set_aspect('equal'); ax.axis('off')
    return fig, min(1.1 * g * (w + 0.4) * 0.6 + 0.6, 5.8)

def _style_axes(ax, spec):
    ax.spines[['top', 'right']].set_visible(False)
    ax.yaxis.grid(True, color='#E3E8EF', lw=0.8); ax.set_axisbelow(True)
    if spec.get('xlabel'): ax.set_xlabel(spec['xlabel'][0], fontsize=10)
    if spec.get('ylabel'): ax.set_ylabel(spec['ylabel'][0], fontsize=10)
    if spec.get('title'): ax.set_title(spec['title'][0], fontsize=11, fontweight='bold', color=NAVY_HEX)

def _labels_values(spec):
    labels = [x.strip() for x in spec.get('labels', [''])[0].split(',')]
    values = _nums(spec.get('values', ['0'])[0])
    return labels, values

def _fig_bar(plt, spec):
    from matplotlib.ticker import MaxNLocator
    labels, values = _labels_values(spec)
    fig, ax = plt.subplots(figsize=(5.2, 3.0))
    bars = ax.bar(labels, values, color=NAVY_HEX, width=0.6)
    for b, v in zip(bars, values):
        ax.text(b.get_x() + b.get_width() / 2, v, _fmt(v), ha='center', va='bottom', fontsize=9)
    ax.set_ylim(0, max(values) * 1.15 if values else 1)
    ax.yaxis.set_major_locator(MaxNLocator(integer=True))
    _style_axes(ax, spec)
    return fig, 4.6

def _fig_line(plt, spec):
    from matplotlib.ticker import MaxNLocator
    labels, values = _labels_values(spec)
    fig, ax = plt.subplots(figsize=(5.2, 3.0))
    ax.plot(labels, values, color=NAVY_HEX, lw=2, marker='o', mfc=TEAL_HEX, mec='white', ms=8)
    for x, v in zip(labels, values):
        ax.annotate(_fmt(v), (x, v), textcoords='offset points', xytext=(0, 7), ha='center', fontsize=9)
    lo = min(values + [0])
    ax.set_ylim(lo, max(values) * 1.15 if values else 1)
    ax.yaxis.set_major_locator(MaxNLocator(integer=True))
    _style_axes(ax, spec)
    return fig, 4.6

def _fig_pie(plt, spec):
    labels, values = _labels_values(spec)
    total = sum(values) or 1
    show = spec.get('angles', ['no'])[0].lower().startswith('y')
    names = [f"{l}\n{_fmt(round(v / total * 360, 1))}°" if show else l for l, v in zip(labels, values)]
    colors = [NAVY_HEX, TEAL_HEX, '#F2A93B', '#8E6CCF', RED_HEX, '#7FA7D9']
    fig, ax = plt.subplots(figsize=(3.8, 3.4))
    ax.pie(values, labels=names, colors=colors[:len(values)], startangle=90, counterclock=False,
           wedgeprops=dict(ec='white', lw=2), textprops=dict(fontsize=9.5))
    ax.set_aspect('equal')
    if spec.get('title'):
        ax.set_title(spec['title'][0], fontsize=11, fontweight='bold', color=NAVY_HEX)
    return fig, 3.4

def _fig_stemleaf(plt, spec):
    data = sorted(int(x) for x in _nums(spec.get('data', [''])[0]))
    stems = {}
    for d in data:
        stems.setdefault(d // 10, []).append(d % 10)
    rows = list(range(min(stems), max(stems) + 1)) if stems else []
    fig, ax = plt.subplots(figsize=(3.0, 0.42 * len(rows) + 1.0))
    ax.text(0.9, len(rows) + 0.2, 'Stem', ha='right', fontsize=10, fontweight='bold', color=NAVY_HEX)
    ax.text(1.3, len(rows) + 0.2, 'Leaf', ha='left', fontsize=10, fontweight='bold', color=NAVY_HEX)
    for k, s in enumerate(rows):
        y = len(rows) - 1 - k
        ax.text(0.9, y, str(s), ha='right', va='center', fontsize=11)
        ax.text(1.3, y, '  '.join(str(l) for l in stems.get(s, [])), ha='left', va='center', fontsize=11)
    ax.plot([1.1, 1.1], [-0.5, len(rows) - 0.4], color=NAVY_HEX, lw=1.5)
    if spec.get('key'):
        ax.text(0.0, -1.0, 'Key: ' + spec['key'][0], fontsize=9, style='italic', color=GRAY_HEX)
    ax.set_xlim(0, 4); ax.set_ylim(-1.3, len(rows) + 0.7); ax.axis('off')
    return fig, 2.6

def _fig_coordinate(plt, spec):
    lo, hi = _nums(spec.get('range', ['-5, 5'])[0])[:2]
    fig, ax = plt.subplots(figsize=(3.8, 3.8))
    ax.set_xlim(lo - 0.5, hi + 0.5); ax.set_ylim(lo - 0.5, hi + 0.5)
    ax.set_xticks(range(int(lo), int(hi) + 1)); ax.set_yticks(range(int(lo), int(hi) + 1))
    ax.grid(True, color='#E3E8EF', lw=0.8)
    for s in ('left', 'bottom'):
        ax.spines[s].set_position('zero'); ax.spines[s].set_color(NAVY_HEX)
    ax.spines[['top', 'right']].set_visible(False)
    ax.tick_params(labelsize=8)
    for segtxt in spec.get('segment', []):
        x1, y1, x2, y2 = _nums(segtxt)[:4]
        ax.plot([x1, x2], [y1, y2], color=NAVY_HEX, lw=1.8)
    for ptxt in spec.get('point', []):
        nums = _nums(ptxt); lab = re.sub(r'^[\s\-−\d.,]+', '', ptxt).strip(' ,')
        ax.plot(nums[0], nums[1], 'o', color=TEAL_HEX, ms=7, zorder=5)
        if lab:
            left = nums[0] < 0
            ax.annotate(f"{lab}({_fmt(nums[0])}, {_fmt(nums[1])})", (nums[0], nums[1]), textcoords='offset points',
                        xytext=(-6, 6) if left else (6, 6), ha='right' if left else 'left',
                        fontsize=9, fontweight='bold', color=TEAL_HEX,
                        bbox=dict(boxstyle='round,pad=0.15', fc='white', ec='none', alpha=0.85))
    ax.set_aspect('equal')
    return fig, 3.4

def _fig_polygon(plt, spec):
    from matplotlib.patches import Polygon
    verts = re.findall(r'([A-Z])\(\s*(-?[\d.]+)\s*,\s*(-?[\d.]+)\s*\)', spec.get('vertices', [''])[0])
    names = [v[0] for v in verts]; pts = [(float(v[1]), float(v[2])) for v in verts]
    cx = sum(p[0] for p in pts) / len(pts); cy = sum(p[1] for p in pts) / len(pts)
    fig, ax = plt.subplots(figsize=(4.0, 3.2))
    ax.add_patch(Polygon(pts, closed=True, fc='#EAF2FB', ec=NAVY_HEX, lw=2))
    span = max(max(p[0] for p in pts) - min(p[0] for p in pts), max(p[1] for p in pts) - min(p[1] for p in pts)) or 1
    for n, (x, y) in zip(names, pts):
        dx, dy = x - cx, y - cy; d = (dx * dx + dy * dy) ** 0.5 or 1
        ax.text(x + dx / d * span * 0.07, y + dy / d * span * 0.07, n, ha='center', va='center', fontsize=11, fontweight='bold')
    for stxt in spec.get('side', []):
        seg, _, lab = stxt.partition(',')
        a, b = names.index(seg.strip()[0]), names.index(seg.strip()[1])
        (x1, y1), (x2, y2) = pts[a], pts[b]
        mx, my = (x1 + x2) / 2, (y1 + y2) / 2; dx, dy = mx - cx, my - cy; d = (dx * dx + dy * dy) ** 0.5 or 1
        ax.text(mx + dx / d * span * 0.13, my + dy / d * span * 0.13, lab.strip(), ha='center', va='center', fontsize=10, color=GRAY_HEX)
    for rtxt in spec.get('right', []):
        i = names.index(rtxt.strip()[0]); x, y = pts[i]
        (ax_, ay_), (bx_, by_) = pts[i - 1], pts[(i + 1) % len(pts)]
        s = span * 0.08
        u = ((ax_ - x), (ay_ - y)); v = ((bx_ - x), (by_ - y))
        lu = (u[0] ** 2 + u[1] ** 2) ** 0.5; lv = (v[0] ** 2 + v[1] ** 2) ** 0.5
        u = (u[0] / lu * s, u[1] / lu * s); v = (v[0] / lv * s, v[1] / lv * s)
        ax.plot([x + u[0], x + u[0] + v[0], x + v[0]], [y + u[1], y + u[1] + v[1], y + v[1]], color=NAVY_HEX, lw=1.3)
    ax.set_xlim(min(p[0] for p in pts) - span * 0.2, max(p[0] for p in pts) + span * 0.2)
    ax.set_ylim(min(p[1] for p in pts) - span * 0.2, max(p[1] for p in pts) + span * 0.2)
    ax.set_aspect('equal'); ax.axis('off')
    return fig, 3.2

FIGURE_TYPES = {
    'numberline': lambda plt, s: _fig_numberline(plt, s),
    'vnumberline': lambda plt, s: _fig_numberline(plt, s, vertical=True),
    'counters': _fig_counters, 'groups': _fig_groups, 'bar': _fig_bar, 'line': _fig_line,
    'pie': _fig_pie, 'stemleaf': _fig_stemleaf, 'coordinate': _fig_coordinate, 'polygon': _fig_polygon,
}

def add_figure(container, spec_lines, max_width=5.8):
    """Draws an exact math figure with matplotlib, inserts it centered, and adds a numbered caption."""
    import tempfile
    spec = _parse_figure_spec(spec_lines)
    _MATH_STATE['fig'] += 1
    num = f"{_MATH_STATE['lesson']}.{_MATH_STATE['fig']}" if _MATH_STATE['lesson'] else str(_MATH_STATE['fig'])
    kind = spec.get('type', ['?'])[0].lower()
    p = container.add_paragraph()
    format_paragraph(p, space_before=4, space_after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.keep_with_next = True
    try:
        import matplotlib
        matplotlib.use('Agg')
        import matplotlib.pyplot as plt
        plt.rcParams.update({'font.family': 'DejaVu Sans', 'axes.unicode_minus': True})
        fig, width = FIGURE_TYPES[kind](plt, spec)
        path = os.path.join(tempfile.mkdtemp(), f'figure_{num}.png')
        fig.savefig(path, dpi=220, bbox_inches='tight', pad_inches=0.06, facecolor='white')
        plt.close(fig)
        p.add_run().add_picture(path, width=Inches(min(width, max_width)))
    except Exception as e:
        r = p.add_run(f"[FIGURE {num} COULD NOT BE DRAWN: type '{kind}' — {e}]")
        r.font.bold = True; r.font.color.rgb = RGBColor(0xCC, 0x00, 0x00); r.font.size = Pt(9)
    cap = container.add_paragraph()
    format_paragraph(cap, space_before=0, space_after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = cap.add_run(f"Figure {num}")
    r.font.name = 'Cambria'; r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    if spec.get('caption'):
        format_inline_runs(cap, '  ' + _true_minus(spec['caption'][0]), default_size=9.0,
                           default_color=RGBColor(0x5A, 0x62, 0x70), default_italic=True)

def append_markdown_to_doc(doc, markdown_text):
    """
    Appends raw or cleaned markdown lines directly into an existing docx.Document.
    Preserves 100% of text, headings, callout boxes, and tables with zero truncation.
    """
    combined_md = extract_clean_markdown(markdown_text)
    lines = combined_md.splitlines()
    i = 0
    total = len(lines)

    while i < total:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if '[PAGE BREAK]' in stripped:
            if hasattr(doc, 'add_page_break'):
                doc.add_page_break()
            i += 1
            continue

        # Math fenced blocks: ```figure / ```steps / ```exercises (file 07)
        fence = re.match(r'^```\s*(figure|steps|exercises)\b(.*)$', stripped)
        if fence:
            kind, args = fence.group(1), fence.group(2)
            block = []
            i += 1
            while i < total and not lines[i].strip().startswith('```'):
                block.append(lines[i])
                i += 1
            i += 1
            if kind == 'figure':
                add_figure(doc, block)
            elif kind == 'steps':
                add_steps(doc, block)
            else:
                cols = re.search(r'cols\s*=\s*(\d+)', args)
                start = re.search(r'start\s*=\s*(\d+)', args)
                add_exercises(doc, block, cols=int(cols.group(1)) if cols else 2, start=int(start.group(1)) if start else 1)
            continue

        # Boxed containers: ::: concept|example|watchout|note Title ... ::: (file 07)
        box = re.match(r'^:::\s*(concept|example|watchout|note)\b\s*(.*)$', stripped)
        if box:
            block = []
            i += 1
            while i < total and lines[i].strip() != ':::':
                block.append(lines[i])
                i += 1
            i += 1
            add_box(doc, box.group(1), box.group(2).strip(), '\n'.join(block))
            continue

        # Display formula: $$ ... $$ or \[ ... \], on one line or across several lines (file 07)
        disp = _display_math_block(lines, i)
        if disp:
            expr, i = disp
            add_display_math(doc, expr)
            continue

        # Blockquote Containers (> ...)
        if stripped.startswith('>'):
            callout_lines = []
            while i < total and lines[i].strip().startswith('>'):
                callout_lines.append(lines[i])
                i += 1
            render_callout_box(doc, callout_lines)
            continue

        # Markdown Tables (| ... |)
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines = []
            while i < total and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i])
                i += 1
            render_markdown_table(doc, table_lines)
            continue

        # Heading 1 (# ...)
        if stripped.startswith('# '):
            h_text = stripped[2:].strip()
            if re.match(r'^PAGE \d+ -', h_text, re.IGNORECASE):
                pass
            else:
                lesson = re.search(r'LESSON\s+(\d+(?:\.\d+)*)', h_text, re.IGNORECASE)
                if lesson:
                    _MATH_STATE.update(lesson=lesson.group(1), fig=0)
                add_heading_1(doc, h_text)
            i += 1
            continue

        # Heading 2 (## ...)
        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            add_heading_2(doc, h_text)
            i += 1
            continue

        # Heading 3 (### ...)
        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            add_heading_3(doc, h_text)
            i += 1
            continue

        # Bulleted Lists (* or - or •)
        if re.match(r'^[\*\-•]\s+', stripped):
            b_text = re.sub(r'^[\*\-•]\s+', '', stripped)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            format_paragraph(p, space_before=1, space_after=2, line_spacing=1.15)
            r_bullet = p.add_run("• ")
            r_bullet.font.name = 'Cambria'
            r_bullet.font.size = Pt(10.5)
            r_bullet.font.bold = True
            r_bullet.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            format_inline_runs(p, b_text, default_size=10.5)
            i += 1
            continue

        # Numbered items / Quiz options ([  ] A., 1., etc.)
        if re.match(r'^(?:\[\s*\]|[0-9]+\.|\([a-z]\))\s+', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            format_paragraph(p, space_before=1, space_after=2, line_spacing=1.15)
            format_inline_runs(p, stripped, default_size=10.5)
            i += 1
            continue

        # Normal Paragraph
        p = doc.add_paragraph()
        format_paragraph(p, space_before=0, space_after=3.5, line_spacing=1.15)
        format_inline_runs(p, stripped, default_size=10.5)
        i += 1

def initialize_document(subject="Mathematics", grade="7", quarter="Quarter 3"):
    """
    Initializes a publication-grade Document with DepEd standard margins,
    running headers, and running footers.
    """
    doc = Document()
    for s in doc.sections:
        s.page_width = Inches(8.5)
        s.page_height = Inches(11.0)
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        s.different_first_page_header_footer = True

        # Running header on pages 2+
        hp = s.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run(f"{subject} {grade} | {quarter}")
        hr.font.name = 'Cambria'
        hr.font.size = Pt(8.5)
        hr.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

        # Running footer on pages 2+
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fr = fp.add_run(f"Grade {grade} {subject} • {quarter} • Page ")
        fr.font.name = 'Cambria'
        fr.font.size = Pt(8.5)
        fr.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
        add_page_number_field(fp)

    return doc

def compile_markdown_to_docx(source_md_path_or_text, output_docx_path, subject="Mathematics", grade="7", quarter="Quarter 3"):
    """
    Main entry point: Compiles any markdown textbook file into a publication-grade DOCX.
    Preserves 100% of text and tables with zero truncation.
    """
    if os.path.exists(source_md_path_or_text):
        with open(source_md_path_or_text, 'r', encoding='utf-8') as f:
            raw_text = f.read()
    else:
        raw_text = source_md_path_or_text

    doc = initialize_document(subject=subject, grade=grade, quarter=quarter)
    _MATH_STATE.update(lesson=None, fig=0)
    append_markdown_to_doc(doc, raw_text)
    doc.save(output_docx_path)
    print(f"Successfully compiled {output_docx_path} ({len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables)")
    return output_docx_path

def append_markdown_to_docx(doc_path_or_doc, markdown_text, output_docx_path=None):
    """
    Progressively appends a markdown lesson chunk to an existing DOCX file.
    Enables multi-turn progressive compilation without token limits.
    """
    if isinstance(doc_path_or_doc, str):
        doc = Document(doc_path_or_doc)
        out_path = output_docx_path or doc_path_or_doc
    else:
        doc = doc_path_or_doc
        out_path = output_docx_path or "MATATAG_Textbook.docx"
        
    append_markdown_to_doc(doc, markdown_text)
    doc.save(out_path)
    return out_path

# ==========================================
# 5. CLI EXECUTION
# ==========================================

if __name__ == '__main__':
    if len(sys.argv) > 1:
        in_path = sys.argv[1]
        out_path = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(in_path)[0] + ".docx"
        compile_markdown_to_docx(in_path, out_path)
    else:
        print("Usage: python 03_DOCX_BUILDER_SCRIPT.py <input_manuscript.md> [output_textbook.docx]")
